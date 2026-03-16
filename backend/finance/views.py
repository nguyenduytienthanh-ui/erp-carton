from decimal import Decimal
from datetime import date, datetime, timedelta
import json

from django.db import transaction
from django.db.models import Q, Sum
from django.http import HttpResponse
from django.utils import timezone
from rest_framework import filters, viewsets
from rest_framework.decorators import action
from rest_framework.exceptions import PermissionDenied, ValidationError
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.filters import SearchFilter, OrderingFilter
from django_filters.rest_framework import DjangoFilterBackend
from unidecode import unidecode
from openpyxl import Workbook

from core.permissions import check_action_permission
from core.models import AuditLog, Notification, Permission, Role, Setting, User, WorkflowPipelineEvent
from workforce.models import Employee, PayrollRecord
from sales.document_policy import round_money

from .models import (
    AdvanceSettlement,
    AdvanceTransaction,
    BankAccount,
    BankReconciliation,
    BankReconciliationStatus,
    CashAccount,
    CashTransaction,
    PayableDocument,
    PayableSettlement,
    PayableStatus,
    ReceivableDocument,
    ReceivableSettlement,
    ReceivableStatus,
    TransactionCategory,
    GeneralLedgerAccount,
    GeneralLedgerEntry,
)
from .services import (
    ensure_system_transaction_category,
    refresh_payable_status,
    refresh_receivable_status,
)
from .reminders import (
    REMINDER_POLICY_SETTING_KEY,
    build_overdue_snapshot,
    build_finance_approval_sla_overview,
    get_approval_sla_policy,
    get_reminder_policy,
    get_reminder_policy_presets,
    run_daily_overdue_reminder_job,
    run_finance_approval_sla_reminder_job,
    save_approval_sla_policy,
    save_reminder_policy,
)
from workforce.reminders import build_salary_advance_approval_sla_overview, run_salary_advance_approval_sla_reminder_job
from workforce.models import SalaryAdvanceRecord
from .serializers import (
    AdvanceSettlementSerializer,
    AdvanceTransactionSerializer,
    BankAccountSerializer,
    BankReconciliationSerializer,
    CashAccountSerializer,
    CashTransactionSerializer,
    PayableDocumentSerializer,
    PayableSettlementSerializer,
    ReceivableDocumentSerializer,
    ReceivableSettlementSerializer,
    TransactionCategorySerializer,
    GeneralLedgerAccountSerializer,
    GeneralLedgerEntrySerializer,
)


def _user_role_names(user):
    try:
        pairs = user.roles.values_list('name', 'code')
    except Exception:
        return set()
    names = set()
    for name, code in pairs:
        if name:
            names.add(str(name).strip().lower())
        if code:
            names.add(str(code).strip().lower())
    return names


def _can_manage_finance(user):
    if getattr(user, 'is_superuser', False) or getattr(user, 'is_staff', False):
        return True
    if check_action_permission(user, 'FINANCE', 'MANAGE', strict=True):
        return True
    role_names = _user_role_names(user)
    allowed_roles = {
        'admin',
        'manager',
        'accountant',
        'finance',
        'finance-manager',
        'quan-ly',
        'quanly',
    }
    return any(role in allowed_roles for role in role_names)


FINANCE_LOCKED_MONTHS_KEY = 'FINANCE_LOCKED_MONTHS'
ADVANCE_APPROVAL_LEVEL2_THRESHOLD_KEY = 'FINANCE_ADVANCE_APPROVAL_LEVEL2_THRESHOLD'
FINANCE_EXECUTIVE_AUTO_POLICY_KEY = 'FINANCE_EXECUTIVE_AUTO_POLICY'
CROSS_MODULE_BOOTSTRAP_AUDIT_ENTITY_TYPE = 'CrossModuleBootstrapRun'
CROSS_MODULE_BOOTSTRAP_AUDIT_ENTITY_CODE = 'CROSS_MODULE_BOOTSTRAP'


def _normalize_month(value: str) -> str:
    raw = (value or '').strip()
    if len(raw) != 7 or '-' not in raw:
        return ''
    y, m = raw.split('-', 1)
    if not (y.isdigit() and m.isdigit()):
        return ''
    y_int = int(y)
    m_int = int(m)
    if y_int < 2000 or y_int > 3000 or m_int < 1 or m_int > 12:
        return ''
    return f'{y_int}-{str(m_int).zfill(2)}'


def _get_next_bank_reconciliation_code(statement_date_value: date | None = None) -> str:
    dt = statement_date_value or timezone.localdate()
    prefix = f'BR{dt.strftime("%Y%m")}'
    latest = (
        BankReconciliation.objects.filter(code__startswith=prefix)
        .order_by('-code')
        .values_list('code', flat=True)
        .first()
    )
    seq = 1
    if latest and latest[-4:].isdigit():
        seq = int(latest[-4:]) + 1
    return f'{prefix}-{seq:04d}'


def _month_from_date_obj(value) -> str:
    if not value:
        return ''
    try:
        return f'{value.year}-{str(value.month).zfill(2)}'
    except Exception:
        return ''


def _get_locked_finance_months() -> set[str]:
    row = Setting.objects.filter(key=FINANCE_LOCKED_MONTHS_KEY, is_active=True).first()
    if not row:
        return set()
    try:
        parsed = json.loads(str(row.value or '{}'))
    except Exception:
        return set()
    months = parsed.get('months') if isinstance(parsed, dict) else []
    if not isinstance(months, list):
        return set()
    return {m for m in (_normalize_month(str(item)) for item in months) if m}


def _save_locked_finance_months(months: set[str]):
    payload = json.dumps({'months': sorted(months)})
    Setting.objects.update_or_create(
        key=FINANCE_LOCKED_MONTHS_KEY,
        defaults={
            'value': payload,
            'data_type': 'json',
            'description': 'Danh sách tháng tài chính đã khóa (YYYY-MM)',
            'is_active': True,
        },
    )


def _ensure_finance_month_unlocked(month: str, message: str):
    normalized = _normalize_month(month)
    if not normalized:
        return
    if normalized in _get_locked_finance_months():
        raise PermissionDenied(message)


def _ensure_finance_months_unlocked(months: list[str], message: str):
    for month in months:
        _ensure_finance_month_unlocked(month, message)


def _log_finance_audit(user, action: str, entity_type: str, entity_id: int, entity_code: str, old_values: dict, new_values: dict, changed_fields: list[str]):
    AuditLog.objects.create(
        user=user,
        action=action,
        entity_type=entity_type,
        entity_id=int(entity_id or 0),
        entity_id_str=str(entity_id or ''),
        entity_code=str(entity_code or ''),
        old_values=old_values if isinstance(old_values, dict) else {},
        new_values=new_values if isinstance(new_values, dict) else {},
        changed_fields=changed_fields if isinstance(changed_fields, list) else [],
    )


def _approved_advance_queryset():
    return AdvanceTransaction.objects.filter(
        is_active=True,
        approval_status=AdvanceTransaction.APPROVAL_APPROVED,
    )


def _validate_settlement_total_under_advance(advance: AdvanceTransaction, spent, refund, exclude_settlement_id: int | None = None) -> str:
    spent_value = Decimal(str(spent or 0))
    refund_value = Decimal(str(refund or 0))
    existing_qs = advance.settlements.all()
    if exclude_settlement_id:
        existing_qs = existing_qs.exclude(pk=exclude_settlement_id)
    existing_totals = existing_qs.aggregate(
        existing_spent=Sum('spent_amount'),
        existing_refund=Sum('refund_amount'),
    )
    existing_total = Decimal(str(existing_totals.get('existing_spent') or 0)) + Decimal(str(existing_totals.get('existing_refund') or 0))
    new_total = existing_total + spent_value + refund_value
    advance_amount = Decimal(str(advance.amount or 0))
    if new_total > advance_amount:
        remaining = max(Decimal('0'), advance_amount - existing_total)
        return (
            f'Tổng quyết toán ({new_total:,.0f}) vượt quá số tiền tạm ứng ({advance_amount:,.0f}). '
            f'Còn được quyết toán tối đa: {remaining:,.0f}.'
        )
    return ''


def _ensure_cash_account_has_available_balance(account: CashAccount | None, amount, *, up_to_date=None, message_prefix: str = 'Quỹ nguồn'):
    if account is None:
        return
    available = account.current_balance_as_of(up_to_date=up_to_date)
    amount_value = Decimal(str(amount or 0))
    if amount_value > available:
        raise ValidationError({'amount': f'{message_prefix} không đủ số dư khả dụng. Khả dụng hiện tại: {available:,.0f}.'})


def _finance_close_check_entry(code: str, severity: str, title: str, message: str, count: int = 0, items: list[dict] | None = None) -> dict:
    return {
        'code': code,
        'severity': severity,
        'title': title,
        'message': message,
        'count': int(count or 0),
        'items': items or [],
    }


def _build_finance_month_close_check(month: str) -> dict:
    month = _normalize_month(month)
    blockers: list[dict] = []
    warnings: list[dict] = []
    if not month:
        return {
            'month': '',
            'is_ready': False,
            'blockers': [
                _finance_close_check_entry(
                    code='INVALID_MONTH',
                    severity='blocker',
                    title='Tháng không hợp lệ',
                    message='month phải có dạng YYYY-MM.',
                )
            ],
            'warnings': [],
        }

    y, m = month.split('-', 1)
    year = int(y)
    month_num = int(m)

    payroll_qs = PayrollRecord.objects.filter(month=month, status=PayrollRecord.STATUS_LOCKED)
    payroll_total = Decimal(str(payroll_qs.aggregate(total=Sum('net_pay')).get('total') or 0))
    tx_qs = CashTransaction.objects.filter(
        transaction_type=CashTransaction.TYPE_EXPENSE,
        transaction_date__year=year,
        transaction_date__month=month_num,
        reason__icontains='[PAYROLL:',
    )
    posted_total = Decimal(str(tx_qs.aggregate(total=Sum('amount')).get('total') or 0))
    reconciliation_delta = payroll_total - posted_total
    if reconciliation_delta != 0:
        blockers.append(_finance_close_check_entry(
            code='PAYROLL_RECONCILIATION_DELTA',
            severity='blocker',
            title='Lệch đối soát lương',
            message='Số tiền lương đã khóa và số tiền đã hạch toán sang tài chính chưa khớp.',
            count=abs(int(reconciliation_delta)),
            items=[{
                'payroll_total': str(payroll_total),
                'posted_total': str(posted_total),
                'delta': str(reconciliation_delta),
            }],
        ))

    pending_l1_qs = AdvanceTransaction.objects.filter(
        approval_status=AdvanceTransaction.APPROVAL_PENDING_L1,
        is_active=True,
        advance_date__year=year,
        advance_date__month=month_num,
    ).order_by('-advance_date', '-id')
    pending_l2_qs = AdvanceTransaction.objects.filter(
        approval_status=AdvanceTransaction.APPROVAL_PENDING_L2,
        is_active=True,
        advance_date__year=year,
        advance_date__month=month_num,
    ).order_by('-advance_date', '-id')
    pending_count = pending_l1_qs.count() + pending_l2_qs.count()
    if pending_count > 0:
        sample_items = [
            {
                'id': int(adv.id),
                'code': adv.code,
                'recipient_name': adv.recipient_name,
                'approval_status': adv.approval_status,
                'amount': str(adv.amount),
            }
            for adv in list(pending_l1_qs[:5]) + list(pending_l2_qs[:5])
        ]
        blockers.append(_finance_close_check_entry(
            code='PENDING_ADVANCE_APPROVALS',
            severity='blocker',
            title='Còn phiếu tạm ứng chờ duyệt',
            message='Cần xử lý xong các phiếu tạm ứng đang chờ duyệt trong tháng trước khi khóa sổ.',
            count=pending_count,
            items=sample_items[:10],
        ))

    open_advances_qs = _approved_advance_queryset().filter(
        advance_date__year=year,
        advance_date__month=month_num,
        status__in=[AdvanceTransaction.STATUS_OPEN, AdvanceTransaction.STATUS_PARTIAL],
    ).order_by('advance_date', 'code')
    if open_advances_qs.exists():
        sample_items = [
            {
                'id': int(adv.id),
                'code': adv.code,
                'recipient_name': adv.recipient_name,
                'status': adv.status,
                'amount': str(adv.amount),
            }
            for adv in open_advances_qs[:10]
        ]
        blockers.append(_finance_close_check_entry(
            code='OPEN_ADVANCES_IN_MONTH',
            severity='blocker',
            title='Còn tạm ứng chưa xử lý hết',
            message='Trong tháng vẫn còn phiếu tạm ứng đã duyệt nhưng chưa quyết toán xong.',
            count=open_advances_qs.count(),
            items=sample_items,
        ))

    overdue_entries = build_overdue_snapshot(as_of=date.today(), threshold_days=30)
    overdue_count = int(overdue_entries.get('count') or 0) if isinstance(overdue_entries, dict) else 0
    if overdue_count > 0:
        warnings.append(_finance_close_check_entry(
            code='OVERDUE_ADVANCES_EXIST',
            severity='warning',
            title='Còn tạm ứng quá hạn',
            message='Hệ thống đang ghi nhận phiếu tạm ứng quá hạn chưa xử lý, nên rà soát trước khi khóa sổ.',
            count=overdue_count,
            items=[{
                'threshold_days': int(overdue_entries.get('threshold_days') or 30),
                'total_remaining': str(overdue_entries.get('total_remaining') or '0'),
            }],
        ))

    return {
        'month': month,
        'is_ready': len(blockers) == 0,
        'blockers': blockers,
        'warnings': warnings,
    }


def _can_approve_level2(user) -> bool:
    if getattr(user, 'is_superuser', False):
        return True
    role_names = _user_role_names(user)
    l2_roles = {
        'admin',
        'manager',
        'finance-manager',
        'finance_director',
        'giam-doc',
        'pho-giam-doc',
    }
    return any(role in l2_roles for role in role_names)


def _get_approval_level2_threshold() -> Decimal:
    row = Setting.objects.filter(key=ADVANCE_APPROVAL_LEVEL2_THRESHOLD_KEY, is_active=True).first()
    if not row:
        return Decimal('50000000')
    try:
        raw = Decimal(str(row.value or '50000000'))
        if raw <= 0:
            return Decimal('50000000')
        return raw
    except Exception:
        return Decimal('50000000')


def _required_approval_level_for_amount(amount: Decimal) -> int:
    threshold = _get_approval_level2_threshold()
    return 2 if Decimal(str(amount or 0)) >= threshold else 1


def _default_executive_auto_policy() -> dict:
    return {
        'enabled': False,
        'cooldown_minutes': 60,
        'auto_run_finance_sla': True,
        'auto_run_workforce_sla': True,
        'only_when_early_warning': True,
        'last_run_at': '',
    }


def _get_executive_auto_policy() -> dict:
    row = Setting.objects.filter(key=FINANCE_EXECUTIVE_AUTO_POLICY_KEY, is_active=True).first()
    if not row:
        return _default_executive_auto_policy()
    try:
        parsed = json.loads(str(row.value or '{}'))
    except Exception:
        return _default_executive_auto_policy()
    if not isinstance(parsed, dict):
        return _default_executive_auto_policy()
    policy = _default_executive_auto_policy()
    policy['enabled'] = bool(parsed.get('enabled', policy['enabled']))
    policy['auto_run_finance_sla'] = bool(parsed.get('auto_run_finance_sla', policy['auto_run_finance_sla']))
    policy['auto_run_workforce_sla'] = bool(parsed.get('auto_run_workforce_sla', policy['auto_run_workforce_sla']))
    policy['only_when_early_warning'] = bool(parsed.get('only_when_early_warning', policy['only_when_early_warning']))
    try:
        policy['cooldown_minutes'] = max(5, min(int(parsed.get('cooldown_minutes', policy['cooldown_minutes'])), 1440))
    except Exception:
        policy['cooldown_minutes'] = 60
    policy['last_run_at'] = str(parsed.get('last_run_at') or '')
    return policy


def _save_executive_auto_policy(policy: dict) -> dict:
    current = _default_executive_auto_policy()
    if isinstance(policy, dict):
        current['enabled'] = bool(policy.get('enabled', current['enabled']))
        current['auto_run_finance_sla'] = bool(policy.get('auto_run_finance_sla', current['auto_run_finance_sla']))
        current['auto_run_workforce_sla'] = bool(policy.get('auto_run_workforce_sla', current['auto_run_workforce_sla']))
        current['only_when_early_warning'] = bool(policy.get('only_when_early_warning', current['only_when_early_warning']))
        try:
            current['cooldown_minutes'] = max(5, min(int(policy.get('cooldown_minutes', current['cooldown_minutes'])), 1440))
        except Exception:
            current['cooldown_minutes'] = 60
        current['last_run_at'] = str(policy.get('last_run_at') or current['last_run_at'])
    Setting.objects.update_or_create(
        key=FINANCE_EXECUTIVE_AUTO_POLICY_KEY,
        defaults={
            'value': json.dumps(current),
            'data_type': 'json',
            'description': 'Policy auto execute cho executive cockpit',
            'is_active': True,
        },
    )
    return current


def _compute_executive_risk_score(finance_sla: dict, workforce_sla: dict, overdue_90: dict) -> dict:
    risk_score = int(
        (int(finance_sla.get('overdue_l1_count') or 0) * 2)
        + (int(finance_sla.get('overdue_l2_count') or 0) * 3)
        + (int(workforce_sla.get('overdue_l1_count') or 0) * 2)
        + (int(workforce_sla.get('overdue_l2_count') or 0) * 3)
        + (int(overdue_90.get('count') or 0) * 2)
    )
    risk_level = 'HIGH' if risk_score >= 25 else ('MEDIUM' if risk_score >= 10 else 'LOW')
    early_warning_threshold = 25 if risk_level == 'MEDIUM' else (10 if risk_level == 'LOW' else 25)
    score_to_next = max(0, int(early_warning_threshold - risk_score))
    early_warning = {
        'is_triggered': bool(risk_level != 'HIGH' and score_to_next <= 3),
        'score_to_next_level': int(score_to_next),
        'target_level': 'HIGH' if risk_level == 'MEDIUM' else ('MEDIUM' if risk_level == 'LOW' else 'HIGH'),
        'hint': (
            'Rủi ro đang tiệm cận ngưỡng cấp kế tiếp, cần chạy nhắc SLA và review nhóm tắc nghẽn.'
            if risk_level != 'HIGH' and score_to_next <= 3
            else 'Rủi ro còn khoảng đệm an toàn.'
        ),
    }
    return {'risk_score': risk_score, 'risk_level': risk_level, 'early_warning': early_warning}


def _parse_iso_datetime(raw_value: str):
    if not raw_value:
        return None
    try:
        value = datetime.fromisoformat(str(raw_value))
        if timezone.is_naive(value):
            value = timezone.make_aware(value, timezone.get_current_timezone())
        return value
    except Exception:
        return None


def _execute_executive_auto(actor=None, force_run: bool = False, source: str = 'api') -> dict:
    policy = _get_executive_auto_policy()
    now_dt = timezone.now()
    if not bool(policy.get('enabled')) and not force_run:
        result = {'success': False, 'skipped': True, 'reason': 'POLICY_DISABLED', 'policy': policy}
    else:
        cooldown_minutes = int(policy.get('cooldown_minutes') or 60)
        last_run_at = _parse_iso_datetime(str(policy.get('last_run_at') or ''))
        if last_run_at and not force_run:
            seconds_since_last = (now_dt - last_run_at).total_seconds()
            if seconds_since_last < cooldown_minutes * 60:
                result = {
                    'success': False,
                    'skipped': True,
                    'reason': 'COOLDOWN_ACTIVE',
                    'remaining_seconds': int(cooldown_minutes * 60 - seconds_since_last),
                    'policy': policy,
                }
            else:
                result = None
        else:
            result = None
        if result is None:
            finance_sla = build_finance_approval_sla_overview()
            workforce_sla = build_salary_advance_approval_sla_overview()
            overdue_90 = build_overdue_snapshot(as_of=timezone.localdate(), threshold_days=90)
            risk_snapshot = _compute_executive_risk_score(finance_sla, workforce_sla, overdue_90)
            early_warning = risk_snapshot.get('early_warning') if isinstance(risk_snapshot, dict) else {}
            if bool(policy.get('only_when_early_warning')) and not bool((early_warning or {}).get('is_triggered')) and not force_run:
                result = {'success': False, 'skipped': True, 'reason': 'EARLY_WARNING_NOT_TRIGGERED', 'policy': policy}
            else:
                finance_result = {'success': False, 'sent_count': 0}
                workforce_result = {'success': False, 'sent_count': 0}
                if bool(policy.get('auto_run_finance_sla')):
                    finance_result = run_finance_approval_sla_reminder_job(dry_run=False)
                if bool(policy.get('auto_run_workforce_sla')):
                    workforce_result = run_salary_advance_approval_sla_reminder_job(dry_run=False)
                policy['last_run_at'] = now_dt.isoformat()
                saved_policy = _save_executive_auto_policy(policy)
                result = {
                    'success': True,
                    'skipped': False,
                    'policy': saved_policy,
                    'finance_result': finance_result,
                    'workforce_result': workforce_result,
                    'risk_snapshot': risk_snapshot,
                }
    _log_finance_audit(
        actor,
        action='UPDATE',
        entity_type='FinanceExecutiveAutoExecute',
        entity_id=0,
        entity_code='FINANCE_EXECUTIVE_AUTO_EXECUTE',
        old_values={},
        new_values={
            'source': source,
            'force_run': bool(force_run),
            'success': bool(result.get('success')),
            'skipped': bool(result.get('skipped')),
            'reason': str(result.get('reason') or ''),
            'finance_sent_count': int((result.get('finance_result') or {}).get('sent_count') or 0),
            'workforce_sent_count': int((result.get('workforce_result') or {}).get('sent_count') or 0),
            'policy_snapshot': result.get('policy') if isinstance(result.get('policy'), dict) else policy,
        },
        changed_fields=['source', 'force_run', 'success', 'skipped', 'reason', 'finance_sent_count', 'workforce_sent_count'],
    )
    return result


def run_executive_auto_execute_job() -> dict:
    return _execute_executive_auto(actor=None, force_run=False, source='scheduler')


def _executive_auto_history_qs():
    return (
        AuditLog.objects
        .filter(entity_type='FinanceExecutiveAutoExecute', entity_code='FINANCE_EXECUTIVE_AUTO_EXECUTE')
        .select_related('user')
    )


def _serialize_executive_auto_history_row(row: AuditLog) -> dict:
    payload = row.new_values if isinstance(row.new_values, dict) else {}
    return {
        'id': int(row.id),
        'created_at': row.created_at.isoformat() if row.created_at else '',
        'username': str(row.user.username) if row.user else '',
        'source': str(payload.get('source') or ''),
        'force_run': bool(payload.get('force_run')),
        'success': bool(payload.get('success')),
        'skipped': bool(payload.get('skipped')),
        'reason': str(payload.get('reason') or ''),
        'finance_sent_count': int(payload.get('finance_sent_count') or 0),
        'workforce_sent_count': int(payload.get('workforce_sent_count') or 0),
    }


def _executive_auto_period_info(dt: datetime, group_by: str) -> tuple[str, str, str]:
    value_date = timezone.localtime(dt).date()
    if group_by == 'week':
        weekday = value_date.weekday()
        start_date = value_date - timedelta(days=weekday)
        end_date = start_date + timedelta(days=6)
        iso_year, iso_week, _ = start_date.isocalendar()
        return (f'{iso_year}-W{str(iso_week).zfill(2)}', start_date.isoformat(), end_date.isoformat())
    text = value_date.isoformat()
    return (text, text, text)


def _cross_module_bootstrap_history_qs():
    return (
        AuditLog.objects
        .filter(
            entity_type=CROSS_MODULE_BOOTSTRAP_AUDIT_ENTITY_TYPE,
            entity_code=CROSS_MODULE_BOOTSTRAP_AUDIT_ENTITY_CODE,
        )
        .select_related('user')
    )


def _serialize_cross_module_bootstrap_history_row(row: AuditLog) -> dict:
    payload = row.new_values if isinstance(row.new_values, dict) else {}
    readiness_before = payload.get('readiness_before') if isinstance(payload.get('readiness_before'), dict) else {}
    readiness_after = payload.get('readiness_after') if isinstance(payload.get('readiness_after'), dict) else {}
    before_score = int(readiness_before.get('readiness_score') or 0)
    after_score = int(readiness_after.get('readiness_score') or 0)
    readiness_delta = int(after_score - before_score)
    level_before = str(readiness_before.get('readiness_level') or '')
    level_after = str(readiness_after.get('readiness_level') or '')
    return {
        'id': int(row.id),
        'created_at': row.created_at.isoformat() if row.created_at else '',
        'username': str(row.user.username) if row.user else '',
        'dry_run': bool(payload.get('dry_run')),
        'created_total': int(payload.get('created_total') or 0),
        'skipped_total': int(payload.get('skipped_total') or 0),
        'created': payload.get('created') if isinstance(payload.get('created'), dict) else {},
        'skipped': payload.get('skipped') if isinstance(payload.get('skipped'), dict) else {},
        'readiness_before': readiness_before,
        'readiness_after': readiness_after,
        'readiness_delta': readiness_delta,
        'improved': bool(readiness_delta > 0),
        'level_before': level_before,
        'level_after': level_after,
        'notes': payload.get('notes') if isinstance(payload.get('notes'), list) else [],
    }


def _build_cross_module_readiness_payload() -> dict:
    finance_advances_total = int(AdvanceTransaction.objects.count())
    finance_advances_active = int(AdvanceTransaction.objects.filter(is_active=True).count())
    workforce_salary_advances_total = int(SalaryAdvanceRecord.objects.count())
    payroll_records_total = int(PayrollRecord.objects.count())
    roles_total = int(Role.objects.filter(deleted_at__isnull=True).count())
    users_total = int(User.objects.count())
    users_staff_total = int(User.objects.filter(is_staff=True).count())
    role_permission_audits_total = int(AuditLog.objects.filter(entity_type='RoleModulePermission').count())
    operations_log_total = int(
        AuditLog.objects.filter(entity_type__in=['TaskBulk', 'WorkflowAnalytics', 'Task', 'WorkflowAutomation']).count()
    )
    pipeline_events_total = int(WorkflowPipelineEvent.objects.count())

    checks = [
        {
            'code': 'roles_configured',
            'label': 'Đã cấu hình role RBAC',
            'ok': roles_total > 0,
            'current': roles_total,
            'recommended_min': 1,
            'weight': 25,
            'hint': 'Tạo ít nhất 1 role để phân quyền module và tránh màn lịch sử phân quyền trống.',
        },
        {
            'code': 'module_permission_history',
            'label': 'Có lịch sử thay đổi phân quyền module',
            'ok': role_permission_audits_total > 0,
            'current': role_permission_audits_total,
            'recommended_min': 1,
            'weight': 20,
            'hint': 'Thực hiện ít nhất 1 lần cập nhật phân quyền module để sinh lịch sử audit.',
        },
        {
            'code': 'finance_advances_data',
            'label': 'Có dữ liệu tạm ứng Tài chính',
            'ok': finance_advances_total > 0,
            'current': finance_advances_total,
            'recommended_min': 1,
            'weight': 15,
            'hint': 'Tạo/nhập dữ liệu tạm ứng để dashboard Finance hiển thị đầy đủ.',
        },
        {
            'code': 'workforce_salary_advances_data',
            'label': 'Có dữ liệu ứng lương Nhân sự',
            'ok': workforce_salary_advances_total > 0,
            'current': workforce_salary_advances_total,
            'recommended_min': 1,
            'weight': 15,
            'hint': 'Tạo/nhập dữ liệu ứng lương để SLA Workforce và KPI điều hành có số liệu.',
        },
        {
            'code': 'payroll_data',
            'label': 'Có dữ liệu bảng lương',
            'ok': payroll_records_total > 0,
            'current': payroll_records_total,
            'recommended_min': 1,
            'weight': 10,
            'hint': 'Tạo kỳ lương đầu tiên để kích hoạt đối soát payroll-finance.',
        },
        {
            'code': 'operations_log_data',
            'label': 'Có dữ liệu nhật ký vận hành',
            'ok': (operations_log_total + pipeline_events_total) > 0,
            'current': int(operations_log_total + pipeline_events_total),
            'recommended_min': 1,
            'weight': 15,
            'hint': 'Thực hiện một số thao tác workflow/bulk để nhật ký vận hành có dữ liệu.',
        },
    ]

    readiness_score = 0
    warnings = []
    for check in checks:
        if bool(check.get('ok')):
            readiness_score += int(check.get('weight') or 0)
        else:
            warnings.append({
                'code': check['code'],
                'label': check['label'],
                'hint': check['hint'],
            })
    readiness_score = max(0, min(int(readiness_score), 100))
    readiness_level = 'READY' if readiness_score >= 85 else ('PARTIAL' if readiness_score >= 50 else 'BOOTSTRAP_NEEDED')
    return {
        'as_of': timezone.now().isoformat(),
        'readiness_score': readiness_score,
        'readiness_level': readiness_level,
        'checks': checks,
        'warnings': warnings,
        'summary': {
            'users_total': users_total,
            'users_staff_total': users_staff_total,
            'roles_total': roles_total,
            'finance_advances_total': finance_advances_total,
            'finance_advances_active': finance_advances_active,
            'workforce_salary_advances_total': workforce_salary_advances_total,
            'payroll_records_total': payroll_records_total,
            'role_permission_audits_total': role_permission_audits_total,
            'operations_log_total': operations_log_total,
            'pipeline_events_total': pipeline_events_total,
        },
    }


class SearchTextMixin:
    search_text_field = 'search_text'

    def check_module_read_permission(self):
        if not _can_manage_finance(self.request.user):
            raise PermissionDenied('Bạn không có quyền xem dữ liệu tài chính.')

    def apply_search(self, queryset):
        self.check_module_read_permission()
        search_raw = (self.request.query_params.get('q') or self.request.query_params.get('search') or '').strip()
        if not search_raw:
            return queryset

        normalized = unidecode(search_raw).lower().strip()
        exact_search = self.request.query_params.get('exact_search') in ('1', 'true', 'True')
        field_name = self.search_text_field

        if exact_search:
            query = Q(**{f'{field_name}__icontains': normalized})
            if search_raw != normalized:
                query |= Q(**{f'{field_name}__icontains': search_raw})
            return queryset.filter(query).distinct()

        tokens = [token for token in normalized.split() if token]
        if not tokens:
            return queryset

        query = Q(**{f'{field_name}__icontains': tokens[0]})
        for token in tokens[1:]:
            query &= Q(**{f'{field_name}__icontains': token})
        return queryset.filter(query).distinct()


class TransactionCategoryViewSet(SearchTextMixin, viewsets.ModelViewSet):
    queryset = TransactionCategory.objects.all()
    serializer_class = TransactionCategorySerializer
    permission_classes = [IsAuthenticated]
    filter_backends = [filters.OrderingFilter]
    ordering_fields = ['code', 'name', 'category_type', 'created_at']
    ordering = ['code']

    def get_queryset(self):
        queryset = super().get_queryset()
        category_type = (self.request.query_params.get('category_type') or '').strip()
        if category_type:
            queryset = queryset.filter(category_type=category_type)
        is_active = self.request.query_params.get('is_active')
        if is_active in ('true', 'false'):
            queryset = queryset.filter(is_active=(is_active == 'true'))
        return self.apply_search(queryset)

    def perform_create(self, serializer):
        if not _can_manage_finance(self.request.user):
            raise PermissionDenied('Bạn không có quyền tạo loại thu chi.')
        instance = serializer.save(created_by=self.request.user, updated_by=self.request.user)
        _log_finance_audit(
            self.request.user,
            action='CREATE',
            entity_type='FinanceTransactionCategory',
            entity_id=int(instance.id),
            entity_code=instance.code or str(instance.id),
            old_values={},
            new_values={'code': instance.code, 'name': instance.name, 'category_type': instance.category_type, 'is_system': instance.is_system},
            changed_fields=['code', 'name', 'category_type', 'is_system'],
        )

    def perform_update(self, serializer):
        if not _can_manage_finance(self.request.user):
            raise PermissionDenied('Bạn không có quyền cập nhật loại thu chi.')
        if serializer.instance.is_system and 'is_system' in serializer.validated_data:
            raise ValidationError({'is_system': 'Không được thay đổi cờ danh mục hệ thống.'})
        old = serializer.instance
        old_snap = {
            'code': old.code,
            'name': old.name,
            'category_type': old.category_type,
            'is_active': old.is_active,
            'is_system': old.is_system,
        }
        instance = serializer.save(updated_by=self.request.user)
        _log_finance_audit(
            self.request.user,
            action='UPDATE',
            entity_type='FinanceTransactionCategory',
            entity_id=int(instance.id),
            entity_code=instance.code or str(instance.id),
            old_values=old_snap,
            new_values={
                'code': instance.code,
                'name': instance.name,
                'category_type': instance.category_type,
                'is_active': instance.is_active,
                'is_system': instance.is_system,
            },
            changed_fields=list(serializer.validated_data.keys()),
        )

    def destroy(self, request, *args, **kwargs):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền xóa loại thu chi.'}, status=403)
        instance = self.get_object()
        if instance.is_system:
            return Response({'error': 'Danh mục hệ thống không được phép xóa.'}, status=400)
        if instance.transactions.exists():
            return Response({'error': 'Danh mục đã phát sinh giao dịch, chỉ nên ngừng sử dụng thay vì xóa.'}, status=400)
        snap = {'code': instance.code, 'name': instance.name, 'category_type': instance.category_type}
        response = super().destroy(request, *args, **kwargs)
        _log_finance_audit(
            request.user,
            action='DELETE',
            entity_type='FinanceTransactionCategory',
            entity_id=int(instance.id),
            entity_code=instance.code or str(instance.id),
            old_values=snap,
            new_values={},
            changed_fields=['deleted'],
        )
        return response


class BankAccountViewSet(SearchTextMixin, viewsets.ModelViewSet):
    queryset = BankAccount.objects.all()
    serializer_class = BankAccountSerializer
    permission_classes = [IsAuthenticated]
    filter_backends = [filters.OrderingFilter]
    ordering_fields = ['code', 'account_number', 'bank_name', 'created_at']
    ordering = ['code']

    def get_queryset(self):
        queryset = super().get_queryset()
        is_active = self.request.query_params.get('is_active')
        if is_active in ('true', 'false'):
            queryset = queryset.filter(is_active=(is_active == 'true'))
        return self.apply_search(queryset)

    def perform_create(self, serializer):
        if not _can_manage_finance(self.request.user):
            raise PermissionDenied('Bạn không có quyền tạo tài khoản ngân hàng.')
        instance = serializer.save(created_by=self.request.user, updated_by=self.request.user)
        _log_finance_audit(
            self.request.user, action='CREATE', entity_type='FinanceBankAccount',
            entity_id=int(instance.id), entity_code=instance.code or str(instance.id),
            old_values={}, new_values={'code': instance.code, 'bank_name': instance.bank_name, 'is_active': instance.is_active},
            changed_fields=['code', 'bank_name', 'is_active'],
        )

    def perform_update(self, serializer):
        if not _can_manage_finance(self.request.user):
            raise PermissionDenied('Bạn không có quyền cập nhật tài khoản ngân hàng.')
        old = serializer.instance
        old_snap = {'code': old.code, 'bank_name': old.bank_name, 'is_active': old.is_active}
        instance = serializer.save(updated_by=self.request.user)
        _log_finance_audit(
            self.request.user, action='UPDATE', entity_type='FinanceBankAccount',
            entity_id=int(instance.id), entity_code=instance.code or str(instance.id),
            old_values=old_snap,
            new_values={'code': instance.code, 'bank_name': instance.bank_name, 'is_active': instance.is_active},
            changed_fields=list(serializer.validated_data.keys()),
        )

    def destroy(self, request, *args, **kwargs):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền xóa tài khoản ngân hàng.'}, status=403)
        instance = self.get_object()
        if (
            CashTransaction.objects.filter(source_bank_account=instance).exists()
            or AdvanceTransaction.objects.filter(source_bank_account=instance).exists()
        ):
            return Response({'error': 'Tài khoản ngân hàng đã phát sinh chứng từ, không thể xóa.'}, status=400)
        snap = {'code': instance.code, 'bank_name': instance.bank_name}
        response = super().destroy(request, *args, **kwargs)
        _log_finance_audit(
            request.user, action='DELETE', entity_type='FinanceBankAccount',
            entity_id=int(instance.id), entity_code=instance.code or str(instance.id),
            old_values=snap, new_values={}, changed_fields=['deleted'],
        )
        return response


class BankReconciliationViewSet(SearchTextMixin, viewsets.ModelViewSet):
    queryset = BankReconciliation.objects.select_related('bank_account', 'approved_by', 'posted_by')
    serializer_class = BankReconciliationSerializer
    permission_classes = [IsAuthenticated]
    filter_backends = [filters.OrderingFilter]
    ordering_fields = ['statement_date', 'status', 'created_at']
    ordering = ['-statement_date', '-id']

    def get_queryset(self):
        queryset = super().get_queryset()
        status_value = (self.request.query_params.get('status') or '').strip()
        if status_value:
            queryset = queryset.filter(status=status_value)
        return self.apply_search(queryset)

    def perform_create(self, serializer):
        if not _can_manage_finance(self.request.user):
            raise PermissionDenied('Bạn không có quyền tạo phiếu đối soát ngân hàng.')
        statement_date_value = serializer.validated_data.get('statement_date')
        instance = serializer.save(
            code=_get_next_bank_reconciliation_code(statement_date_value),
            created_by=self.request.user,
            updated_by=self.request.user,
        )
        _log_finance_audit(
            self.request.user,
            action='CREATE',
            entity_type='FinanceBankReconciliation',
            entity_id=int(instance.id),
            entity_code=instance.code,
            old_values={},
            new_values={
                'statement_date': instance.statement_date.isoformat() if instance.statement_date else '',
                'status': instance.status,
                'statement_balance': str(instance.statement_balance),
                'book_balance': str(instance.book_balance),
                'delta': str(instance.delta),
            },
            changed_fields=['statement_date', 'status', 'statement_balance', 'book_balance', 'delta'],
        )

    def perform_update(self, serializer):
        if not _can_manage_finance(self.request.user):
            raise PermissionDenied('Bạn không có quyền cập nhật phiếu đối soát ngân hàng.')
        if serializer.instance.status != BankReconciliationStatus.DRAFT:
            raise ValidationError({'status': 'Chỉ được sửa phiếu đối soát ở trạng thái nháp.'})
        old = serializer.instance
        old_snap = {
            'statement_date': old.statement_date.isoformat() if old.statement_date else '',
            'status': old.status,
            'statement_balance': str(old.statement_balance),
            'book_balance': str(old.book_balance),
            'delta': str(old.delta),
        }
        instance = serializer.save(updated_by=self.request.user)
        _log_finance_audit(
            self.request.user,
            action='UPDATE',
            entity_type='FinanceBankReconciliation',
            entity_id=int(instance.id),
            entity_code=instance.code,
            old_values=old_snap,
            new_values={
                'statement_date': instance.statement_date.isoformat() if instance.statement_date else '',
                'status': instance.status,
                'statement_balance': str(instance.statement_balance),
                'book_balance': str(instance.book_balance),
                'delta': str(instance.delta),
            },
            changed_fields=list(serializer.validated_data.keys()),
        )

    def destroy(self, request, *args, **kwargs):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền xóa phiếu đối soát ngân hàng.'}, status=403)
        instance = self.get_object()
        if instance.status != BankReconciliationStatus.DRAFT:
            return Response({'error': 'Chỉ được xóa phiếu đối soát ở trạng thái nháp.'}, status=400)
        snap = {'code': instance.code, 'status': instance.status}
        response = super().destroy(request, *args, **kwargs)
        _log_finance_audit(
            request.user,
            action='DELETE',
            entity_type='FinanceBankReconciliation',
            entity_id=int(instance.id),
            entity_code=instance.code,
            old_values=snap,
            new_values={},
            changed_fields=['deleted'],
        )
        return response

    @action(detail=True, methods=['post'])
    def approve(self, request, pk=None):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền duyệt phiếu đối soát.'}, status=403)
        instance = self.get_object()
        if instance.status != BankReconciliationStatus.DRAFT:
            return Response({'error': 'Chỉ được duyệt phiếu ở trạng thái nháp.'}, status=400)
        instance.status = BankReconciliationStatus.APPROVED
        instance.approved_by = request.user
        instance.approved_at = timezone.now()
        instance.save(update_fields=['status', 'approved_by', 'approved_at', 'updated_at'])
        _log_finance_audit(
            request.user,
            action='APPROVE',
            entity_type='FinanceBankReconciliation',
            entity_id=int(instance.id),
            entity_code=instance.code,
            old_values={'status': BankReconciliationStatus.DRAFT},
            new_values={'status': instance.status},
            changed_fields=['status'],
        )
        return Response({'status': instance.status})

    @action(detail=True, methods=['post'])
    def post(self, request, pk=None):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền post phiếu đối soát.'}, status=403)
        instance = self.get_object()
        if instance.status != BankReconciliationStatus.APPROVED:
            return Response({'error': 'Chỉ được post phiếu đã duyệt.'}, status=400)
        instance.status = BankReconciliationStatus.POSTED
        instance.posted_by = request.user
        instance.posted_at = timezone.now()
        instance.save(update_fields=['status', 'posted_by', 'posted_at', 'updated_at'])
        _log_finance_audit(
            request.user,
            action='POST',
            entity_type='FinanceBankReconciliation',
            entity_id=int(instance.id),
            entity_code=instance.code,
            old_values={'status': BankReconciliationStatus.APPROVED},
            new_values={'status': instance.status},
            changed_fields=['status'],
        )
        return Response({'status': instance.status})


class CashAccountViewSet(SearchTextMixin, viewsets.ModelViewSet):
    queryset = CashAccount.objects.all()
    serializer_class = CashAccountSerializer
    permission_classes = [IsAuthenticated]
    filter_backends = [filters.OrderingFilter]
    ordering_fields = ['name', 'account_type', 'balance', 'created_at']
    ordering = ['name']

    def get_queryset(self):
        queryset = super().get_queryset()
        account_type = (self.request.query_params.get('account_type') or '').strip()
        if account_type:
            queryset = queryset.filter(account_type=account_type)
        is_active = self.request.query_params.get('is_active')
        if is_active in ('true', 'false'):
            queryset = queryset.filter(is_active=(is_active == 'true'))
        return self.apply_search(queryset)

    def perform_create(self, serializer):
        if not _can_manage_finance(self.request.user):
            raise PermissionDenied('Bạn không có quyền tạo tài khoản quỹ.')
        instance = serializer.save(created_by=self.request.user, updated_by=self.request.user)
        _log_finance_audit(
            self.request.user, action='CREATE', entity_type='FinanceCashAccount',
            entity_id=int(instance.id), entity_code=instance.name or str(instance.id),
            old_values={}, new_values={'name': instance.name, 'account_type': instance.account_type, 'is_active': instance.is_active},
            changed_fields=['name', 'account_type', 'is_active'],
        )

    def perform_update(self, serializer):
        if not _can_manage_finance(self.request.user):
            raise PermissionDenied('Bạn không có quyền cập nhật tài khoản quỹ.')
        if 'balance' in serializer.validated_data and serializer.validated_data.get('balance') != serializer.instance.balance:
            raise ValidationError({'balance': 'Không cho phép sửa trực tiếp số dư mở sổ sau khi đã tạo tài khoản.'})
        old = serializer.instance
        old_snap = {'name': old.name, 'account_type': old.account_type, 'is_active': old.is_active}
        instance = serializer.save(updated_by=self.request.user)
        _log_finance_audit(
            self.request.user, action='UPDATE', entity_type='FinanceCashAccount',
            entity_id=int(instance.id), entity_code=instance.name or str(instance.id),
            old_values=old_snap,
            new_values={'name': instance.name, 'account_type': instance.account_type, 'is_active': instance.is_active},
            changed_fields=list(serializer.validated_data.keys()),
        )

    def destroy(self, request, *args, **kwargs):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền xóa tài khoản quỹ.'}, status=403)
        instance = self.get_object()
        if (
            CashTransaction.objects.filter(Q(source_cash_account=instance) | Q(target_cash_account=instance)).exists()
            or AdvanceTransaction.objects.filter(source_cash_account=instance).exists()
        ):
            return Response({'error': 'Tài khoản quỹ đã phát sinh chứng từ, không thể xóa.'}, status=400)
        snap = {'name': instance.name, 'account_type': instance.account_type}
        response = super().destroy(request, *args, **kwargs)
        _log_finance_audit(
            request.user, action='DELETE', entity_type='FinanceCashAccount',
            entity_id=int(instance.id), entity_code=instance.name or str(instance.id),
            old_values=snap, new_values={}, changed_fields=['deleted'],
        )
        return response


class CashTransactionViewSet(SearchTextMixin, viewsets.ModelViewSet):
    queryset = CashTransaction.objects.select_related(
        'category',
        'source_cash_account',
        'source_bank_account',
        'target_cash_account',
    )
    serializer_class = CashTransactionSerializer
    permission_classes = [IsAuthenticated]
    filter_backends = [filters.OrderingFilter]
    ordering_fields = ['transaction_date', 'amount', 'transaction_type', 'created_at']
    ordering = ['-transaction_date', '-id']

    def get_queryset(self):
        queryset = super().get_queryset()
        params = self.request.query_params

        transaction_type = (params.get('transaction_type') or '').strip()
        if transaction_type:
            queryset = queryset.filter(transaction_type=transaction_type)

        source_type = (params.get('source_type') or '').strip()
        if source_type:
            queryset = queryset.filter(source_type=source_type)

        source_cash_account = (params.get('source_cash_account') or '').strip()
        if source_cash_account.isdigit():
            queryset = queryset.filter(source_cash_account_id=int(source_cash_account))

        source_bank_account = (params.get('source_bank_account') or '').strip()
        if source_bank_account.isdigit():
            queryset = queryset.filter(source_bank_account_id=int(source_bank_account))

        date_from_s = (params.get('transaction_date__gte') or params.get('date_from') or '').strip()
        date_to_s = (params.get('transaction_date__lte') or params.get('date_to') or '').strip()
        if date_from_s:
            try:
                queryset = queryset.filter(transaction_date__gte=date.fromisoformat(date_from_s))
            except ValueError:
                pass
        if date_to_s:
            try:
                queryset = queryset.filter(transaction_date__lte=date.fromisoformat(date_to_s))
            except ValueError:
                pass

        month = (params.get('month') or '').strip()
        if len(month) == 7 and '-' in month:
            year_str, month_str = month.split('-', 1)
            if year_str.isdigit() and month_str.isdigit():
                queryset = queryset.filter(
                    transaction_date__year=int(year_str),
                    transaction_date__month=int(month_str),
                )

        return self.apply_search(queryset)

    def perform_create(self, serializer):
        if not _can_manage_finance(self.request.user):
            raise PermissionDenied('Bạn không có quyền tạo giao dịch tài chính.')
        tx_date = serializer.validated_data.get('transaction_date')
        _ensure_finance_month_unlocked(
            _month_from_date_obj(tx_date),
            'Tháng tài chính đã khóa, không thể thêm giao dịch.',
        )
        instance = serializer.save(created_by=self.request.user)
        _log_finance_audit(
            self.request.user, action='CREATE', entity_type='FinanceCashTransaction',
            entity_id=int(instance.id), entity_code=str(instance.id),
            old_values={},
            new_values={
                'transaction_type': instance.transaction_type,
                'amount': str(instance.amount),
                'transaction_date': instance.transaction_date.isoformat() if instance.transaction_date else '',
                'source_type': instance.source_type,
            },
            changed_fields=['transaction_type', 'amount', 'transaction_date', 'source_type'],
        )

    def perform_update(self, serializer):
        if not _can_manage_finance(self.request.user):
            raise PermissionDenied('Bạn không có quyền cập nhật giao dịch tài chính.')
        old_month = _month_from_date_obj(serializer.instance.transaction_date)
        tx_date = serializer.validated_data.get('transaction_date', serializer.instance.transaction_date)
        _ensure_finance_months_unlocked(
            [old_month, _month_from_date_obj(tx_date)],
            'Tháng tài chính đã khóa, không thể cập nhật giao dịch.',
        )
        old = serializer.instance
        old_snap = {
            'transaction_type': old.transaction_type, 'amount': str(old.amount),
            'transaction_date': old.transaction_date.isoformat() if old.transaction_date else '',
        }
        instance = serializer.save()
        _log_finance_audit(
            self.request.user, action='UPDATE', entity_type='FinanceCashTransaction',
            entity_id=int(instance.id), entity_code=str(instance.id),
            old_values=old_snap,
            new_values={
                'transaction_type': instance.transaction_type,
                'amount': str(instance.amount),
                'transaction_date': instance.transaction_date.isoformat() if instance.transaction_date else '',
            },
            changed_fields=list(serializer.validated_data.keys()),
        )

    @action(detail=False, methods=['post'])
    def bulk_delete(self, request):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền xóa giao dịch tài chính.'}, status=403)
        ids = request.data.get('ids', [])
        if not isinstance(ids, list) or not ids:
            return Response({'error': 'No IDs provided'}, status=400)
        locked_months = _get_locked_finance_months()
        if locked_months:
            txs = CashTransaction.objects.filter(id__in=ids)
            for tx in txs:
                month = _month_from_date_obj(tx.transaction_date)
                if month in locked_months:
                    return Response({'error': f'Tháng {month} đã khóa, không thể xóa giao dịch.'}, status=400)
        deleted_ids = list(CashTransaction.objects.filter(id__in=ids).values_list('id', flat=True))
        deleted_count, _ = CashTransaction.objects.filter(id__in=ids).delete()
        if deleted_ids:
            _log_finance_audit(
                request.user,
                action='DELETE',
                entity_type='FinanceCashTransactionBulkDelete',
                entity_id=0,
                entity_code='bulk_delete',
                old_values={'ids': deleted_ids},
                new_values={},
                changed_fields=['deleted_ids'],
            )
        return Response({'success': True, 'count': deleted_count})

    def destroy(self, request, *args, **kwargs):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền xóa giao dịch tài chính.'}, status=403)
        instance = self.get_object()
        _ensure_finance_month_unlocked(
            _month_from_date_obj(instance.transaction_date),
            'Tháng tài chính đã khóa, không thể xóa giao dịch.',
        )
        snap = {
            'transaction_type': instance.transaction_type, 'amount': str(instance.amount),
            'transaction_date': instance.transaction_date.isoformat() if instance.transaction_date else '',
        }
        response = super().destroy(request, *args, **kwargs)
        _log_finance_audit(
            request.user, action='DELETE', entity_type='FinanceCashTransaction',
            entity_id=int(instance.id), entity_code=str(instance.id),
            old_values=snap, new_values={}, changed_fields=['deleted'],
        )
        return response

    @action(detail=False, methods=['get'])
    def locked_months(self, request):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền xem danh sách tháng đã khóa.'}, status=403)
        months = sorted(_get_locked_finance_months())
        return Response({'months': months})

    @action(detail=False, methods=['get'])
    def preclose_check(self, request):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền kiểm tra đóng sổ tài chính.'}, status=403)
        month = _normalize_month(str(request.query_params.get('month') or ''))
        if not month:
            return Response({'error': 'month phải có dạng YYYY-MM'}, status=400)
        return Response(_build_finance_month_close_check(month))

    @action(detail=False, methods=['post'])
    def lock_month(self, request):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền khóa sổ tài chính.'}, status=403)
        month = _normalize_month(str(request.data.get('month') or ''))
        if not month:
            return Response({'error': 'month phải có dạng YYYY-MM'}, status=400)
        checklist = _build_finance_month_close_check(month)
        if checklist.get('blockers'):
            return Response(
                {
                    'error': 'Chưa thể khóa sổ tài chính vì còn vướng kiểm tra đóng kỳ.',
                    'checklist': checklist,
                },
                status=400,
            )
        months = _get_locked_finance_months()
        old_months = sorted(months)
        months.add(month)
        _save_locked_finance_months(months)
        _log_finance_audit(
            request.user,
            action='UPDATE',
            entity_type='FinanceMonthLock',
            entity_id=0,
            entity_code=FINANCE_LOCKED_MONTHS_KEY,
            old_values={'months': old_months},
            new_values={'months': sorted(months), 'action': 'lock', 'month': month},
            changed_fields=['months'],
        )
        return Response({'success': True, 'months': sorted(months)})

    @action(detail=False, methods=['post'])
    def unlock_month(self, request):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền mở khóa sổ tài chính.'}, status=403)
        month = _normalize_month(str(request.data.get('month') or ''))
        if not month:
            return Response({'error': 'month phải có dạng YYYY-MM'}, status=400)
        force_unlock = str(request.data.get('force') or '').strip().lower() in {'1', 'true', 'yes'}
        y, m = month.split('-', 1)
        if not force_unlock:
            has_data = (
                CashTransaction.objects.filter(transaction_date__year=int(y), transaction_date__month=int(m)).exists()
                or AdvanceTransaction.objects.filter(advance_date__year=int(y), advance_date__month=int(m)).exists()
                or AdvanceSettlement.objects.filter(settlement_date__year=int(y), settlement_date__month=int(m)).exists()
            )
            if has_data:
                return Response(
                    {'error': 'Tháng đã có dữ liệu tài chính. Chỉ superuser mới được force mở khóa với force=true.'},
                    status=400,
                )
        elif not getattr(request.user, 'is_superuser', False):
            return Response({'error': 'Chỉ superuser mới được force mở khóa sổ.'}, status=403)
        months = _get_locked_finance_months()
        old_months = sorted(months)
        months.discard(month)
        _save_locked_finance_months(months)
        _log_finance_audit(
            request.user,
            action='UPDATE',
            entity_type='FinanceMonthLock',
            entity_id=0,
            entity_code=FINANCE_LOCKED_MONTHS_KEY,
            old_values={'months': old_months},
            new_values={'months': sorted(months), 'action': 'unlock', 'month': month, 'force': force_unlock},
            changed_fields=['months'],
        )
        return Response({'success': True, 'months': sorted(months)})

    @action(detail=False, methods=['get'])
    def payroll_reconciliation(self, request):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền xem đối soát bảng lương.'}, status=403)
        month = _normalize_month(str(request.query_params.get('month') or ''))
        if not month:
            return Response({'error': 'month phải có dạng YYYY-MM'}, status=400)

        payroll_qs = PayrollRecord.objects.filter(month=month, status=PayrollRecord.STATUS_LOCKED)
        payroll_total = Decimal(str(payroll_qs.aggregate(total=Sum('net_pay')).get('total') or 0))
        payroll_count = payroll_qs.count()

        y, m = month.split('-', 1)
        tx_qs = CashTransaction.objects.filter(
            transaction_type=CashTransaction.TYPE_EXPENSE,
            transaction_date__year=int(y),
            transaction_date__month=int(m),
            reason__icontains='[PAYROLL:',
        )
        posted_total = Decimal(str(tx_qs.aggregate(total=Sum('amount')).get('total') or 0))
        posted_count = tx_qs.count()
        delta = payroll_total - posted_total

        return Response({
            'month': month,
            'payroll_total': str(payroll_total),
            'payroll_count': payroll_count,
            'posted_total': str(posted_total),
            'posted_count': posted_count,
            'delta': str(delta),
            'is_balanced': delta == 0,
        })

    @action(detail=False, methods=['get'])
    def monthly_summary(self, request):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền xem tổng hợp tài chính tháng.'}, status=403)
        month = (request.query_params.get('month') or '').strip()
        if len(month) != 7 or '-' not in month:
            return Response({'error': 'month phải có dạng YYYY-MM'}, status=400)
        year_str, month_str = month.split('-', 1)
        if not (year_str.isdigit() and month_str.isdigit()):
            return Response({'error': 'month phải có dạng YYYY-MM'}, status=400)
        year = int(year_str)
        month_num = int(month_str)
        if month_num < 1 or month_num > 12:
            return Response({'error': 'month không hợp lệ'}, status=400)

        tx_qs = CashTransaction.objects.filter(
            transaction_date__year=year,
            transaction_date__month=month_num,
        )
        total_income = Decimal(str(
            tx_qs.filter(transaction_type=CashTransaction.TYPE_INCOME).aggregate(total=Sum('amount')).get('total') or 0
        ))
        total_expense = Decimal(str(
            tx_qs.filter(transaction_type=CashTransaction.TYPE_EXPENSE).aggregate(total=Sum('amount')).get('total') or 0
        ))
        total_transfer = Decimal(str(
            tx_qs.filter(transaction_type=CashTransaction.TYPE_TRANSFER).aggregate(total=Sum('amount')).get('total') or 0
        ))

        advance_qs = _approved_advance_queryset().filter(
            advance_date__year=year,
            advance_date__month=month_num,
        )
        total_advance = Decimal(str(advance_qs.aggregate(total=Sum('amount')).get('total') or 0))

        settlement_qs = AdvanceSettlement.objects.filter(
            settlement_date__year=year,
            settlement_date__month=month_num,
        )
        total_settlement_spent = Decimal(str(settlement_qs.aggregate(total=Sum('spent_amount')).get('total') or 0))
        total_settlement_refund = Decimal(str(settlement_qs.aggregate(total=Sum('refund_amount')).get('total') or 0))

        data = {
            'month': month,
            'total_income': str(total_income),
            'total_expense': str(total_expense),
            'total_transfer': str(total_transfer),
            'cash_delta': str(total_income - total_expense),
            'total_advance': str(total_advance),
            'total_settlement_spent': str(total_settlement_spent),
            'total_settlement_refund': str(total_settlement_refund),
            'advance_net_delta': str(total_advance - total_settlement_spent - total_settlement_refund),
            'transactions_count': tx_qs.count(),
            'advances_count': advance_qs.count(),
            'settlements_count': settlement_qs.count(),
        }

        if request.query_params.get('export') == 'excel':
            wb = Workbook()
            ws = wb.active
            ws.title = f'BaoCao-{month}'
            ws.append(['Chỉ tiêu', 'Giá trị'])
            ws.append(['Tháng', month])
            ws.append(['Tổng thu', float(total_income)])
            ws.append(['Tổng chi', float(total_expense)])
            ws.append(['Tổng chuyển', float(total_transfer)])
            ws.append(['Chênh lệch thu-chi', float(total_income - total_expense)])
            ws.append(['Tổng tạm ứng', float(total_advance)])
            ws.append(['Tổng chi quyết toán', float(total_settlement_spent)])
            ws.append(['Tổng hoàn ứng', float(total_settlement_refund)])
            ws.append(['Chênh lệch tạm ứng', float(total_advance - total_settlement_spent - total_settlement_refund)])
            ws.append(['Số giao dịch', tx_qs.count()])
            ws.append(['Số phiếu tạm ứng', advance_qs.count()])
            ws.append(['Số phiếu quyết toán', settlement_qs.count()])

            response = HttpResponse(
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )
            response['Content-Disposition'] = f'attachment; filename="bao_cao_tai_chinh_{month}.xlsx"'
            wb.save(response)
            return response

        return Response(data)

    @action(detail=False, methods=['get'])
    def cash_flow_summary(self, request):
        """Báo cáo thu chi theo kỳ (date_from, date_to). GET ?date_from=YYYY-MM-DD&date_to=YYYY-MM-DD"""
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền xem báo cáo thu chi.'}, status=403)
        date_from_s = (request.query_params.get('date_from') or '').strip()
        date_to_s = (request.query_params.get('date_to') or '').strip()
        if not date_from_s or not date_to_s:
            return Response({'error': 'date_from và date_to bắt buộc (YYYY-MM-DD).'}, status=400)
        try:
            date_from = date.fromisoformat(date_from_s)
            date_to = date.fromisoformat(date_to_s)
        except ValueError:
            return Response({'error': 'date_from, date_to phải đúng định dạng YYYY-MM-DD.'}, status=400)
        if date_from > date_to:
            return Response({'error': 'date_from không được lớn hơn date_to.'}, status=400)

        tx_qs = CashTransaction.objects.filter(
            transaction_date__gte=date_from,
            transaction_date__lte=date_to,
        )
        total_income = Decimal(str(
            tx_qs.filter(transaction_type=CashTransaction.TYPE_INCOME).aggregate(total=Sum('amount')).get('total') or 0
        ))
        total_expense = Decimal(str(
            tx_qs.filter(transaction_type=CashTransaction.TYPE_EXPENSE).aggregate(total=Sum('amount')).get('total') or 0
        ))
        return Response({
            'date_from': date_from_s,
            'date_to': date_to_s,
            'total_income': str(total_income),
            'total_expense': str(total_expense),
            'cash_delta': str(total_income - total_expense),
            'transactions_count': tx_qs.count(),
        })

    @action(detail=False, methods=['get'])
    def general_ledger(self, request):
        """Sổ cái: danh sách phát sinh từ giao dịch quỹ (date_from, date_to, category tùy chọn)."""
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền xem sổ cái.'}, status=403)
        date_from_s = (request.query_params.get('date_from') or '').strip()
        date_to_s = (request.query_params.get('date_to') or '').strip()
        if not date_from_s or not date_to_s:
            return Response({'error': 'date_from và date_to bắt buộc (YYYY-MM-DD).'}, status=400)
        try:
            date_from = date.fromisoformat(date_from_s)
            date_to = date.fromisoformat(date_to_s)
        except ValueError:
            return Response({'error': 'date_from, date_to phải đúng định dạng YYYY-MM-DD.'}, status=400)
        if date_from > date_to:
            return Response({'error': 'date_from không được lớn hơn date_to.'}, status=400)
        category_id = request.query_params.get('category')
        tx_qs = CashTransaction.objects.filter(
            transaction_date__gte=date_from,
            transaction_date__lte=date_to,
        ).select_related('category', 'source_cash_account', 'target_cash_account')
        if category_id:
            tx_qs = tx_qs.filter(category_id=category_id)
        tx_qs = tx_qs.order_by('transaction_date', 'id')
        results = []
        for tx in tx_qs:
            amount = tx.amount or Decimal('0')
            debit = '0'
            credit = '0'
            if tx.transaction_type == CashTransaction.TYPE_INCOME:
                credit = str(amount)
            elif tx.transaction_type == CashTransaction.TYPE_EXPENSE:
                debit = str(amount)
            else:
                debit = str(amount)
                credit = str(amount)
            results.append({
                'id': tx.id,
                'transaction_date': tx.transaction_date.isoformat(),
                'reference': tx.reference or '',
                'category_id': tx.category_id,
                'category_code': getattr(tx.category, 'code', '') if tx.category_id else '',
                'category_name': getattr(tx.category, 'name', '') if tx.category_id else '',
                'description': tx.reason or '',
                'transaction_type': tx.transaction_type,
                'debit': debit,
                'credit': credit,
            })
        return Response({
            'date_from': date_from_s,
            'date_to': date_to_s,
            'results': results,
        })

    @action(detail=False, methods=['get'])
    def trend_12m(self, request):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền xem xu hướng tài chính 12 tháng.'}, status=403)
        end_month = (request.query_params.get('end_month') or '').strip()
        if not end_month:
            today = date.today()
            end_year = today.year
            end_month_num = today.month
            end_month = f'{end_year}-{str(end_month_num).zfill(2)}'
        else:
            if len(end_month) != 7 or '-' not in end_month:
                return Response({'error': 'end_month phải có dạng YYYY-MM'}, status=400)
            y, m = end_month.split('-', 1)
            if not (y.isdigit() and m.isdigit()):
                return Response({'error': 'end_month phải có dạng YYYY-MM'}, status=400)
            end_year = int(y)
            end_month_num = int(m)
            if end_month_num < 1 or end_month_num > 12:
                return Response({'error': 'end_month không hợp lệ'}, status=400)

        months: list[tuple[int, int, str]] = []
        y = end_year
        m = end_month_num
        for _ in range(12):
            months.append((y, m, f'{y}-{str(m).zfill(2)}'))
            m -= 1
            if m == 0:
                y -= 1
                m = 12
        months.reverse()

        items = []
        for y_item, m_item, month_key in months:
            tx_qs = CashTransaction.objects.filter(
                transaction_date__year=y_item,
                transaction_date__month=m_item,
            )
            total_income = Decimal(str(
                tx_qs.filter(transaction_type=CashTransaction.TYPE_INCOME).aggregate(total=Sum('amount')).get('total') or 0
            ))
            total_expense = Decimal(str(
                tx_qs.filter(transaction_type=CashTransaction.TYPE_EXPENSE).aggregate(total=Sum('amount')).get('total') or 0
            ))

            advance_qs = _approved_advance_queryset().filter(
                advance_date__year=y_item,
                advance_date__month=m_item,
            )
            total_advance = Decimal(str(advance_qs.aggregate(total=Sum('amount')).get('total') or 0))

            settlement_qs = AdvanceSettlement.objects.filter(
                settlement_date__year=y_item,
                settlement_date__month=m_item,
            )
            total_settlement_spent = Decimal(str(settlement_qs.aggregate(total=Sum('spent_amount')).get('total') or 0))
            total_settlement_refund = Decimal(str(settlement_qs.aggregate(total=Sum('refund_amount')).get('total') or 0))

            items.append({
                'month': month_key,
                'total_income': str(total_income),
                'total_expense': str(total_expense),
                'cash_delta': str(total_income - total_expense),
                'total_advance': str(total_advance),
                'total_settlement_spent': str(total_settlement_spent),
                'total_settlement_refund': str(total_settlement_refund),
                'advance_net_delta': str(total_advance - total_settlement_spent - total_settlement_refund),
            })

        return Response({'end_month': end_month, 'items': items})


def refresh_advance_status(advance: AdvanceTransaction, user):
    totals = advance.settlements.aggregate(
        spent_total=Sum('spent_amount'),
        refund_total=Sum('refund_amount'),
    )
    spent_total = Decimal(str(totals.get('spent_total') or 0))
    refund_total = Decimal(str(totals.get('refund_total') or 0))
    handled_total = spent_total + refund_total
    amount = Decimal(str(advance.amount or 0))

    if advance.status == AdvanceTransaction.STATUS_CANCELLED:
        return
    if handled_total <= 0:
        new_status = AdvanceTransaction.STATUS_OPEN
    elif handled_total < amount:
        new_status = AdvanceTransaction.STATUS_PARTIAL
    else:
        new_status = AdvanceTransaction.STATUS_SETTLED
    if advance.status != new_status:
        advance.status = new_status
        advance.updated_by = user
        advance.save(update_fields=['status', 'updated_by', 'updated_at', 'search_text'])


def _advance_disbursement_marker(advance: AdvanceTransaction) -> str:
    return f'[ADVANCE:{advance.id}]'


def _get_advance_disbursement_tx(advance: AdvanceTransaction, *, lock_for_update: bool = False):
    qs = CashTransaction.objects.filter(reason__icontains=_advance_disbursement_marker(advance)).order_by('-id')
    if lock_for_update:
        qs = qs.select_for_update()
    return qs.first()


def _ensure_advance_disbursement_category(actor):
    category, _ = TransactionCategory.objects.get_or_create(
        code='ADVANCE_DISB',
        defaults={
            'name': 'Chi tạm ứng',
            'category_type': TransactionCategory.TYPE_EXPENSE,
            'color': '#d46b08',
            'note': 'Danh mục hệ thống cho chi tiền tạm ứng',
            'is_system': True,
            'is_active': True,
            'created_by': actor,
            'updated_by': actor,
        },
    )
    if not category.is_active:
        category.is_active = True
        category.updated_by = actor
        category.save(update_fields=['is_active', 'updated_by', 'updated_at', 'search_text'])
    return category


def _recommend_reminder_policy_preset() -> dict:
    presets = get_reminder_policy_presets()
    cutoff = date.today().fromordinal(date.today().toordinal() - 30)
    qs = Notification.objects.filter(
        notification_type='due_date',
        entity_type='FinanceAdvanceOverdueDaily',
        created_at__date__gte=cutoff,
    )
    total_sent = int(qs.count() or 0)
    total_unread = int(qs.filter(is_read=False).count() or 0)
    unread_rate = round((total_unread / total_sent) * 100, 2) if total_sent > 0 else 0

    recommended_key = 'balanced'
    reason = 'Dữ liệu ổn định, ưu tiên cân bằng tần suất và độ phủ nhắc.'
    if total_sent < 10:
        recommended_key = 'balanced'
        reason = 'Dữ liệu lịch sử còn ít, ưu tiên preset cân bằng.'
    elif unread_rate >= 55:
        recommended_key = 'conservative'
        reason = 'Tỷ lệ chưa đọc cao, nên giảm tần suất để tránh quá tải thông báo.'
    elif unread_rate <= 20 and total_sent >= 20:
        recommended_key = 'aggressive'
        reason = 'Tỷ lệ đọc cao, có thể tăng tần suất để bám sát xử lý quá hạn.'

    return {
        'recommended_preset_key': recommended_key,
        'reason': reason,
        'metrics': {
            'total_sent_30d': total_sent,
            'total_unread_30d': total_unread,
            'unread_rate_30d': unread_rate,
        },
        'preset': dict(presets.get(recommended_key) or {}),
    }


class AdvanceTransactionViewSet(SearchTextMixin, viewsets.ModelViewSet):
    queryset = AdvanceTransaction.objects.select_related('source_cash_account', 'source_bank_account')
    serializer_class = AdvanceTransactionSerializer
    permission_classes = [IsAuthenticated]
    filter_backends = [filters.OrderingFilter]
    ordering_fields = ['advance_date', 'code', 'amount', 'status', 'created_at']
    ordering = ['-advance_date', '-id']

    def get_queryset(self):
        queryset = super().get_queryset()
        params = self.request.query_params

        advance_type = (params.get('advance_type') or '').strip()
        if advance_type:
            queryset = queryset.filter(advance_type=advance_type)

        status_param = (params.get('status') or '').strip()
        if status_param:
            queryset = queryset.filter(status=status_param)
        approval_status = (params.get('approval_status') or '').strip()
        if approval_status:
            queryset = queryset.filter(approval_status=approval_status)

        month = (params.get('month') or '').strip()
        if len(month) == 7 and '-' in month:
            year_str, month_str = month.split('-', 1)
            if year_str.isdigit() and month_str.isdigit():
                queryset = queryset.filter(
                    advance_date__year=int(year_str),
                    advance_date__month=int(month_str),
                )

        is_active = params.get('is_active')
        if is_active in ('true', 'false'):
            queryset = queryset.filter(is_active=(is_active == 'true'))

        return self.apply_search(queryset)

    def perform_create(self, serializer):
        if not _can_manage_finance(self.request.user):
            raise PermissionDenied('Bạn không có quyền tạo phiếu tạm ứng.')
        adv_date = serializer.validated_data.get('advance_date')
        _ensure_finance_month_unlocked(
            _month_from_date_obj(adv_date),
            'Tháng tài chính đã khóa, không thể thêm phiếu tạm ứng.',
        )
        amount = Decimal(str(serializer.validated_data.get('amount') or 0))
        required_level = _required_approval_level_for_amount(amount)
        advance = serializer.save(
            created_by=self.request.user,
            updated_by=self.request.user,
            required_approval_level=required_level,
            approval_status=AdvanceTransaction.APPROVAL_DRAFT,
        )
        _log_finance_audit(
            self.request.user,
            action='CREATE',
            entity_type='FinanceAdvanceApproval',
            entity_id=int(advance.id),
            entity_code=advance.code,
            old_values={},
            new_values={
                'approval_status': advance.approval_status,
                'required_approval_level': advance.required_approval_level,
            },
            changed_fields=['approval_status', 'required_approval_level'],
        )

    def perform_update(self, serializer):
        if not _can_manage_finance(self.request.user):
            raise PermissionDenied('Bạn không có quyền cập nhật phiếu tạm ứng.')
        old_month = _month_from_date_obj(serializer.instance.advance_date)
        adv_date = serializer.validated_data.get('advance_date', serializer.instance.advance_date)
        _ensure_finance_months_unlocked(
            [old_month, _month_from_date_obj(adv_date)],
            'Tháng tài chính đã khóa, không thể cập nhật phiếu tạm ứng.',
        )
        current_status = serializer.instance.approval_status
        if current_status in {
            AdvanceTransaction.APPROVAL_PENDING_L1,
            AdvanceTransaction.APPROVAL_PENDING_L2,
            AdvanceTransaction.APPROVAL_APPROVED,
        }:
            raise PermissionDenied('Phiếu đã gửi duyệt/đã duyệt, không thể chỉnh sửa trực tiếp.')
        amount = Decimal(str(serializer.validated_data.get('amount', serializer.instance.amount) or 0))
        required_level = _required_approval_level_for_amount(amount)
        old = serializer.instance
        old_snap = {
            'code': old.code,
            'advance_date': old.advance_date.isoformat() if old.advance_date else '',
            'recipient_name': old.recipient_name,
            'amount': str(old.amount),
            'status': old.status,
            'approval_status': old.approval_status,
        }
        instance = serializer.save(
            updated_by=self.request.user,
            required_approval_level=required_level,
        )
        _log_finance_audit(
            self.request.user,
            action='UPDATE',
            entity_type='FinanceAdvanceApproval',
            entity_id=int(instance.id),
            entity_code=instance.code,
            old_values=old_snap,
            new_values={
                'code': instance.code,
                'advance_date': instance.advance_date.isoformat() if instance.advance_date else '',
                'recipient_name': instance.recipient_name,
                'amount': str(instance.amount),
                'status': instance.status,
                'approval_status': instance.approval_status,
                'required_approval_level': instance.required_approval_level,
            },
            changed_fields=list(serializer.validated_data.keys()),
        )

    @action(detail=False, methods=['post'])
    def bulk_delete(self, request):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền xóa phiếu tạm ứng.'}, status=403)
        ids = request.data.get('ids', [])
        if not isinstance(ids, list) or not ids:
            return Response({'error': 'No IDs provided'}, status=400)
        locked_months = _get_locked_finance_months()
        if locked_months:
            txs = AdvanceTransaction.objects.filter(id__in=ids)
            for tx in txs:
                month = _month_from_date_obj(tx.advance_date)
                if month in locked_months:
                    return Response({'error': f'Tháng {month} đã khóa, không thể xóa phiếu tạm ứng.'}, status=400)
                if tx.approval_status in {AdvanceTransaction.APPROVAL_PENDING_L1, AdvanceTransaction.APPROVAL_PENDING_L2, AdvanceTransaction.APPROVAL_APPROVED}:
                    return Response({'error': f'Phiếu {tx.code} đã gửi duyệt/đã duyệt, không thể xóa.'}, status=400)
        deleted_rows = list(AdvanceTransaction.objects.filter(id__in=ids).values_list('id', 'code'))
        deleted_count, _ = AdvanceTransaction.objects.filter(id__in=ids).delete()
        if deleted_rows:
            _log_finance_audit(
                request.user,
                action='DELETE',
                entity_type='FinanceAdvanceApprovalBulkDelete',
                entity_id=0,
                entity_code='bulk_delete',
                old_values={'items': [{'id': int(row_id), 'code': code} for row_id, code in deleted_rows]},
                new_values={},
                changed_fields=['deleted_items'],
            )
        return Response({'success': True, 'count': deleted_count})

    def destroy(self, request, *args, **kwargs):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền xóa phiếu tạm ứng.'}, status=403)
        instance = self.get_object()
        _ensure_finance_month_unlocked(
            _month_from_date_obj(instance.advance_date),
            'Tháng tài chính đã khóa, không thể xóa phiếu tạm ứng.',
        )
        if instance.approval_status in {
            AdvanceTransaction.APPROVAL_PENDING_L1,
            AdvanceTransaction.APPROVAL_PENDING_L2,
            AdvanceTransaction.APPROVAL_APPROVED,
        }:
            return Response({'error': 'Phiếu đã gửi duyệt/đã duyệt, không thể xóa.'}, status=400)
        snap = {
            'code': instance.code,
            'advance_date': instance.advance_date.isoformat() if instance.advance_date else '',
            'recipient_name': instance.recipient_name,
            'amount': str(instance.amount),
        }
        response = super().destroy(request, *args, **kwargs)
        _log_finance_audit(
            request.user,
            action='DELETE',
            entity_type='FinanceAdvanceApproval',
            entity_id=int(instance.id),
            entity_code=instance.code,
            old_values=snap,
            new_values={},
            changed_fields=['deleted'],
        )
        return response

    @action(detail=False, methods=['get'])
    def approval_queue(self, request):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền xem hàng đợi duyệt.'}, status=403)
        pending_l1_qs = AdvanceTransaction.objects.filter(approval_status=AdvanceTransaction.APPROVAL_PENDING_L1, is_active=True)
        pending_l2_qs = AdvanceTransaction.objects.filter(approval_status=AdvanceTransaction.APPROVAL_PENDING_L2, is_active=True)
        if not _can_approve_level2(request.user):
            pending_l2_qs = pending_l2_qs.none()
        items = []
        for adv in list(pending_l1_qs.order_by('-advance_date', '-id')[:10]) + list(pending_l2_qs.order_by('-advance_date', '-id')[:10]):
            items.append({
                'id': int(adv.id),
                'code': adv.code,
                'recipient_name': adv.recipient_name,
                'amount': str(adv.amount),
                'advance_date': adv.advance_date.isoformat() if adv.advance_date else '',
                'approval_status': adv.approval_status,
                'required_approval_level': int(adv.required_approval_level or 1),
            })
        return Response({
            'pending_l1_count': int(pending_l1_qs.count()),
            'pending_l2_count': int(pending_l2_qs.count()),
            'items': items,
        })

    @action(detail=False, methods=['get'])
    def approval_sla_overview(self, request):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền xem SLA duyệt phiếu tạm ứng.'}, status=403)
        overview = build_finance_approval_sla_overview()
        return Response(overview)

    @action(detail=False, methods=['get'])
    def executive_kpi(self, request):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền xem KPI điều hành tài chính.'}, status=403)
        finance_sla = build_finance_approval_sla_overview()
        workforce_sla = build_salary_advance_approval_sla_overview()
        overdue_90 = build_overdue_snapshot(as_of=timezone.localdate(), threshold_days=90)
        today = timezone.now()
        y = today.year
        m = today.month
        months: list[tuple[int, int, str]] = []
        for _ in range(6):
            months.append((y, m, f'{y}-{str(m).zfill(2)}'))
            m -= 1
            if m == 0:
                y -= 1
                m = 12
        months.reverse()
        points = []
        for year_value, month_value, key in months:
            fin_pending_l1 = AdvanceTransaction.objects.filter(
                approval_status=AdvanceTransaction.APPROVAL_PENDING_L1,
                submitted_at__year=year_value,
                submitted_at__month=month_value,
                is_active=True,
            ).count()
            fin_pending_l2 = AdvanceTransaction.objects.filter(
                approval_status=AdvanceTransaction.APPROVAL_PENDING_L2,
                submitted_at__year=year_value,
                submitted_at__month=month_value,
                is_active=True,
            ).count()
            wf_pending_l1 = SalaryAdvanceRecord.objects.filter(
                approval_status=SalaryAdvanceRecord.APPROVAL_PENDING_L1,
                submitted_at__year=year_value,
                submitted_at__month=month_value,
                is_active=True,
            ).count()
            wf_pending_l2 = SalaryAdvanceRecord.objects.filter(
                approval_status=SalaryAdvanceRecord.APPROVAL_PENDING_L2,
                submitted_at__year=year_value,
                submitted_at__month=month_value,
                is_active=True,
            ).count()
            points.append({
                'month': key,
                'finance_pending': int(fin_pending_l1 + fin_pending_l2),
                'workforce_pending': int(wf_pending_l1 + wf_pending_l2),
            })
        risk_score = int(
            (int(finance_sla.get('overdue_l1_count') or 0) * 2)
            + (int(finance_sla.get('overdue_l2_count') or 0) * 3)
            + (int(workforce_sla.get('overdue_l1_count') or 0) * 2)
            + (int(workforce_sla.get('overdue_l2_count') or 0) * 3)
            + (int(overdue_90.get('count') or 0) * 2)
        )
        contributors = [
            {
                'key': 'finance_overdue_l2',
                'label': 'Finance overdue L2',
                'count': int(finance_sla.get('overdue_l2_count') or 0),
                'weight': 3,
            },
            {
                'key': 'finance_overdue_l1',
                'label': 'Finance overdue L1',
                'count': int(finance_sla.get('overdue_l1_count') or 0),
                'weight': 2,
            },
            {
                'key': 'workforce_overdue_l2',
                'label': 'Workforce overdue L2',
                'count': int(workforce_sla.get('overdue_l2_count') or 0),
                'weight': 3,
            },
            {
                'key': 'workforce_overdue_l1',
                'label': 'Workforce overdue L1',
                'count': int(workforce_sla.get('overdue_l1_count') or 0),
                'weight': 2,
            },
            {
                'key': 'finance_advance_overdue_90d',
                'label': 'Advance overdue >= 90 ngày',
                'count': int(overdue_90.get('count') or 0),
                'weight': 2,
            },
        ]
        for item in contributors:
            item['impact_score'] = int(item['count']) * int(item['weight'])
        contributors = sorted(contributors, key=lambda x: (-int(x['impact_score']), x['key']))
        top_contributors = [item for item in contributors if int(item.get('impact_score') or 0) > 0][:5]

        risk_level = 'HIGH' if risk_score >= 25 else ('MEDIUM' if risk_score >= 10 else 'LOW')
        current_pending_total = int(finance_sla.get('pending_l1_count') or 0) + int(finance_sla.get('pending_l2_count') or 0) + int(workforce_sla.get('pending_l1_count') or 0) + int(workforce_sla.get('pending_l2_count') or 0)
        previous_month_pending_total = 0
        previous_week_pending_total = 0
        previous_week_start = timezone.now() - timedelta(days=14)
        previous_week_end = timezone.now() - timedelta(days=7)
        previous_month = timezone.now() - timedelta(days=30)
        previous_month_pending_total += AdvanceTransaction.objects.filter(
            approval_status__in=[AdvanceTransaction.APPROVAL_PENDING_L1, AdvanceTransaction.APPROVAL_PENDING_L2],
            submitted_at__year=previous_month.year,
            submitted_at__month=previous_month.month,
            is_active=True,
        ).count()
        previous_month_pending_total += SalaryAdvanceRecord.objects.filter(
            approval_status__in=[SalaryAdvanceRecord.APPROVAL_PENDING_L1, SalaryAdvanceRecord.APPROVAL_PENDING_L2],
            submitted_at__year=previous_month.year,
            submitted_at__month=previous_month.month,
            is_active=True,
        ).count()
        previous_week_pending_total += AdvanceTransaction.objects.filter(
            approval_status__in=[AdvanceTransaction.APPROVAL_PENDING_L1, AdvanceTransaction.APPROVAL_PENDING_L2],
            submitted_at__gte=previous_week_start,
            submitted_at__lt=previous_week_end,
            is_active=True,
        ).count()
        previous_week_pending_total += SalaryAdvanceRecord.objects.filter(
            approval_status__in=[SalaryAdvanceRecord.APPROVAL_PENDING_L1, SalaryAdvanceRecord.APPROVAL_PENDING_L2],
            submitted_at__gte=previous_week_start,
            submitted_at__lt=previous_week_end,
            is_active=True,
        ).count()
        mom_delta_pending = int(current_pending_total - previous_month_pending_total)
        wow_delta_pending = int(current_pending_total - previous_week_pending_total)
        early_warning_threshold = 25 if risk_level == 'MEDIUM' else (10 if risk_level == 'LOW' else 25)
        risk_to_next_level = max(0, int(early_warning_threshold - risk_score))
        early_warning = {
            'is_triggered': bool(risk_level != 'HIGH' and risk_to_next_level <= 3),
            'score_to_next_level': int(risk_to_next_level),
            'target_level': 'HIGH' if risk_level == 'MEDIUM' else ('MEDIUM' if risk_level == 'LOW' else 'HIGH'),
            'hint': (
                'Rủi ro đang tiệm cận ngưỡng cấp kế tiếp, cần chạy nhắc SLA và review nhóm tắc nghẽn.'
                if risk_level != 'HIGH' and risk_to_next_level <= 3
                else 'Rủi ro còn khoảng đệm an toàn.'
            ),
        }
        priority_queue = []
        if int(finance_sla.get('overdue_l2_count') or 0) > 0:
            priority_queue.append({
                'code': 'finance-l2-overdue',
                'title': 'Xử lý pending Finance L2 quá hạn',
                'owner': 'Finance Manager',
                'impact_score': int(finance_sla.get('overdue_l2_count') or 0) * 3,
                'quick_action': 'RUN_FINANCE_SLA_REMINDER',
            })
        if int(workforce_sla.get('overdue_l2_count') or 0) > 0:
            priority_queue.append({
                'code': 'workforce-l2-overdue',
                'title': 'Xử lý pending Workforce L2 quá hạn',
                'owner': 'HR/Payroll Manager',
                'impact_score': int(workforce_sla.get('overdue_l2_count') or 0) * 3,
                'quick_action': 'RUN_WORKFORCE_SLA_REMINDER',
            })
        if int(overdue_90.get('count') or 0) > 0:
            priority_queue.append({
                'code': 'advance-overdue-90d',
                'title': 'Rà soát tạm ứng tồn đọng >= 90 ngày',
                'owner': 'Finance + Ops',
                'impact_score': int(overdue_90.get('count') or 0) * 2,
                'quick_action': 'ESCALATE_MANAGEMENT_REVIEW',
            })
        priority_queue = sorted(priority_queue, key=lambda x: (-int(x['impact_score']), x['code']))[:5]
        recommendations: list[dict] = []
        if risk_level == 'HIGH':
            recommendations.append({
                'code': 'RUN_FINANCE_SLA_REMINDER',
                'title': 'Kích hoạt nhắc SLA Finance ngay',
                'priority': 'P0',
                'description': 'Gửi nhắc toàn bộ pending Finance để kéo giảm backlog duyệt cấp cao.',
            })
            recommendations.append({
                'code': 'RUN_WORKFORCE_SLA_REMINDER',
                'title': 'Kích hoạt nhắc SLA Workforce ngay',
                'priority': 'P0',
                'description': 'Gửi nhắc ứng lương pending để giảm tắc nghẽn HR/Payroll.',
            })
            recommendations.append({
                'code': 'ESCALATE_MANAGEMENT_REVIEW',
                'title': 'Tổ chức review escalation trong ngày',
                'priority': 'P1',
                'description': 'Xử lý nhóm L2 overdue và case >= 90 ngày theo danh sách ưu tiên.',
            })
        elif risk_level == 'MEDIUM':
            recommendations.append({
                'code': 'RUN_TARGETED_SLA_REMINDER',
                'title': 'Nhắc có chọn lọc theo nhóm quá hạn',
                'priority': 'P1',
                'description': 'Ưu tiên nhóm submitter có pending nhiều và max wait cao.',
            })
            recommendations.append({
                'code': 'TUNE_SLA_POLICY',
                'title': 'Tinh chỉnh policy SLA/cooldown',
                'priority': 'P2',
                'description': 'Giảm ngưỡng cảnh báo hoặc rút ngắn remind interval cho giai đoạn cao điểm.',
            })
        else:
            recommendations.append({
                'code': 'KEEP_MONITORING',
                'title': 'Duy trì giám sát hiện tại',
                'priority': 'P3',
                'description': 'Theo dõi trend 6 tháng và giữ cadence nhắc định kỳ.',
            })
        return Response({
            'as_of': timezone.localdate().isoformat(),
            'finance_sla': finance_sla,
            'workforce_sla': workforce_sla,
            'finance_overdue_90': overdue_90,
            'trend_6m': points,
            'risk_score': risk_score,
            'risk_level': risk_level,
            'risk_contributors': top_contributors,
            'recommendations': recommendations,
            'risk_trend': {
                'mom_delta_pending': int(mom_delta_pending),
                'wow_delta_pending': int(wow_delta_pending),
                'current_pending_total': int(current_pending_total),
                'previous_month_pending_total': int(previous_month_pending_total),
                'previous_week_pending_total': int(previous_week_pending_total),
            },
            'early_warning': early_warning,
            'priority_queue': priority_queue,
            'auto_policy': _get_executive_auto_policy(),
        })

    @action(detail=False, methods=['get'])
    def cross_module_readiness(self, request):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền xem readiness vận hành.'}, status=403)
        return Response(_build_cross_module_readiness_payload())

    @action(detail=False, methods=['post'])
    def cross_module_bootstrap(self, request):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền bootstrap dữ liệu vận hành.'}, status=403)

        dry_run = str(request.data.get('dry_run') or '').strip().lower() in {'1', 'true', 'yes'}
        actor = request.user
        readiness_before = _build_cross_module_readiness_payload()
        today = timezone.localdate()
        month = f'{today.year}-{str(today.month).zfill(2)}'
        result = {
            'dry_run': bool(dry_run),
            'created': {
                'permissions': 0,
                'roles': 0,
                'employees': 0,
                'finance_advances': 0,
                'workforce_advances': 0,
                'payroll_records': 0,
                'role_permission_audits': 0,
                'operations_audits': 0,
                'pipeline_events': 0,
            },
            'skipped': {
                'permissions': 0,
                'roles': 0,
                'employees': 0,
                'finance_advances': 0,
                'workforce_advances': 0,
                'payroll_records': 0,
                'role_permission_audits': 0,
                'operations_audits': 0,
                'pipeline_events': 0,
            },
            'notes': [],
        }

        perm_specs = [
            {'resource': 'WORKFORCE', 'action': 'MANAGE', 'code': 'WORKFORCE_MANAGE', 'name': 'Workforce Manage'},
            {'resource': 'FINANCE', 'action': 'MANAGE', 'code': 'FINANCE_MANAGE', 'name': 'Finance Manage'},
            {'resource': 'CORE', 'action': 'MANAGE_RBAC', 'code': 'CORE_MANAGE_RBAC', 'name': 'RBAC Manage'},
        ]
        perm_map = {}
        for spec in perm_specs:
            existing = Permission.objects.filter(resource=spec['resource'], action=spec['action']).first()
            if existing:
                perm_map[f"{spec['resource']}:{spec['action']}"] = existing
                result['skipped']['permissions'] += 1
                continue
            if dry_run:
                result['created']['permissions'] += 1
                continue
            row, created = Permission.objects.get_or_create(
                resource=spec['resource'],
                action=spec['action'],
                defaults={'name': spec['name'], 'code': spec['code'], 'description': f"Permission {spec['code']}"},
            )
            perm_map[f"{spec['resource']}:{spec['action']}"] = row
            if created:
                result['created']['permissions'] += 1
            else:
                result['skipped']['permissions'] += 1

        role_specs = [
            {'code': 'MANAGER', 'name': 'Manager', 'sort_order': 10, 'perms': ['WORKFORCE:MANAGE', 'FINANCE:MANAGE', 'CORE:MANAGE_RBAC']},
            {'code': 'FINANCE_MANAGER', 'name': 'Finance Manager', 'sort_order': 20, 'perms': ['FINANCE:MANAGE']},
            {'code': 'HR_MANAGER', 'name': 'HR Manager', 'sort_order': 30, 'perms': ['WORKFORCE:MANAGE']},
        ]
        role_rows = []
        for spec in role_specs:
            existing = Role.objects.filter(code=spec['code'], deleted_at__isnull=True).first()
            if existing:
                role_rows.append(existing)
                result['skipped']['roles'] += 1
                continue
            if dry_run:
                result['created']['roles'] += 1
                continue
            role_row = Role.objects.create(
                code=spec['code'],
                name=spec['name'],
                description='Bootstrap seed role',
                is_active=True,
                sort_order=spec['sort_order'],
                created_by=actor,
                updated_by=actor,
            )
            role_rows.append(role_row)
            result['created']['roles'] += 1

        if not dry_run:
            role_map = {r.code: r for r in role_rows}
            for spec in role_specs:
                role_row = role_map.get(spec['code'])
                if not role_row:
                    continue
                for perm_key in spec['perms']:
                    perm_row = perm_map.get(perm_key)
                    if perm_row:
                        role_row.permissions.add(perm_row)

        cash_account = None
        cash_account = CashAccount.objects.filter(account_type=CashAccount.TYPE_CASH, is_active=True).order_by('id').first()
        if not cash_account and not dry_run:
            cash_account = CashAccount.objects.create(
                name='Quy bootstrap',
                account_type=CashAccount.TYPE_CASH,
                balance=Decimal('50000000'),
                note='Seed by cross_module_bootstrap',
                is_active=True,
                created_by=actor,
                updated_by=actor,
            )

        employee_specs = [
            {'code': 'EMP-DEMO-01', 'name': 'Nhan vien Demo 01', 'salary_basic': Decimal('9000000')},
            {'code': 'EMP-DEMO-02', 'name': 'Nhan vien Demo 02', 'salary_basic': Decimal('10500000')},
        ]
        employee_rows = []
        for spec in employee_specs:
            existing = Employee.objects.filter(code=spec['code']).first()
            if existing:
                employee_rows.append(existing)
                result['skipped']['employees'] += 1
                continue
            if dry_run:
                result['created']['employees'] += 1
                continue
            emp = Employee.objects.create(
                code=spec['code'],
                name=spec['name'],
                status=Employee.STATUS_ACTIVE,
                salary_basic=spec['salary_basic'],
                start_date=today,
                is_active=True,
                created_by=actor,
                updated_by=actor,
            )
            employee_rows.append(emp)
            result['created']['employees'] += 1

        finance_advance_specs = [
            {'code': 'ADV-DEMO-001', 'recipient_name': 'Nguoi nhan demo 01', 'amount': Decimal('2000000'), 'approval_status': AdvanceTransaction.APPROVAL_PENDING_L1},
            {'code': 'ADV-DEMO-002', 'recipient_name': 'Nguoi nhan demo 02', 'amount': Decimal('65000000'), 'approval_status': AdvanceTransaction.APPROVAL_PENDING_L2},
        ]
        for spec in finance_advance_specs:
            existing = AdvanceTransaction.objects.filter(code=spec['code']).first()
            if existing:
                result['skipped']['finance_advances'] += 1
                continue
            if dry_run:
                result['created']['finance_advances'] += 1
                continue
            if not cash_account:
                result['notes'].append('Không có tài khoản quỹ đang dùng để tạo phiếu tạm ứng demo.')
                result['skipped']['finance_advances'] += 1
                continue
            required_level = _required_approval_level_for_amount(spec['amount'])
            submitted_at = timezone.now() - timedelta(hours=12 if spec['approval_status'] == AdvanceTransaction.APPROVAL_PENDING_L1 else 30)
            AdvanceTransaction.objects.create(
                code=spec['code'],
                advance_type=AdvanceTransaction.TYPE_OTHER,
                advance_date=today,
                recipient_name=spec['recipient_name'],
                source_type=AdvanceTransaction.SOURCE_CASH,
                source_cash_account=cash_account,
                amount=spec['amount'],
                purpose='Bootstrap demo advance',
                status=AdvanceTransaction.STATUS_OPEN,
                approval_status=spec['approval_status'],
                required_approval_level=required_level,
                submitted_at=submitted_at,
                submitted_by=actor,
                is_active=True,
                created_by=actor,
                updated_by=actor,
            )
            result['created']['finance_advances'] += 1

        if not employee_rows:
            employee_rows = list(Employee.objects.filter(code__in=[item['code'] for item in employee_specs]))
        for idx, emp in enumerate(employee_rows[:2]):
            existing = SalaryAdvanceRecord.objects.filter(
                employee=emp,
                advance_date=today,
                month=month,
                amount=Decimal('1500000') + Decimal(str(idx * 500000)),
            ).first()
            if existing:
                result['skipped']['workforce_advances'] += 1
            elif dry_run:
                result['created']['workforce_advances'] += 1
            else:
                approval_status = SalaryAdvanceRecord.APPROVAL_APPROVED if idx == 0 else SalaryAdvanceRecord.APPROVAL_PENDING_L1
                SalaryAdvanceRecord.objects.create(
                    employee=emp,
                    advance_date=today,
                    month=month,
                    amount=Decimal('1500000') + Decimal(str(idx * 500000)),
                    reason='Bootstrap salary advance',
                    status=SalaryAdvanceRecord.STATUS_UNDEDUCTED,
                    approval_status=approval_status,
                    required_approval_level=1 if idx == 0 else 2,
                    submitted_at=timezone.now() - timedelta(hours=8),
                    submitted_by=actor,
                    is_active=True,
                    created_by=actor,
                    updated_by=actor,
                )
                result['created']['workforce_advances'] += 1

            existing_payroll = PayrollRecord.objects.filter(employee=emp, month=month).first()
            if existing_payroll:
                result['skipped']['payroll_records'] += 1
            elif dry_run:
                result['created']['payroll_records'] += 1
            else:
                basic_salary = Decimal(str(emp.salary_basic or 0))
                PayrollRecord.objects.create(
                    employee=emp,
                    month=month,
                    standard_days=Decimal('26'),
                    actual_days=Decimal('26'),
                    basic_salary=basic_salary,
                    salary_by_attendance=basic_salary,
                    overtime_pay=Decimal('0'),
                    total_bonus=Decimal('500000'),
                    total_penalty=Decimal('0'),
                    advance_deduction=Decimal('0'),
                    total_income=basic_salary + Decimal('500000'),
                    total_deductions=Decimal('0'),
                    net_pay=basic_salary + Decimal('500000'),
                    status=PayrollRecord.STATUS_UNLOCKED,
                    note='Bootstrap payroll record',
                    created_by=actor,
                    updated_by=actor,
                )
                result['created']['payroll_records'] += 1

        if AuditLog.objects.filter(entity_type='RoleModulePermission').exists():
            result['skipped']['role_permission_audits'] += 1
        elif dry_run:
            result['created']['role_permission_audits'] += 1
        else:
            role_items = []
            for role_row in Role.objects.filter(code__in=['MANAGER', 'FINANCE_MANAGER', 'HR_MANAGER'], deleted_at__isnull=True).prefetch_related('permissions'):
                perm_codes = {f'{p.resource}:{p.action}' for p in role_row.permissions.all()}
                role_items.append({
                    'role_id': int(role_row.id),
                    'role_code': role_row.code,
                    'role_name': role_row.name,
                    'is_active': bool(role_row.is_active),
                    'workforce_manage': 'WORKFORCE:MANAGE' in perm_codes,
                    'finance_manage': 'FINANCE:MANAGE' in perm_codes,
                    'rbac_manage': 'CORE:MANAGE_RBAC' in perm_codes,
                })
            AuditLog.objects.create(
                user=actor,
                action='UPDATE',
                entity_type='RoleModulePermission',
                entity_id=0,
                entity_code='ROLE_MODULE_PERMISSIONS',
                old_values={'items': []},
                new_values={'items': role_items, 'source': 'cross_module_bootstrap'},
                changed_fields=['items'],
            )
            result['created']['role_permission_audits'] += 1

        has_ops_seed = AuditLog.objects.filter(entity_type__in=['TaskBulk', 'WorkflowAnalytics', 'Task', 'WorkflowAutomation']).exists()
        if has_ops_seed:
            result['skipped']['operations_audits'] += 1
        elif dry_run:
            result['created']['operations_audits'] += 1
        else:
            AuditLog.objects.create(
                user=actor,
                action='UPDATE',
                entity_type='TaskBulk',
                entity_id=0,
                entity_code='TASK_BULK',
                old_values={},
                new_values={
                    'action': 'REMIND_OVERDUE',
                    'success_count': 4,
                    'failed_count': 0,
                    'processed_count': 4,
                    'total_requested': 4,
                    'reminder_sent_count': 3,
                },
                changed_fields=['bulk_action'],
            )
            AuditLog.objects.create(
                user=actor,
                action='UPDATE',
                entity_type='WorkflowAnalytics',
                entity_id=1,
                entity_code='ANL-BOOTSTRAP',
                old_values={},
                new_values={
                    'insight_type': 'READINESS_BOOTSTRAP',
                    'suggested_action': 'RUN_AUTOMATION',
                    'success': True,
                    'message': 'Bootstrap readiness data seeded',
                },
                changed_fields=['execute_insight'],
            )
            result['created']['operations_audits'] += 2

        if WorkflowPipelineEvent.objects.exists():
            result['skipped']['pipeline_events'] += 1
        elif dry_run:
            result['created']['pipeline_events'] += 1
        else:
            WorkflowPipelineEvent.objects.create(
                entity_type='SalesOrder',
                entity_id=999,
                entity_code='SO-BOOTSTRAP',
                trigger='SUBMIT',
                action=WorkflowPipelineEvent.ACTION_ADVANCE,
                from_step='NEW',
                to_step='PLAN',
                note='Seed pipeline event for operations log',
                actor=actor,
            )
            result['created']['pipeline_events'] += 1

        result['created_total'] = int(sum(int(v) for v in result['created'].values()))
        result['skipped_total'] = int(sum(int(v) for v in result['skipped'].values()))
        result['readiness_after'] = _build_cross_module_readiness_payload()
        _log_finance_audit(
            actor,
            action='UPDATE',
            entity_type=CROSS_MODULE_BOOTSTRAP_AUDIT_ENTITY_TYPE,
            entity_id=0,
            entity_code=CROSS_MODULE_BOOTSTRAP_AUDIT_ENTITY_CODE,
            old_values={},
            new_values={
                'dry_run': bool(dry_run),
                'created': result.get('created') if isinstance(result.get('created'), dict) else {},
                'skipped': result.get('skipped') if isinstance(result.get('skipped'), dict) else {},
                'created_total': int(result.get('created_total') or 0),
                'skipped_total': int(result.get('skipped_total') or 0),
                'notes': result.get('notes') if isinstance(result.get('notes'), list) else [],
                'readiness_before': readiness_before,
                'readiness_after': result.get('readiness_after') if isinstance(result.get('readiness_after'), dict) else {},
            },
            changed_fields=['dry_run', 'created_total', 'skipped_total', 'created', 'skipped', 'readiness_before', 'readiness_after'],
        )
        return Response(result)

    @action(detail=False, methods=['get'])
    def cross_module_bootstrap_history(self, request):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền xem lịch sử bootstrap dữ liệu.'}, status=403)
        limit_raw = str(request.query_params.get('limit') or '20').strip()
        days_raw = str(request.query_params.get('days') or '').strip()
        username_filter = str(request.query_params.get('username') or '').strip()
        dry_run_raw = str(request.query_params.get('dry_run') or '').strip().lower()
        try:
            limit = max(1, min(int(limit_raw), 200))
        except Exception:
            limit = 20
        logs = _cross_module_bootstrap_history_qs()
        if days_raw:
            try:
                days = max(1, min(int(days_raw), 365))
                cutoff = timezone.localdate() - timedelta(days=days - 1)
                logs = logs.filter(created_at__date__gte=cutoff)
            except Exception:
                pass
        if username_filter:
            logs = logs.filter(user__username__icontains=username_filter)
        if dry_run_raw in {'true', '1', 'yes'}:
            logs = logs.filter(new_values__dry_run=True)
        elif dry_run_raw in {'false', '0', 'no'}:
            logs = logs.filter(new_values__dry_run=False)
        logs = logs.order_by('-created_at', '-id')
        items = [_serialize_cross_module_bootstrap_history_row(row) for row in logs]
        export_excel = str(request.query_params.get('export') or '').strip().lower() == 'excel'
        if export_excel:
            wb = Workbook()
            ws = wb.active
            ws.title = 'BootstrapHistory'
            ws.append([
                'Created At',
                'Username',
                'Dry Run',
                'Created Total',
                'Skipped Total',
                'Readiness Before',
                'Readiness After',
                'Readiness Delta',
                'Improved',
                'Level Before',
                'Level After',
                'Notes',
            ])
            for row in items:
                ws.append([
                    row.get('created_at') or '',
                    row.get('username') or '',
                    'Yes' if bool(row.get('dry_run')) else 'No',
                    int(row.get('created_total') or 0),
                    int(row.get('skipped_total') or 0),
                    int((row.get('readiness_before') or {}).get('readiness_score') or 0),
                    int((row.get('readiness_after') or {}).get('readiness_score') or 0),
                    int(row.get('readiness_delta') or 0),
                    'Yes' if bool(row.get('improved')) else 'No',
                    str(row.get('level_before') or ''),
                    str(row.get('level_after') or ''),
                    ' | '.join(str(note) for note in (row.get('notes') or [])),
                ])
            response = HttpResponse(
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )
            filename = f'cross_module_bootstrap_history_{timezone.localdate().isoformat()}.xlsx'
            response['Content-Disposition'] = f'attachment; filename="{filename}"'
            wb.save(response)
            return response
        items = items[:limit]
        dry_run_filter_value = ''
        if dry_run_raw in {'true', '1', 'yes'}:
            dry_run_filter_value = 'true'
        elif dry_run_raw in {'false', '0', 'no'}:
            dry_run_filter_value = 'false'
        return Response({
            'count': len(items),
            'filters': {
                'days': int(days_raw) if days_raw.isdigit() else None,
                'username': username_filter,
                'dry_run': dry_run_filter_value,
            },
            'items': items,
        })

    @action(detail=False, methods=['get'])
    def executive_auto_policy(self, request):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền xem policy tự động điều hành.'}, status=403)
        return Response(_get_executive_auto_policy())

    @executive_auto_policy.mapping.post
    def save_executive_auto_policy(self, request):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền cập nhật policy tự động điều hành.'}, status=403)
        payload = request.data if isinstance(request.data, dict) else {}
        old_policy = _get_executive_auto_policy()
        policy = _save_executive_auto_policy(payload)
        _log_finance_audit(
            request.user,
            action='UPDATE',
            entity_type='FinanceExecutiveAutoPolicy',
            entity_id=0,
            entity_code=FINANCE_EXECUTIVE_AUTO_POLICY_KEY,
            old_values=old_policy,
            new_values=policy,
            changed_fields=['enabled', 'cooldown_minutes', 'auto_run_finance_sla', 'auto_run_workforce_sla', 'only_when_early_warning', 'last_run_at'],
        )
        return Response({'success': True, 'policy': policy})

    @action(detail=False, methods=['post'])
    def executive_auto_execute(self, request):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền chạy tự động điều hành.'}, status=403)
        force_run = str(request.data.get('force') or '').strip().lower() in {'1', 'true', 'yes'}
        return Response(_execute_executive_auto(actor=request.user, force_run=force_run, source='api'))

    @action(detail=False, methods=['get'])
    def executive_auto_history(self, request):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền xem lịch sử auto-execute.'}, status=403)
        limit_raw = str(request.query_params.get('limit') or '30').strip()
        try:
            limit = max(1, min(int(limit_raw), 200))
        except Exception:
            limit = 30
        logs = _executive_auto_history_qs().order_by('-created_at', '-id')[:limit]
        items = [_serialize_executive_auto_history_row(row) for row in logs]
        return Response({'count': len(items), 'items': items})

    @action(detail=False, methods=['get'])
    def executive_auto_governance(self, request):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền xem governance auto-execute.'}, status=403)
        days_raw = str(request.query_params.get('days') or '30').strip()
        try:
            days = max(7, min(int(days_raw), 365))
        except Exception:
            days = 30
        group_by = str(request.query_params.get('group_by') or 'day').strip().lower()
        if group_by not in {'day', 'week'}:
            group_by = 'day'
        cutoff = timezone.localdate() - timedelta(days=days - 1)
        logs = _executive_auto_history_qs().filter(created_at__date__gte=cutoff).order_by('created_at', 'id')

        def _empty_summary() -> dict:
            return {
                'total_runs': 0,
                'success_runs': 0,
                'skipped_runs': 0,
                'failed_runs': 0,
                'finance_sent_total': 0,
                'workforce_sent_total': 0,
                'sent_total': 0,
                'success_rate': 0.0,
                'skipped_rate': 0.0,
                'failed_rate': 0.0,
                'avg_sent_per_run': 0.0,
            }

        summary = _empty_summary()
        buckets: dict[str, dict] = {}
        skip_reason_counts: dict[str, int] = {}
        action_effectiveness_map: dict[str, dict] = {
            'BOTH': {'action': 'BOTH', 'total_runs': 0, 'success_runs': 0, 'skipped_runs': 0, 'failed_runs': 0, 'finance_sent_total': 0, 'workforce_sent_total': 0, 'sent_total': 0, 'avg_sent_per_run': 0.0},
            'FINANCE_ONLY': {'action': 'FINANCE_ONLY', 'total_runs': 0, 'success_runs': 0, 'skipped_runs': 0, 'failed_runs': 0, 'finance_sent_total': 0, 'workforce_sent_total': 0, 'sent_total': 0, 'avg_sent_per_run': 0.0},
            'WORKFORCE_ONLY': {'action': 'WORKFORCE_ONLY', 'total_runs': 0, 'success_runs': 0, 'skipped_runs': 0, 'failed_runs': 0, 'finance_sent_total': 0, 'workforce_sent_total': 0, 'sent_total': 0, 'avg_sent_per_run': 0.0},
            'NO_SENT': {'action': 'NO_SENT', 'total_runs': 0, 'success_runs': 0, 'skipped_runs': 0, 'failed_runs': 0, 'finance_sent_total': 0, 'workforce_sent_total': 0, 'sent_total': 0, 'avg_sent_per_run': 0.0},
        }

        for row in logs:
            item = _serialize_executive_auto_history_row(row)
            created_at = row.created_at
            if not created_at:
                continue
            period_key, period_start, period_end = _executive_auto_period_info(created_at, group_by)
            bucket = buckets.get(period_key)
            if bucket is None:
                bucket = {
                    'period_key': period_key,
                    'period_start': period_start,
                    'period_end': period_end,
                    'total_runs': 0,
                    'success_runs': 0,
                    'skipped_runs': 0,
                    'failed_runs': 0,
                    'finance_sent_total': 0,
                    'workforce_sent_total': 0,
                    'sent_total': 0,
                    'success_rate': 0.0,
                    'skipped_rate': 0.0,
                    'failed_rate': 0.0,
                    'avg_sent_per_run': 0.0,
                }
                buckets[period_key] = bucket

            finance_sent = int(item['finance_sent_count'])
            workforce_sent = int(item['workforce_sent_count'])
            sent_total = int(finance_sent + workforce_sent)
            is_success = bool(item['success'])
            is_skipped = bool(item['skipped'])
            is_failed = bool((not is_success) and (not is_skipped))
            reason = str(item.get('reason') or '')

            summary['total_runs'] += 1
            bucket['total_runs'] += 1
            if is_success:
                summary['success_runs'] += 1
                bucket['success_runs'] += 1
            elif is_skipped:
                summary['skipped_runs'] += 1
                bucket['skipped_runs'] += 1
            else:
                summary['failed_runs'] += 1
                bucket['failed_runs'] += 1
            summary['finance_sent_total'] += finance_sent
            summary['workforce_sent_total'] += workforce_sent
            summary['sent_total'] += sent_total
            bucket['finance_sent_total'] += finance_sent
            bucket['workforce_sent_total'] += workforce_sent
            bucket['sent_total'] += sent_total

            if is_skipped and reason:
                skip_reason_counts[reason] = int(skip_reason_counts.get(reason) or 0) + 1

            if finance_sent > 0 and workforce_sent > 0:
                action_key = 'BOTH'
            elif finance_sent > 0:
                action_key = 'FINANCE_ONLY'
            elif workforce_sent > 0:
                action_key = 'WORKFORCE_ONLY'
            else:
                action_key = 'NO_SENT'
            action_bucket = action_effectiveness_map[action_key]
            action_bucket['total_runs'] += 1
            if is_success:
                action_bucket['success_runs'] += 1
            elif is_skipped:
                action_bucket['skipped_runs'] += 1
            else:
                action_bucket['failed_runs'] += 1
            action_bucket['finance_sent_total'] += finance_sent
            action_bucket['workforce_sent_total'] += workforce_sent
            action_bucket['sent_total'] += sent_total

        def _finalize_rates(entry: dict):
            total_runs = int(entry.get('total_runs') or 0)
            success_runs = int(entry.get('success_runs') or 0)
            skipped_runs = int(entry.get('skipped_runs') or 0)
            failed_runs = int(entry.get('failed_runs') or 0)
            sent_total = int(entry.get('sent_total') or 0)
            entry['success_rate'] = round((success_runs / total_runs) * 100, 2) if total_runs > 0 else 0.0
            entry['skipped_rate'] = round((skipped_runs / total_runs) * 100, 2) if total_runs > 0 else 0.0
            entry['failed_rate'] = round((failed_runs / total_runs) * 100, 2) if total_runs > 0 else 0.0
            entry['avg_sent_per_run'] = round((sent_total / total_runs), 2) if total_runs > 0 else 0.0

        _finalize_rates(summary)

        previous_cutoff_end = cutoff - timedelta(days=1)
        previous_cutoff_start = previous_cutoff_end - timedelta(days=days - 1)
        previous_logs = _executive_auto_history_qs().filter(
            created_at__date__gte=previous_cutoff_start,
            created_at__date__lte=previous_cutoff_end,
        )
        previous_summary = _empty_summary()
        for row in previous_logs:
            item = _serialize_executive_auto_history_row(row)
            finance_sent = int(item['finance_sent_count'])
            workforce_sent = int(item['workforce_sent_count'])
            sent_total = int(finance_sent + workforce_sent)
            is_success = bool(item['success'])
            is_skipped = bool(item['skipped'])
            previous_summary['total_runs'] += 1
            if is_success:
                previous_summary['success_runs'] += 1
            elif is_skipped:
                previous_summary['skipped_runs'] += 1
            else:
                previous_summary['failed_runs'] += 1
            previous_summary['finance_sent_total'] += finance_sent
            previous_summary['workforce_sent_total'] += workforce_sent
            previous_summary['sent_total'] += sent_total
        _finalize_rates(previous_summary)

        by_period = sorted(buckets.values(), key=lambda x: x['period_key'])
        for row in by_period:
            _finalize_rates(row)
        skip_reasons = sorted(
            [{'reason': key, 'count': int(value)} for key, value in skip_reason_counts.items()],
            key=lambda x: (-int(x['count']), x['reason']),
        )
        action_effectiveness = []
        for key in ['BOTH', 'FINANCE_ONLY', 'WORKFORCE_ONLY', 'NO_SENT']:
            row = action_effectiveness_map[key]
            _finalize_rates(row)
            action_effectiveness.append(row)

        comparison = {
            'previous_from_date': previous_cutoff_start.isoformat(),
            'previous_to_date': previous_cutoff_end.isoformat(),
            'previous_summary': previous_summary,
            'delta_total_runs': int(summary['total_runs'] - previous_summary['total_runs']),
            'delta_success_rate': round(float(summary['success_rate']) - float(previous_summary['success_rate']), 2),
            'delta_skipped_rate': round(float(summary['skipped_rate']) - float(previous_summary['skipped_rate']), 2),
            'delta_failed_rate': round(float(summary['failed_rate']) - float(previous_summary['failed_rate']), 2),
            'delta_avg_sent_per_run': round(float(summary['avg_sent_per_run']) - float(previous_summary['avg_sent_per_run']), 2),
            'delta_sent_total': int(summary['sent_total'] - previous_summary['sent_total']),
        }

        insights = []
        if float(summary['success_rate']) < 70:
            insights.append({
                'code': 'LOW_SUCCESS_RATE',
                'severity': 'HIGH',
                'message': 'Tỷ lệ success auto-run dưới 70%, cần review policy và điều kiện trigger.',
            })
        if float(summary['skipped_rate']) >= 40:
            top_reason = skip_reasons[0]['reason'] if skip_reasons else ''
            insights.append({
                'code': 'HIGH_SKIPPED_RATE',
                'severity': 'MEDIUM',
                'message': f'Tỷ lệ skipped cao ({summary["skipped_rate"]}%). Lý do chính: {top_reason or "N/A"}.',
            })
        no_sent_bucket = next((item for item in action_effectiveness if item.get('action') == 'NO_SENT'), None)
        if no_sent_bucket and int(no_sent_bucket.get('total_runs') or 0) > 0:
            no_sent_rate = round((int(no_sent_bucket.get('total_runs') or 0) / max(1, int(summary.get('total_runs') or 0))) * 100, 2)
            if no_sent_rate >= 50:
                insights.append({
                    'code': 'HIGH_NO_SENT_RUNS',
                    'severity': 'MEDIUM',
                    'message': f'Có {no_sent_rate}% run không gửi reminder nào; kiểm tra bật/tắt Finance/Workforce SLA trong policy.',
                })
        if float(comparison['delta_success_rate']) >= 5:
            insights.append({
                'code': 'SUCCESS_TREND_UP',
                'severity': 'LOW',
                'message': f'Tỷ lệ success cải thiện {comparison["delta_success_rate"]}% so với kỳ trước.',
            })
        elif float(comparison['delta_success_rate']) <= -5:
            insights.append({
                'code': 'SUCCESS_TREND_DOWN',
                'severity': 'HIGH',
                'message': f'Tỷ lệ success giảm {abs(float(comparison["delta_success_rate"]))}% so với kỳ trước.',
            })

        payload = {
            'filters': {
                'days': days,
                'group_by': group_by,
                'from_date': cutoff.isoformat(),
                'to_date': timezone.localdate().isoformat(),
            },
            'summary': summary,
            'comparison': comparison,
            'insights': insights[:5],
            'by_period': by_period,
            'skip_reasons': skip_reasons[:10],
            'action_effectiveness': action_effectiveness,
        }
        if str(request.query_params.get('export') or '').strip().lower() == 'excel':
            wb = Workbook()
            ws_summary = wb.active
            ws_summary.title = 'Summary'
            ws_summary.append(['Metric', 'Value'])
            ws_summary.append(['days', days])
            ws_summary.append(['group_by', group_by])
            ws_summary.append(['from_date', payload['filters']['from_date']])
            ws_summary.append(['to_date', payload['filters']['to_date']])
            ws_summary.append(['total_runs', summary['total_runs']])
            ws_summary.append(['success_runs', summary['success_runs']])
            ws_summary.append(['skipped_runs', summary['skipped_runs']])
            ws_summary.append(['failed_runs', summary['failed_runs']])
            ws_summary.append(['success_rate_percent', float(summary['success_rate'])])
            ws_summary.append(['skipped_rate_percent', float(summary['skipped_rate'])])
            ws_summary.append(['failed_rate_percent', float(summary['failed_rate'])])
            ws_summary.append(['finance_sent_total', summary['finance_sent_total']])
            ws_summary.append(['workforce_sent_total', summary['workforce_sent_total']])
            ws_summary.append(['sent_total', summary['sent_total']])
            ws_summary.append(['avg_sent_per_run', float(summary['avg_sent_per_run'])])
            ws_summary.append(['delta_success_rate_vs_prev_percent', float(comparison['delta_success_rate'])])
            ws_summary.append(['delta_skipped_rate_vs_prev_percent', float(comparison['delta_skipped_rate'])])
            ws_summary.append(['delta_failed_rate_vs_prev_percent', float(comparison['delta_failed_rate'])])
            ws_summary.append(['delta_avg_sent_per_run_vs_prev', float(comparison['delta_avg_sent_per_run'])])

            ws_period = wb.create_sheet('ByPeriod')
            ws_period.append([
                'period_key', 'period_start', 'period_end', 'total_runs', 'success_runs', 'skipped_runs', 'failed_runs',
                'success_rate_percent', 'skipped_rate_percent', 'failed_rate_percent',
                'finance_sent_total', 'workforce_sent_total', 'sent_total', 'avg_sent_per_run',
            ])
            for row in by_period:
                ws_period.append([
                    row['period_key'],
                    row['period_start'],
                    row['period_end'],
                    int(row['total_runs']),
                    int(row['success_runs']),
                    int(row['skipped_runs']),
                    int(row['failed_runs']),
                    float(row['success_rate']),
                    float(row['skipped_rate']),
                    float(row['failed_rate']),
                    int(row['finance_sent_total']),
                    int(row['workforce_sent_total']),
                    int(row['sent_total']),
                    float(row['avg_sent_per_run']),
                ])

            ws_reasons = wb.create_sheet('SkipReasons')
            ws_reasons.append(['reason', 'count'])
            for row in skip_reasons:
                ws_reasons.append([row['reason'], int(row['count'])])

            ws_action = wb.create_sheet('ActionEffectiveness')
            ws_action.append([
                'action', 'total_runs', 'success_runs', 'skipped_runs', 'failed_runs',
                'success_rate_percent', 'skipped_rate_percent', 'failed_rate_percent',
                'finance_sent_total', 'workforce_sent_total', 'sent_total', 'avg_sent_per_run',
            ])
            for row in action_effectiveness:
                ws_action.append([
                    row['action'],
                    int(row['total_runs']),
                    int(row['success_runs']),
                    int(row['skipped_runs']),
                    int(row['failed_runs']),
                    float(row['success_rate']),
                    float(row['skipped_rate']),
                    float(row['failed_rate']),
                    int(row['finance_sent_total']),
                    int(row['workforce_sent_total']),
                    int(row['sent_total']),
                    float(row['avg_sent_per_run']),
                ])

            ws_insights = wb.create_sheet('Insights')
            ws_insights.append(['code', 'severity', 'message'])
            for row in insights[:20]:
                ws_insights.append([row.get('code') or '', row.get('severity') or '', row.get('message') or ''])

            response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            response['Content-Disposition'] = f'attachment; filename="executive_auto_governance_{group_by}_{days}d.xlsx"'
            wb.save(response)
            return response
        return Response(payload)

    @action(detail=False, methods=['post'])
    def remind_pending_approvals(self, request):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền gửi nhắc SLA duyệt phiếu tạm ứng.'}, status=403)
        dry_run = str(request.data.get('dry_run') or '').strip().lower() in {'1', 'true', 'yes'}
        result = run_finance_approval_sla_reminder_job(dry_run=dry_run)
        return Response(result)

    @action(detail=False, methods=['get'])
    def approval_sla_reminder_history(self, request):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền xem lịch sử nhắc SLA duyệt phiếu tạm ứng.'}, status=403)
        days_raw = str(request.query_params.get('days') or '30').strip()
        try:
            days = max(1, min(int(days_raw), 180))
        except Exception:
            days = 30
        cutoff = timezone.now() - timedelta(days=days)
        qs = (
            Notification.objects
            .filter(
                notification_type='due_date',
                entity_type='FinanceAdvanceApprovalPending',
                created_at__gte=cutoff,
            )
            .select_related('recipient')
            .order_by('-created_at', '-id')
        )
        grouped = {}
        for row in qs:
            level_key = int(row.entity_id or 0)
            item = grouped.get(level_key)
            if item is None:
                level_label = 'ESCALATION' if level_key == 99 else ('L2' if level_key == 2 else 'L1')
                item = {
                    'level_key': level_key,
                    'level_label': level_label,
                    'latest_created_at': row.created_at,
                    'sent_count': 0,
                    'unread_count': 0,
                    'sample_recipients': [],
                }
                grouped[level_key] = item
            item['sent_count'] += 1
            if not row.is_read:
                item['unread_count'] += 1
            if len(item['sample_recipients']) < 8:
                item['sample_recipients'].append(str(row.recipient.username))
            if row.created_at and row.created_at > item['latest_created_at']:
                item['latest_created_at'] = row.created_at

        items = []
        total_sent = 0
        total_unread = 0
        for item in sorted(grouped.values(), key=lambda x: (-int(x['sent_count']), int(x['level_key']))):
            sent_count = int(item['sent_count'])
            unread_count = int(item['unread_count'])
            read_count = max(0, sent_count - unread_count)
            item['read_count'] = read_count
            item['read_rate'] = round((read_count / sent_count) * 100, 2) if sent_count > 0 else 0.0
            item['latest_created_at'] = item['latest_created_at'].isoformat() if item.get('latest_created_at') else ''
            total_sent += sent_count
            total_unread += unread_count
            items.append(item)
        total_read = max(0, total_sent - total_unread)
        overall_read_rate = round((total_read / total_sent) * 100, 2) if total_sent > 0 else 0.0
        return Response({
            'days': days,
            'summary': {
                'total_sent': total_sent,
                'total_read': total_read,
                'total_unread': total_unread,
                'overall_read_rate': overall_read_rate,
            },
            'items': items,
        })

    @action(detail=False, methods=['get'])
    def approval_sla_policy(self, request):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền xem policy SLA duyệt.'}, status=403)
        return Response(get_approval_sla_policy())

    @approval_sla_policy.mapping.post
    def save_approval_sla_policy_action(self, request):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền cập nhật policy SLA duyệt.'}, status=403)
        payload = request.data if isinstance(request.data, dict) else {}
        old_policy = get_approval_sla_policy()
        policy = save_approval_sla_policy(payload)
        _log_finance_audit(
            request.user,
            action='UPDATE',
            entity_type='FinanceAdvanceApprovalSlaPolicy',
            entity_id=0,
            entity_code='FINANCE_ADVANCE_APPROVAL_SLA_POLICY',
            old_values=old_policy,
            new_values=policy,
            changed_fields=['sla_hours_l1', 'sla_hours_l2', 'remind_every_hours', 'window_days'],
        )
        return Response({'success': True, 'policy': policy})

    @action(detail=True, methods=['post'])
    def submit_approval(self, request, pk=None):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền gửi duyệt phiếu tạm ứng.'}, status=403)
        with transaction.atomic():
            try:
                advance = AdvanceTransaction.objects.select_for_update().get(pk=pk)
            except AdvanceTransaction.DoesNotExist:
                return Response({'error': 'Không tìm thấy phiếu tạm ứng.'}, status=404)
            if advance.approval_status in {AdvanceTransaction.APPROVAL_PENDING_L1, AdvanceTransaction.APPROVAL_PENDING_L2, AdvanceTransaction.APPROVAL_APPROVED}:
                return Response({'error': 'Phiếu đang chờ duyệt hoặc đã duyệt.'}, status=400)
            _ensure_finance_month_unlocked(_month_from_date_obj(advance.advance_date), 'Tháng tài chính đã khóa, không thể gửi duyệt.')
            old_status = advance.approval_status
            advance.approval_status = AdvanceTransaction.APPROVAL_PENDING_L1
            advance.submitted_at = timezone.now()
            advance.submitted_by = request.user
            advance.rejected_at = None
            advance.rejected_by = None
            advance.rejection_reason = ''
            advance.required_approval_level = _required_approval_level_for_amount(Decimal(str(advance.amount or 0)))
            advance.updated_by = request.user
            advance.save(update_fields=[
                'approval_status',
                'submitted_at',
                'submitted_by',
                'rejected_at',
                'rejected_by',
                'rejection_reason',
                'required_approval_level',
                'updated_by',
                'updated_at',
                'search_text',
            ])
            _log_finance_audit(
                request.user,
                action='UPDATE',
                entity_type='FinanceAdvanceApproval',
                entity_id=int(advance.id),
                entity_code=advance.code,
                old_values={'approval_status': old_status},
                new_values={'approval_status': advance.approval_status, 'required_approval_level': advance.required_approval_level},
                changed_fields=['approval_status', 'required_approval_level'],
            )
        return Response({'success': True, 'approval_status': advance.approval_status, 'required_approval_level': advance.required_approval_level})

    @action(detail=True, methods=['post'])
    def approve_level1(self, request, pk=None):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền duyệt cấp 1.'}, status=403)
        with transaction.atomic():
            try:
                advance = AdvanceTransaction.objects.select_for_update().get(pk=pk)
            except AdvanceTransaction.DoesNotExist:
                return Response({'error': 'Không tìm thấy phiếu tạm ứng.'}, status=404)
            if advance.submitted_by_id and advance.submitted_by_id == request.user.id:
                return Response({'error': 'Người gửi duyệt không thể tự phê duyệt phiếu của mình.'}, status=403)
            if advance.approval_status != AdvanceTransaction.APPROVAL_PENDING_L1:
                return Response({'error': 'Phiếu không ở trạng thái chờ duyệt cấp 1.'}, status=400)
            _ensure_finance_month_unlocked(_month_from_date_obj(advance.advance_date), 'Tháng tài chính đã khóa, không thể duyệt.')
            old_status = advance.approval_status
            next_status = AdvanceTransaction.APPROVAL_APPROVED if int(advance.required_approval_level or 1) <= 1 else AdvanceTransaction.APPROVAL_PENDING_L2
            advance.approval_status = next_status
            advance.approved_level1_at = timezone.now()
            advance.approved_level1_by = request.user
            advance.updated_by = request.user
            advance.save(update_fields=['approval_status', 'approved_level1_at', 'approved_level1_by', 'updated_by', 'updated_at', 'search_text'])
            _log_finance_audit(
                request.user,
                action='UPDATE',
                entity_type='FinanceAdvanceApproval',
                entity_id=int(advance.id),
                entity_code=advance.code,
                old_values={'approval_status': old_status},
                new_values={'approval_status': advance.approval_status},
                changed_fields=['approval_status', 'approved_level1_at'],
            )
        return Response({'success': True, 'approval_status': advance.approval_status})

    @action(detail=True, methods=['post'])
    def approve_level2(self, request, pk=None):
        if not _can_manage_finance(request.user) or not _can_approve_level2(request.user):
            return Response({'error': 'Bạn không có quyền duyệt cấp 2.'}, status=403)
        with transaction.atomic():
            try:
                advance = AdvanceTransaction.objects.select_for_update().get(pk=pk)
            except AdvanceTransaction.DoesNotExist:
                return Response({'error': 'Không tìm thấy phiếu tạm ứng.'}, status=404)
            if advance.submitted_by_id and advance.submitted_by_id == request.user.id:
                return Response({'error': 'Người gửi duyệt không thể tự phê duyệt phiếu của mình.'}, status=403)
            if advance.approved_level1_by_id and advance.approved_level1_by_id == request.user.id and int(advance.required_approval_level or 1) >= 2:
                return Response({'error': 'Người đã duyệt cấp 1 không thể đồng thời duyệt cấp 2 cho cùng phiếu.'}, status=403)
            if advance.approval_status != AdvanceTransaction.APPROVAL_PENDING_L2:
                return Response({'error': 'Phiếu không ở trạng thái chờ duyệt cấp 2.'}, status=400)
            _ensure_finance_month_unlocked(_month_from_date_obj(advance.advance_date), 'Tháng tài chính đã khóa, không thể duyệt.')
            old_status = advance.approval_status
            advance.approval_status = AdvanceTransaction.APPROVAL_APPROVED
            advance.approved_level2_at = timezone.now()
            advance.approved_level2_by = request.user
            advance.updated_by = request.user
            advance.save(update_fields=['approval_status', 'approved_level2_at', 'approved_level2_by', 'updated_by', 'updated_at', 'search_text'])
            _log_finance_audit(
                request.user,
                action='UPDATE',
                entity_type='FinanceAdvanceApproval',
                entity_id=int(advance.id),
                entity_code=advance.code,
                old_values={'approval_status': old_status},
                new_values={'approval_status': advance.approval_status},
                changed_fields=['approval_status', 'approved_level2_at'],
            )
        return Response({'success': True, 'approval_status': advance.approval_status})

    @action(detail=True, methods=['post'])
    def reject_approval(self, request, pk=None):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền từ chối duyệt.'}, status=403)
        reason = str(request.data.get('reason') or '').strip()
        if not reason:
            return Response({'error': 'reason là bắt buộc khi từ chối.'}, status=400)
        with transaction.atomic():
            try:
                advance = AdvanceTransaction.objects.select_for_update().get(pk=pk)
            except AdvanceTransaction.DoesNotExist:
                return Response({'error': 'Không tìm thấy phiếu tạm ứng.'}, status=404)
            if advance.approval_status not in {AdvanceTransaction.APPROVAL_PENDING_L1, AdvanceTransaction.APPROVAL_PENDING_L2}:
                return Response({'error': 'Phiếu không ở trạng thái chờ duyệt.'}, status=400)
            old_status = advance.approval_status
            advance.approval_status = AdvanceTransaction.APPROVAL_REJECTED
            advance.rejected_at = timezone.now()
            advance.rejected_by = request.user
            advance.rejection_reason = reason[:255]
            advance.updated_by = request.user
            advance.save(update_fields=['approval_status', 'rejected_at', 'rejected_by', 'rejection_reason', 'updated_by', 'updated_at', 'search_text'])
            _log_finance_audit(
                request.user,
                action='UPDATE',
                entity_type='FinanceAdvanceApproval',
                entity_id=int(advance.id),
                entity_code=advance.code,
                old_values={'approval_status': old_status},
                new_values={'approval_status': advance.approval_status, 'rejection_reason': advance.rejection_reason},
                changed_fields=['approval_status', 'rejection_reason'],
            )
        return Response({'success': True, 'approval_status': advance.approval_status})

    @action(detail=True, methods=['post'])
    def post_disbursement(self, request, pk=None):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền chi tiền tạm ứng.'}, status=403)
        with transaction.atomic():
            try:
                advance = AdvanceTransaction.objects.select_for_update().get(pk=pk)
            except AdvanceTransaction.DoesNotExist:
                return Response({'error': 'Không tìm thấy phiếu tạm ứng.'}, status=404)
            if advance.approval_status != AdvanceTransaction.APPROVAL_APPROVED:
                return Response({'error': 'Chỉ được chi tiền cho phiếu tạm ứng đã duyệt.'}, status=400)
            _ensure_finance_month_unlocked(_month_from_date_obj(advance.advance_date), 'Tháng tài chính đã khóa, không thể chi tiền tạm ứng.')
            existing = _get_advance_disbursement_tx(advance, lock_for_update=True)
            if existing:
                return Response({'error': 'Phiếu này đã có chứng từ chi tiền.'}, status=400)
            if advance.source_type == AdvanceTransaction.SOURCE_CASH:
                if not advance.source_cash_account_id or not getattr(advance.source_cash_account, 'is_active', False):
                    return Response({'error': 'Tài khoản quỹ nguồn không còn hiệu lực để chi tiền.'}, status=400)
                try:
                    _ensure_cash_account_has_available_balance(
                        advance.source_cash_account,
                        advance.amount,
                        up_to_date=advance.advance_date,
                        message_prefix='Quỹ chi tạm ứng',
                    )
                except ValidationError as exc:
                    return Response({'error': exc.detail.get('amount', exc.detail)}, status=400)
            else:
                if not advance.source_bank_account_id or not getattr(advance.source_bank_account, 'is_active', False):
                    return Response({'error': 'Tài khoản ngân hàng nguồn không còn hiệu lực để chi tiền.'}, status=400)
            category = _ensure_advance_disbursement_category(request.user)
            marker = _advance_disbursement_marker(advance)
            tx = CashTransaction.objects.create(
                transaction_type=CashTransaction.TYPE_EXPENSE,
                source_type=advance.source_type,
                source_cash_account=advance.source_cash_account if advance.source_type == AdvanceTransaction.SOURCE_CASH else None,
                source_bank_account=advance.source_bank_account if advance.source_type == AdvanceTransaction.SOURCE_BANK else None,
                category=category,
                transaction_date=advance.advance_date,
                amount=advance.amount,
                object_name=advance.recipient_name,
                reason=f'{marker} Chi tiền tạm ứng {advance.code}',
                note=f'Tự động sinh từ phiếu tạm ứng {advance.code}',
                created_by=request.user,
            )
            _log_finance_audit(
                request.user,
                action='CREATE',
                entity_type='FinanceAdvanceDisbursement',
                entity_id=int(advance.id),
                entity_code=advance.code,
                old_values={},
                new_values={'cash_transaction_id': int(tx.id), 'amount': str(advance.amount)},
                changed_fields=['cash_transaction_id', 'amount'],
            )
        return Response({'success': True, 'transaction_id': int(tx.id), 'disbursement_status': 'DISBURSED'})

    @action(detail=True, methods=['post'])
    def reverse_disbursement(self, request, pk=None):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền hủy chi tiền tạm ứng.'}, status=403)
        with transaction.atomic():
            try:
                advance = AdvanceTransaction.objects.select_for_update().get(pk=pk)
            except AdvanceTransaction.DoesNotExist:
                return Response({'error': 'Không tìm thấy phiếu tạm ứng.'}, status=404)
            _ensure_finance_month_unlocked(_month_from_date_obj(advance.advance_date), 'Tháng tài chính đã khóa, không thể hủy chi tiền tạm ứng.')
            if advance.settlements.exists():
                return Response({'error': 'Phiếu đã có quyết toán, không thể hủy chứng từ chi tiền.'}, status=400)
            tx = _get_advance_disbursement_tx(advance, lock_for_update=True)
            if tx is None:
                return Response({'error': 'Phiếu này chưa có chứng từ chi tiền.'}, status=400)
            _ensure_finance_month_unlocked(_month_from_date_obj(tx.transaction_date), 'Tháng tài chính của chứng từ chi đã khóa, không thể hủy.')
            tx_id = int(tx.id)
            tx.delete()
            _log_finance_audit(
                request.user,
                action='DELETE',
                entity_type='FinanceAdvanceDisbursement',
                entity_id=int(advance.id),
                entity_code=advance.code,
                old_values={'cash_transaction_id': tx_id},
                new_values={},
                changed_fields=['cash_transaction_id'],
            )
        return Response({'success': True, 'transaction_id': tx_id, 'disbursement_status': 'NOT_DISBURSED'})

    @action(detail=False, methods=['get'])
    def overdue_report(self, request):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền xem báo cáo tạm ứng quá hạn.'}, status=403)
        as_of_raw = (request.query_params.get('as_of') or '').strip()
        overdue_days_raw = (request.query_params.get('overdue_days') or '30').strip()
        try:
            overdue_days = int(overdue_days_raw)
        except ValueError:
            return Response({'error': 'overdue_days phải là số nguyên.'}, status=400)
        overdue_days = max(1, min(3650, overdue_days))

        if as_of_raw:
            try:
                y, m, d = as_of_raw.split('-', 2)
                as_of = date(int(y), int(m), int(d))
            except Exception:
                return Response({'error': 'as_of phải có dạng YYYY-MM-DD'}, status=400)
        else:
            as_of = date.today()
        cutoff_date = as_of.fromordinal(as_of.toordinal() - overdue_days)

        qs = _approved_advance_queryset().filter(
            status__in=[AdvanceTransaction.STATUS_OPEN, AdvanceTransaction.STATUS_PARTIAL],
            advance_date__lte=cutoff_date,
        ).select_related('source_cash_account', 'source_bank_account')

        items = []
        total_amount = Decimal('0')
        total_handled = Decimal('0')
        total_remaining = Decimal('0')
        for adv in qs:
            spent = Decimal(str(adv.settlements.aggregate(total=Sum('spent_amount')).get('total') or 0))
            refund = Decimal(str(adv.settlements.aggregate(total=Sum('refund_amount')).get('total') or 0))
            handled = spent + refund
            amount = Decimal(str(adv.amount or 0))
            remaining = amount - handled
            if remaining <= 0:
                continue
            days_overdue = (as_of - adv.advance_date).days
            total_amount += amount
            total_handled += handled
            total_remaining += remaining
            items.append({
                'id': adv.id,
                'code': adv.code,
                'advance_type': adv.advance_type,
                'advance_date': adv.advance_date.isoformat() if adv.advance_date else '',
                'recipient_name': adv.recipient_name,
                'status': adv.status,
                'amount': str(amount),
                'spent_amount': str(spent),
                'refund_amount': str(refund),
                'handled_amount': str(handled),
                'remaining_amount': str(remaining),
                'days_overdue': days_overdue,
                'source_type': adv.source_type,
                'source_cash_account_name': adv.source_cash_account.name if adv.source_cash_account_id else '',
                'source_bank_account_code': adv.source_bank_account.code if adv.source_bank_account_id else '',
                'purpose': adv.purpose,
            })

        items.sort(key=lambda x: (-int(x['days_overdue']), x['advance_date'], x['code']))
        summary = {
            'as_of': as_of.isoformat(),
            'overdue_days': overdue_days,
            'count': len(items),
            'total_amount': str(total_amount),
            'total_handled': str(total_handled),
            'total_remaining': str(total_remaining),
        }

        if (request.query_params.get('export') or '').strip().lower() == 'excel':
            wb = Workbook()
            ws = wb.active
            ws.title = 'TamUngQuaHan'
            ws.append([
                'Ma tam ung',
                'Ngay tam ung',
                'Nguoi nhan',
                'Loai',
                'Trang thai',
                'So tien',
                'Da xu ly',
                'Con lai',
                'So ngay qua han',
                'Muc dich',
            ])
            for row in items:
                ws.append([
                    row['code'],
                    row['advance_date'],
                    row['recipient_name'],
                    row['advance_type'],
                    row['status'],
                    float(row['amount']),
                    float(row['handled_amount']),
                    float(row['remaining_amount']),
                    int(row['days_overdue']),
                    row['purpose'],
                ])
            response = HttpResponse(
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )
            response['Content-Disposition'] = f'attachment; filename=\"bao_cao_tam_ung_qua_han_{as_of.isoformat()}.xlsx\"'
            wb.save(response)
            return response

        return Response({'summary': summary, 'items': items})

    @action(detail=False, methods=['get'])
    def overdue_overview(self, request):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền xem tổng quan tạm ứng quá hạn.'}, status=403)
        as_of_raw = (request.query_params.get('as_of') or '').strip()
        if as_of_raw:
            try:
                y, m, d = as_of_raw.split('-', 2)
                as_of = date(int(y), int(m), int(d))
            except Exception:
                return Response({'error': 'as_of phải có dạng YYYY-MM-DD'}, status=400)
        else:
            as_of = date.today()

        base_qs = _approved_advance_queryset().filter(
            status__in=[AdvanceTransaction.STATUS_OPEN, AdvanceTransaction.STATUS_PARTIAL],
            advance_date__lt=as_of,
        ).select_related('source_cash_account', 'source_bank_account')

        entries: list[dict] = []
        for adv in base_qs:
            spent = Decimal(str(adv.settlements.aggregate(total=Sum('spent_amount')).get('total') or 0))
            refund = Decimal(str(adv.settlements.aggregate(total=Sum('refund_amount')).get('total') or 0))
            handled = spent + refund
            amount = Decimal(str(adv.amount or 0))
            remaining = amount - handled
            if remaining <= 0:
                continue
            days_overdue = (as_of - adv.advance_date).days
            entries.append({
                'id': adv.id,
                'code': adv.code,
                'advance_date': adv.advance_date.isoformat() if adv.advance_date else '',
                'recipient_name': adv.recipient_name,
                'status': adv.status,
                'amount': str(amount),
                'remaining_amount': str(remaining),
                'days_overdue': days_overdue,
            })

        def bucket_stats(threshold: int) -> dict:
            bucket_items = [item for item in entries if int(item['days_overdue']) >= threshold]
            total_remaining = Decimal('0')
            for item in bucket_items:
                total_remaining += Decimal(str(item['remaining_amount']))
            return {
                'threshold_days': threshold,
                'count': len(bucket_items),
                'total_remaining': str(total_remaining),
            }

        buckets = [
            bucket_stats(30),
            bucket_stats(60),
            bucket_stats(90),
        ]
        top_urgent = sorted(entries, key=lambda x: (-int(x['days_overdue']), x['advance_date'], x['code']))[:10]
        return Response({
            'as_of': as_of.isoformat(),
            'buckets': buckets,
            'top_urgent': top_urgent,
        })

    @action(detail=False, methods=['post'])
    def remind_overdue(self, request):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền gửi nhắc quá hạn.'}, status=403)
        threshold_days_raw = request.data.get('threshold_days')
        threshold_days = None if threshold_days_raw in (None, '') else threshold_days_raw
        as_of = (request.data.get('as_of') or '').strip() or None
        recipient_usernames_raw = request.data.get('recipient_usernames')
        recipient_usernames = None
        if isinstance(recipient_usernames_raw, list):
            recipient_usernames = [str(item).strip() for item in recipient_usernames_raw if str(item).strip()]
        elif isinstance(recipient_usernames_raw, str):
            recipient_usernames = [part.strip() for part in recipient_usernames_raw.split(',') if part.strip()]
        try:
            result = run_daily_overdue_reminder_job(
                threshold_days=None if threshold_days is None else int(threshold_days),
                as_of=as_of,
                recipient_usernames=recipient_usernames,
            )
        except ValueError:
            return Response({'error': 'threshold_days phải là số nguyên.'}, status=400)
        return Response(result)

    @action(detail=False, methods=['get'])
    def reminder_policy(self, request):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền xem policy nhắc quá hạn.'}, status=403)
        policy = get_reminder_policy()
        policy['presets'] = get_reminder_policy_presets()
        policy['recommendation'] = _recommend_reminder_policy_preset()
        return Response(policy)

    @reminder_policy.mapping.post
    def save_reminder_policy_action(self, request):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền cập nhật policy nhắc quá hạn.'}, status=403)
        payload = request.data if isinstance(request.data, dict) else {}
        old_policy = get_reminder_policy()
        try:
            policy = save_reminder_policy(payload)
        except Exception:
            return Response({'error': 'Dữ liệu policy không hợp lệ.'}, status=400)
        setting_row = Setting.objects.filter(key=REMINDER_POLICY_SETTING_KEY).first()
        AuditLog.objects.create(
            user=request.user,
            action='UPDATE',
            entity_type='FinanceAdvanceReminderPolicy',
            entity_id=int(setting_row.id if setting_row else 0),
            entity_id_str=str(setting_row.id if setting_row else ''),
            entity_code=REMINDER_POLICY_SETTING_KEY,
            old_values=old_policy,
            new_values=policy,
            changed_fields=['default_threshold_days', 'cooldown_hours', 'role_threshold_days', 'user_threshold_days'],
        )
        return Response({'success': True, 'policy': policy})

    @action(detail=False, methods=['get'])
    def reminder_policy_history(self, request):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền xem lịch sử policy nhắc quá hạn.'}, status=403)
        limit_raw = (request.query_params.get('limit') or '20').strip()
        try:
            limit = int(limit_raw)
        except ValueError:
            return Response({'error': 'limit phải là số nguyên.'}, status=400)
        limit = max(1, min(limit, 100))
        logs = (
            AuditLog.objects
            .filter(entity_type='FinanceAdvanceReminderPolicy', action='UPDATE')
            .select_related('user')
            .order_by('-created_at', '-id')[:limit]
        )
        items = []
        for row in logs:
            old_values = row.old_values if isinstance(row.old_values, dict) else {}
            new_values = row.new_values if isinstance(row.new_values, dict) else {}
            items.append({
                'id': int(row.id),
                'created_at': row.created_at.isoformat() if row.created_at else '',
                'username': str(row.user.username) if row.user else '',
                'old_values': old_values,
                'new_values': new_values,
                'changed_fields': row.changed_fields if isinstance(row.changed_fields, list) else [],
            })
        return Response({'count': len(items), 'items': items})

    @action(detail=False, methods=['post'])
    def reminder_policy_rollback(self, request):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền rollback policy nhắc quá hạn.'}, status=403)
        audit_log_id_raw = request.data.get('audit_log_id')
        try:
            audit_log_id = int(audit_log_id_raw)
        except (TypeError, ValueError):
            return Response({'error': 'audit_log_id phải là số nguyên.'}, status=400)
        target_log = (
            AuditLog.objects
            .filter(id=audit_log_id, entity_type='FinanceAdvanceReminderPolicy', action='UPDATE')
            .first()
        )
        if not target_log:
            return Response({'error': 'Không tìm thấy bản ghi lịch sử policy.'}, status=404)
        target_policy = target_log.old_values if isinstance(target_log.old_values, dict) else {}
        if not target_policy:
            return Response({'error': 'Bản ghi lịch sử không chứa policy cũ hợp lệ.'}, status=400)
        old_policy = get_reminder_policy()
        try:
            policy = save_reminder_policy(target_policy)
        except Exception:
            return Response({'error': 'Rollback policy thất bại do dữ liệu không hợp lệ.'}, status=400)
        setting_row = Setting.objects.filter(key=REMINDER_POLICY_SETTING_KEY).first()
        AuditLog.objects.create(
            user=request.user,
            action='UPDATE',
            entity_type='FinanceAdvanceReminderPolicy',
            entity_id=int(setting_row.id if setting_row else 0),
            entity_id_str=str(setting_row.id if setting_row else ''),
            entity_code=REMINDER_POLICY_SETTING_KEY,
            old_values=old_policy,
            new_values=policy,
            changed_fields=['default_threshold_days', 'cooldown_hours', 'role_threshold_days', 'user_threshold_days'],
        )
        return Response({
            'success': True,
            'policy': policy,
            'rolled_back_to_audit_log_id': audit_log_id,
        })

    @action(detail=False, methods=['post'])
    def reminder_policy_simulate(self, request):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền mô phỏng policy nhắc quá hạn.'}, status=403)
        payload = request.data if isinstance(request.data, dict) else {}
        threshold_days_raw = payload.get('threshold_days')
        threshold_days = None if threshold_days_raw in (None, '') else threshold_days_raw
        as_of = (payload.get('as_of') or '').strip() or None
        recipient_usernames_raw = payload.get('recipient_usernames')
        policy_override = payload.get('policy')
        if policy_override is not None and not isinstance(policy_override, dict):
            return Response({'error': 'policy phải là object JSON.'}, status=400)

        recipient_usernames = None
        if isinstance(recipient_usernames_raw, list):
            recipient_usernames = [str(item).strip() for item in recipient_usernames_raw if str(item).strip()]
        elif isinstance(recipient_usernames_raw, str):
            recipient_usernames = [part.strip() for part in recipient_usernames_raw.split(',') if part.strip()]
        try:
            result = run_daily_overdue_reminder_job(
                threshold_days=None if threshold_days is None else int(threshold_days),
                as_of=as_of,
                recipient_usernames=recipient_usernames,
                policy_override=policy_override,
                dry_run=True,
            )
        except ValueError:
            return Response({'error': 'threshold_days phải là số nguyên.'}, status=400)
        return Response(result)

    @action(detail=False, methods=['get'])
    def reminder_history(self, request):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền xem lịch sử nhắc quá hạn.'}, status=403)
        days_raw = (request.query_params.get('days') or '30').strip()
        try:
            days = int(days_raw)
        except ValueError:
            return Response({'error': 'days phải là số nguyên.'}, status=400)
        days = max(1, min(days, 365))
        cutoff = date.today().fromordinal(date.today().toordinal() - days)

        as_of_from = (request.query_params.get('as_of_from') or '').strip()
        as_of_to = (request.query_params.get('as_of_to') or '').strip()
        entity_from = None
        entity_to = None
        if as_of_from:
            try:
                y, m, d = as_of_from.split('-', 2)
                entity_from = int(f'{int(y):04d}{int(m):02d}{int(d):02d}')
            except Exception:
                return Response({'error': 'as_of_from phải có dạng YYYY-MM-DD'}, status=400)
        if as_of_to:
            try:
                y, m, d = as_of_to.split('-', 2)
                entity_to = int(f'{int(y):04d}{int(m):02d}{int(d):02d}')
            except Exception:
                return Response({'error': 'as_of_to phải có dạng YYYY-MM-DD'}, status=400)
        unread_only = str(request.query_params.get('unread_only') or '').strip().lower() in {'1', 'true', 'yes'}

        qs = (
            Notification.objects
            .filter(
                notification_type='due_date',
                entity_type='FinanceAdvanceOverdueDaily',
                created_at__date__gte=cutoff,
            )
            .select_related('recipient')
            .order_by('-created_at', '-id')
        )
        if entity_from is not None:
            qs = qs.filter(entity_id__gte=entity_from)
        if entity_to is not None:
            qs = qs.filter(entity_id__lte=entity_to)

        grouped = {}
        recipient_stats = {}
        for n in qs:
            key = int(n.entity_id or 0)
            entry = grouped.get(key)
            if entry is None:
                entry = {
                    'entity_id': key,
                    'title': n.title,
                    'message': n.message,
                    'created_at': n.created_at,
                    'sent_count': 0,
                    'unread_count': 0,
                    'recipient_names': [],
                }
                grouped[key] = entry
            entry['sent_count'] += 1
            if not n.is_read:
                entry['unread_count'] += 1
            if len(entry['recipient_names']) < 5:
                entry['recipient_names'].append(n.recipient.username)
            if n.created_at and n.created_at > entry['created_at']:
                entry['created_at'] = n.created_at

            recipient_key = n.recipient.username
            stat = recipient_stats.get(recipient_key)
            if stat is None:
                stat = {
                    'username': recipient_key,
                    'total_received': 0,
                    'unread_count': 0,
                }
                recipient_stats[recipient_key] = stat
            stat['total_received'] += 1
            if not n.is_read:
                stat['unread_count'] += 1

        items = sorted(grouped.values(), key=lambda x: x['created_at'], reverse=True)
        total_sent = 0
        total_unread = 0
        for item in items:
            item['created_at'] = item['created_at'].isoformat() if item['created_at'] else ''
            if item['entity_id'] > 0:
                entity_str = str(item['entity_id']).zfill(8)
                item['as_of'] = f'{entity_str[0:4]}-{entity_str[4:6]}-{entity_str[6:8]}'
            else:
                item['as_of'] = ''
            sent_count = int(item.get('sent_count') or 0)
            unread_count = int(item.get('unread_count') or 0)
            read_count = max(0, sent_count - unread_count)
            read_rate = round((read_count / sent_count) * 100, 2) if sent_count > 0 else 0
            unread_rate = round((unread_count / sent_count) * 100, 2) if sent_count > 0 else 0
            item['read_count'] = read_count
            item['read_rate'] = read_rate
            item['unread_rate'] = unread_rate
            total_sent += sent_count
            total_unread += unread_count
        if unread_only:
            items = [item for item in items if int(item['unread_count']) > 0]

        if (request.query_params.get('export') or '').strip().lower() == 'excel':
            wb = Workbook()
            ws = wb.active
            ws.title = 'LichSuNhacQuaHan'
            ws.append([
                'Ngay nhac (as_of)',
                'Tieu de',
                'Da gui',
                'Da doc',
                'Chua doc',
                'Ty le da doc (%)',
                'Ty le chua doc (%)',
                'Nguoi nhan mau',
                'Thoi diem tao ban ghi',
            ])
            for item in items[:200]:
                ws.append([
                    item.get('as_of', ''),
                    item.get('title', ''),
                    int(item.get('sent_count') or 0),
                    int(item.get('read_count') or 0),
                    int(item.get('unread_count') or 0),
                    float(item.get('read_rate') or 0),
                    float(item.get('unread_rate') or 0),
                    ', '.join(item.get('recipient_names') or []),
                    item.get('created_at', ''),
                ])
            response = HttpResponse(
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )
            response['Content-Disposition'] = 'attachment; filename="lich_su_nhac_qua_han.xlsx"'
            wb.save(response)
            return response

        total_read = max(0, total_sent - total_unread)
        overall_read_rate = round((total_read / total_sent) * 100, 2) if total_sent > 0 else 0
        top_unread_recipients = sorted(
            recipient_stats.values(),
            key=lambda x: (-int(x.get('unread_count') or 0), -int(x.get('total_received') or 0), str(x.get('username') or '')),
        )
        if unread_only:
            top_unread_recipients = [item for item in top_unread_recipients if int(item.get('unread_count') or 0) > 0]
        for item in top_unread_recipients:
            total_received = int(item.get('total_received') or 0)
            unread_count = int(item.get('unread_count') or 0)
            item['read_count'] = max(0, total_received - unread_count)
            item['unread_rate'] = round((unread_count / total_received) * 100, 2) if total_received > 0 else 0
        return Response({
            'days': days,
            'count': len(items),
            'filters': {
                'as_of_from': as_of_from,
                'as_of_to': as_of_to,
                'unread_only': unread_only,
            },
            'summary': {
                'total_sent': total_sent,
                'total_read': total_read,
                'total_unread': total_unread,
                'overall_read_rate': overall_read_rate,
                'top_unread_recipients': top_unread_recipients[:10],
            },
            'items': items[:50],
        })


class AdvanceSettlementViewSet(SearchTextMixin, viewsets.ModelViewSet):
    queryset = AdvanceSettlement.objects.select_related('advance_transaction')
    serializer_class = AdvanceSettlementSerializer
    permission_classes = [IsAuthenticated]
    filter_backends = [filters.OrderingFilter]
    ordering_fields = ['settlement_date', 'spent_amount', 'refund_amount', 'created_at']
    ordering = ['-settlement_date', '-id']

    def get_queryset(self):
        queryset = super().get_queryset()
        params = self.request.query_params

        advance_id = (params.get('advance_transaction') or '').strip()
        if advance_id.isdigit():
            queryset = queryset.filter(advance_transaction_id=int(advance_id))

        month = (params.get('month') or '').strip()
        if len(month) == 7 and '-' in month:
            year_str, month_str = month.split('-', 1)
            if year_str.isdigit() and month_str.isdigit():
                queryset = queryset.filter(
                    settlement_date__year=int(year_str),
                    settlement_date__month=int(month_str),
                )
        return self.apply_search(queryset)

    def perform_create(self, serializer):
        if not _can_manage_finance(self.request.user):
            raise PermissionDenied('Bạn không có quyền tạo quyết toán.')
        settlement_date = serializer.validated_data.get('settlement_date')
        _ensure_finance_month_unlocked(
            _month_from_date_obj(settlement_date),
            'Tháng tài chính đã khóa, không thể thêm quyết toán.',
        )
        advance = serializer.validated_data.get('advance_transaction')
        if advance and advance.approval_status != AdvanceTransaction.APPROVAL_APPROVED:
            raise PermissionDenied('Phiếu tạm ứng chưa được duyệt đầy đủ, không thể quyết toán.')
        with transaction.atomic():
            locked_advance = (
                AdvanceTransaction.objects
                .select_for_update()
                .get(pk=advance.pk)
            )
            validation_error = _validate_settlement_total_under_advance(
                locked_advance,
                serializer.validated_data.get('spent_amount'),
                serializer.validated_data.get('refund_amount'),
            )
            if validation_error:
                raise ValidationError({'non_field_errors': [validation_error]})
            settlement = serializer.save(
                created_by=self.request.user,
                updated_by=self.request.user,
                advance_transaction=locked_advance,
            )
            refresh_advance_status(locked_advance, self.request.user)
        _log_finance_audit(
            self.request.user, action='CREATE', entity_type='FinanceAdvanceSettlement',
            entity_id=int(settlement.id), entity_code=str(settlement.id),
            old_values={},
            new_values={
                'advance_transaction_id': int(settlement.advance_transaction_id),
                'spent_amount': str(settlement.spent_amount),
                'refund_amount': str(settlement.refund_amount),
                'settlement_date': settlement.settlement_date.isoformat() if settlement.settlement_date else '',
            },
            changed_fields=['spent_amount', 'refund_amount', 'settlement_date'],
        )

    def perform_update(self, serializer):
        if not _can_manage_finance(self.request.user):
            raise PermissionDenied('Bạn không có quyền cập nhật quyết toán.')
        old_month = _month_from_date_obj(serializer.instance.settlement_date)
        settlement_date = serializer.validated_data.get('settlement_date', serializer.instance.settlement_date)
        _ensure_finance_months_unlocked(
            [old_month, _month_from_date_obj(settlement_date)],
            'Tháng tài chính đã khóa, không thể cập nhật quyết toán.',
        )
        advance = serializer.validated_data.get('advance_transaction', serializer.instance.advance_transaction)
        if advance and advance.approval_status != AdvanceTransaction.APPROVAL_APPROVED:
            raise PermissionDenied('Phiếu tạm ứng chưa được duyệt đầy đủ, không thể cập nhật quyết toán.')
        old = serializer.instance
        old_parent_id = old.advance_transaction_id
        old_snap = {
            'advance_transaction_id': int(old.advance_transaction_id),
            'spent_amount': str(old.spent_amount),
            'refund_amount': str(old.refund_amount),
            'settlement_date': old.settlement_date.isoformat() if old.settlement_date else '',
        }
        with transaction.atomic():
            lock_ids = {int(old_parent_id)}
            if advance and advance.pk:
                lock_ids.add(int(advance.pk))
            locked_advances = {
                row.id: row
                for row in AdvanceTransaction.objects.select_for_update().filter(id__in=lock_ids)
            }
            target_advance = locked_advances.get(int(advance.pk if advance else old_parent_id))
            validation_error = _validate_settlement_total_under_advance(
                target_advance,
                serializer.validated_data.get('spent_amount', old.spent_amount),
                serializer.validated_data.get('refund_amount', old.refund_amount),
                exclude_settlement_id=old.id,
            )
            if validation_error:
                raise ValidationError({'non_field_errors': [validation_error]})
            settlement = serializer.save(
                updated_by=self.request.user,
                advance_transaction=target_advance,
            )
            refresh_advance_status(target_advance, self.request.user)
            if int(old_parent_id) != int(target_advance.id):
                refresh_advance_status(locked_advances[int(old_parent_id)], self.request.user)
        _log_finance_audit(
            self.request.user, action='UPDATE', entity_type='FinanceAdvanceSettlement',
            entity_id=int(settlement.id), entity_code=str(settlement.id),
            old_values=old_snap,
            new_values={
                'advance_transaction_id': int(settlement.advance_transaction_id),
                'spent_amount': str(settlement.spent_amount),
                'refund_amount': str(settlement.refund_amount),
                'settlement_date': settlement.settlement_date.isoformat() if settlement.settlement_date else '',
            },
            changed_fields=list(serializer.validated_data.keys()),
        )

    def perform_destroy(self, instance):
        _ensure_finance_month_unlocked(
            _month_from_date_obj(instance.settlement_date),
            'Tháng tài chính đã khóa, không thể xóa quyết toán.',
        )
        advance = instance.advance_transaction
        if advance.approval_status != AdvanceTransaction.APPROVAL_APPROVED:
            raise PermissionDenied('Phiếu tạm ứng chưa được duyệt đầy đủ, không thể xóa quyết toán.')
        snap = {'spent_amount': str(instance.spent_amount), 'refund_amount': str(instance.refund_amount), 'advance_transaction_id': int(instance.advance_transaction_id)}
        with transaction.atomic():
            locked_advance = AdvanceTransaction.objects.select_for_update().get(pk=advance.pk)
            instance.delete()
            refresh_advance_status(locked_advance, self.request.user)
        _log_finance_audit(
            self.request.user, action='DELETE', entity_type='FinanceAdvanceSettlement',
            entity_id=int(instance.id), entity_code=str(instance.id),
            old_values=snap, new_values={}, changed_fields=['deleted'],
        )

    @action(detail=False, methods=['post'])
    def bulk_delete(self, request):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền xóa quyết toán.'}, status=403)
        ids = request.data.get('ids', [])
        if not isinstance(ids, list) or not ids:
            return Response({'error': 'No IDs provided'}, status=400)

        with transaction.atomic():
            settlements = list(AdvanceSettlement.objects.filter(id__in=ids).select_related('advance_transaction'))
            locked_months = _get_locked_finance_months()
            for settlement in settlements:
                month = _month_from_date_obj(settlement.settlement_date)
                if month in locked_months:
                    return Response({'error': f'Tháng {month} đã khóa, không thể xóa quyết toán.'}, status=400)
                if settlement.advance_transaction.approval_status != AdvanceTransaction.APPROVAL_APPROVED:
                    return Response({'error': f'Phiếu {settlement.advance_transaction.code} chưa duyệt đầy đủ, không thể xóa quyết toán.'}, status=400)
            advance_ids = {item.advance_transaction_id for item in settlements}
            locked_advances = {
                row.id: row
                for row in AdvanceTransaction.objects.select_for_update().filter(id__in=advance_ids)
            }
            deleted_ids = [int(item.id) for item in settlements]
            deleted_count, _ = AdvanceSettlement.objects.filter(id__in=ids).delete()
            for advance_id in advance_ids:
                refresh_advance_status(locked_advances[advance_id], request.user)
        if deleted_ids:
            _log_finance_audit(
                request.user,
                action='DELETE',
                entity_type='FinanceAdvanceSettlementBulkDelete',
                entity_id=0,
                entity_code='bulk_delete',
                old_values={'ids': deleted_ids},
                new_values={},
                changed_fields=['deleted_ids'],
            )
        return Response({'success': True, 'count': deleted_count})

    def destroy(self, request, *args, **kwargs):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền xóa quyết toán.'}, status=403)
        return super().destroy(request, *args, **kwargs)


class ReceivableDocumentViewSet(SearchTextMixin, viewsets.ReadOnlyModelViewSet):
    serializer_class = ReceivableDocumentSerializer
    permission_classes = [IsAuthenticated]
    filter_backends = [filters.OrderingFilter]
    ordering_fields = ['document_date', 'due_date', 'total_amount', 'settled_amount', 'created_at']
    ordering = ['due_date', '-id']

    def get_queryset(self):
        queryset = ReceivableDocument.objects.select_related(
            'source_sales_order',
            'customer',
        ).prefetch_related(
            'settlements',
            'settlements__source_cash_account',
            'settlements__source_bank_account',
        )
        params = self.request.query_params
        status_value = (params.get('status') or '').strip()
        if status_value:
            queryset = queryset.filter(status=status_value)
        customer_id = (params.get('customer') or '').strip()
        if customer_id.isdigit():
            queryset = queryset.filter(customer_id=int(customer_id))
        document_date_gte = (params.get('document_date__gte') or '').strip()
        if document_date_gte:
            queryset = queryset.filter(document_date__gte=document_date_gte)
        document_date_lte = (params.get('document_date__lte') or '').strip()
        if document_date_lte:
            queryset = queryset.filter(document_date__lte=document_date_lte)
        month = (params.get('month') or '').strip()
        if len(month) == 7 and '-' in month:
            year_str, month_str = month.split('-', 1)
            if year_str.isdigit() and month_str.isdigit():
                queryset = queryset.filter(
                    due_date__year=int(year_str),
                    due_date__month=int(month_str),
                )
        overdue_only = str(params.get('overdue_only') or '').strip().lower() in {'1', 'true', 'yes'}
        if overdue_only:
            queryset = queryset.exclude(status=ReceivableStatus.CANCELLED).filter(due_date__lt=timezone.localdate())
        return self.apply_search(queryset)

    @action(detail=False, methods=['get'])
    def summary(self, request):
        queryset = self.filter_queryset(self.get_queryset())
        today = timezone.localdate()
        open_qs = queryset.exclude(status=ReceivableStatus.CANCELLED)
        total_amount = Decimal(str(open_qs.aggregate(total=Sum('total_amount')).get('total') or 0))
        settled_amount = Decimal(str(open_qs.aggregate(total=Sum('settled_amount')).get('total') or 0))
        remaining_amount = total_amount - settled_amount
        overdue_qs = open_qs.filter(due_date__lt=today).exclude(status=ReceivableStatus.SETTLED)
        overdue_amount = Decimal(str(overdue_qs.aggregate(total=Sum('total_amount')).get('total') or 0)) - Decimal(str(overdue_qs.aggregate(total=Sum('settled_amount')).get('total') or 0))
        return Response({
            'count': queryset.count(),
            'open_count': int(queryset.filter(status=ReceivableStatus.OPEN).count()),
            'partial_count': int(queryset.filter(status=ReceivableStatus.PARTIAL).count()),
            'settled_count': int(queryset.filter(status=ReceivableStatus.SETTLED).count()),
            'cancelled_count': int(queryset.filter(status=ReceivableStatus.CANCELLED).count()),
            'overdue_count': int(overdue_qs.count()),
            'total_amount': str(round_money(total_amount)),
            'settled_amount': str(round_money(settled_amount)),
            'remaining_amount': str(round_money(remaining_amount if remaining_amount > 0 else 0)),
            'overdue_amount': str(round_money(overdue_amount if overdue_amount > 0 else 0)),
        })

    @action(detail=True, methods=['post'])
    def collect(self, request, pk=None):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền ghi nhận thu tiền công nợ.'}, status=403)
        document = self.get_object()
        settlement_date_raw = request.data.get('settlement_date') or timezone.localdate().isoformat()
        try:
            settlement_date = date.fromisoformat(str(settlement_date_raw))
        except (TypeError, ValueError):
            return Response({'error': 'Ngày thu tiền không hợp lệ.'}, status=400)
        _ensure_finance_month_unlocked(
            _month_from_date_obj(settlement_date),
            'Tháng tài chính đã khóa, không thể ghi nhận thu tiền.',
        )
        try:
            amount = Decimal(str(request.data.get('amount') or 0))
        except (InvalidOperation, TypeError, ValueError):
            return Response({'error': 'Số tiền thu không hợp lệ.'}, status=400)
        if amount <= 0:
            return Response({'error': 'Số tiền thu phải lớn hơn 0.'}, status=400)
        source_type = str(request.data.get('source_type') or CashTransaction.SOURCE_CASH).strip().upper()
        if source_type not in {CashTransaction.SOURCE_CASH, CashTransaction.SOURCE_BANK}:
            return Response({'error': 'Nguồn tiền không hợp lệ.'}, status=400)

        category = ensure_system_transaction_category(
            code='AR_COLLECTION',
            name='Thu công nợ phải thu',
            category_type=CashTransaction.TYPE_INCOME,
            color='#52c41a',
        )
        with transaction.atomic():
            locked_document = ReceivableDocument.objects.select_for_update().get(pk=document.pk)
            if locked_document.status == ReceivableStatus.CANCELLED:
                return Response({'error': 'Chứng từ phải thu đã hủy, không thể thu tiền.'}, status=400)
            if amount > locked_document.remaining_amount:
                return Response({'error': f'Số tiền thu vượt số còn phải thu ({locked_document.remaining_amount:,.0f}).'}, status=400)
            tx_serializer = CashTransactionSerializer(data={
                'transaction_type': CashTransaction.TYPE_INCOME,
                'source_type': source_type,
                'source_cash_account': request.data.get('source_cash_account'),
                'source_bank_account': request.data.get('source_bank_account'),
                'category': category.id,
                'transaction_date': settlement_date_raw,
                'amount': str(amount),
                'object_name': (locked_document.customer_snapshot or {}).get('company_name') or (locked_document.customer_snapshot or {}).get('name') or '',
                'reason': f'[AR:{locked_document.id}] Thu công nợ {locked_document.code}',
                'note': str(request.data.get('note') or '').strip(),
            })
            tx_serializer.is_valid(raise_exception=True)
            cash_tx = tx_serializer.save(created_by=request.user)
            settlement = ReceivableSettlement.objects.create(
                receivable_document=locked_document,
                settlement_date=settlement_date,
                amount=amount,
                source_type=source_type,
                source_cash_account_id=request.data.get('source_cash_account'),
                source_bank_account_id=request.data.get('source_bank_account'),
                cash_transaction=cash_tx,
                note=str(request.data.get('note') or '').strip(),
                created_by=request.user,
                updated_by=request.user,
            )
            refresh_receivable_status(locked_document, actor=request.user)
        _log_finance_audit(
            request.user,
            action='CREATE',
            entity_type='FinanceReceivableSettlement',
            entity_id=int(settlement.id),
            entity_code=locked_document.code,
            old_values={},
            new_values={'receivable_document_id': int(locked_document.id), 'amount': str(amount), 'settlement_date': settlement_date.isoformat()},
            changed_fields=['amount', 'settlement_date'],
        )
        return Response(ReceivableDocumentSerializer(locked_document).data)

    @action(detail=True, methods=['post'])
    def cancel(self, request, pk=None):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền hủy chứng từ phải thu.'}, status=403)
        document = self.get_object()
        reason = str(request.data.get('reason') or request.data.get('cancel_reason') or '').strip()
        if not reason:
            return Response({'error': 'Bắt buộc nhập lý do hủy chứng từ.'}, status=400)
        _ensure_finance_month_unlocked(
            _month_from_date_obj(document.document_date),
            'Tháng tài chính đã khóa, không thể hủy chứng từ phải thu.',
        )
        if document.settled_amount > 0:
            return Response({'error': 'Chứng từ đã phát sinh thu tiền, không thể hủy.'}, status=400)
        old_status = document.status
        document.status = ReceivableStatus.CANCELLED
        document.note = ((document.note or '').strip() + f'\n[Huỷ] {reason}').strip()
        document.updated_by = request.user
        document.save(update_fields=['status', 'note', 'updated_by', 'updated_at'])
        _log_finance_audit(
            request.user,
            action='CANCEL',
            entity_type='FinanceReceivable',
            entity_id=int(document.id),
            entity_code=document.code,
            old_values={'status': old_status},
            new_values={'status': document.status, 'reason': reason},
            changed_fields=['status'],
        )
        return Response({'status': document.status})

    @action(detail=True, methods=['get'])
    def settlements(self, request, pk=None):
        document = self.get_object()
        items = document.settlements.select_related('source_cash_account', 'source_bank_account', 'cash_transaction').order_by('-settlement_date', '-id')
        return Response(ReceivableSettlementSerializer(items, many=True).data)


class PayableDocumentViewSet(SearchTextMixin, viewsets.ReadOnlyModelViewSet):
    serializer_class = PayableDocumentSerializer
    permission_classes = [IsAuthenticated]
    filter_backends = [filters.OrderingFilter]
    ordering_fields = ['document_date', 'due_date', 'total_amount', 'settled_amount', 'created_at']
    ordering = ['due_date', '-id']

    def get_queryset(self):
        queryset = PayableDocument.objects.select_related(
            'source_purchase_receipt',
            'source_purchase_receipt__purchase_order',
            'supplier',
        ).prefetch_related(
            'settlements',
            'settlements__source_cash_account',
            'settlements__source_bank_account',
        )
        params = self.request.query_params
        status_value = (params.get('status') or '').strip()
        if status_value:
            queryset = queryset.filter(status=status_value)
        supplier_id = (params.get('supplier') or '').strip()
        if supplier_id.isdigit():
            queryset = queryset.filter(supplier_id=int(supplier_id))
        document_date_gte = (params.get('document_date__gte') or '').strip()
        if document_date_gte:
            queryset = queryset.filter(document_date__gte=document_date_gte)
        document_date_lte = (params.get('document_date__lte') or '').strip()
        if document_date_lte:
            queryset = queryset.filter(document_date__lte=document_date_lte)
        month = (params.get('month') or '').strip()
        if len(month) == 7 and '-' in month:
            year_str, month_str = month.split('-', 1)
            if year_str.isdigit() and month_str.isdigit():
                queryset = queryset.filter(
                    due_date__year=int(year_str),
                    due_date__month=int(month_str),
                )
        overdue_only = str(params.get('overdue_only') or '').strip().lower() in {'1', 'true', 'yes'}
        if overdue_only:
            queryset = queryset.exclude(status=PayableStatus.CANCELLED).filter(due_date__lt=timezone.localdate())
        return self.apply_search(queryset)

    @action(detail=False, methods=['get'])
    def summary(self, request):
        queryset = self.filter_queryset(self.get_queryset())
        today = timezone.localdate()
        open_qs = queryset.exclude(status=PayableStatus.CANCELLED)
        total_amount = Decimal(str(open_qs.aggregate(total=Sum('total_amount')).get('total') or 0))
        settled_amount = Decimal(str(open_qs.aggregate(total=Sum('settled_amount')).get('total') or 0))
        remaining_amount = total_amount - settled_amount
        overdue_qs = open_qs.filter(due_date__lt=today).exclude(status=PayableStatus.SETTLED)
        overdue_amount = Decimal(str(overdue_qs.aggregate(total=Sum('total_amount')).get('total') or 0)) - Decimal(str(overdue_qs.aggregate(total=Sum('settled_amount')).get('total') or 0))
        return Response({
            'count': queryset.count(),
            'open_count': int(queryset.filter(status=PayableStatus.OPEN).count()),
            'partial_count': int(queryset.filter(status=PayableStatus.PARTIAL).count()),
            'settled_count': int(queryset.filter(status=PayableStatus.SETTLED).count()),
            'cancelled_count': int(queryset.filter(status=PayableStatus.CANCELLED).count()),
            'overdue_count': int(overdue_qs.count()),
            'total_amount': str(round_money(total_amount)),
            'settled_amount': str(round_money(settled_amount)),
            'remaining_amount': str(round_money(remaining_amount if remaining_amount > 0 else 0)),
            'overdue_amount': str(round_money(overdue_amount if overdue_amount > 0 else 0)),
        })

    @action(detail=True, methods=['post'])
    def pay(self, request, pk=None):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền ghi nhận chi trả công nợ.'}, status=403)
        document = self.get_object()
        settlement_date_raw = request.data.get('settlement_date') or timezone.localdate().isoformat()
        try:
            settlement_date = date.fromisoformat(str(settlement_date_raw))
        except (TypeError, ValueError):
            return Response({'error': 'Ngày thanh toán không hợp lệ.'}, status=400)
        _ensure_finance_month_unlocked(
            _month_from_date_obj(settlement_date),
            'Tháng tài chính đã khóa, không thể ghi nhận chi trả.',
        )
        try:
            amount = Decimal(str(request.data.get('amount') or 0))
        except (InvalidOperation, TypeError, ValueError):
            return Response({'error': 'Số tiền chi không hợp lệ.'}, status=400)
        if amount <= 0:
            return Response({'error': 'Số tiền chi phải lớn hơn 0.'}, status=400)
        source_type = str(request.data.get('source_type') or CashTransaction.SOURCE_CASH).strip().upper()
        if source_type not in {CashTransaction.SOURCE_CASH, CashTransaction.SOURCE_BANK}:
            return Response({'error': 'Nguồn tiền không hợp lệ.'}, status=400)

        category = ensure_system_transaction_category(
            code='AP_PAYMENT',
            name='Chi trả công nợ phải trả',
            category_type=CashTransaction.TYPE_EXPENSE,
            color='#faad14',
        )
        with transaction.atomic():
            locked_document = PayableDocument.objects.select_for_update().get(pk=document.pk)
            if locked_document.status == PayableStatus.CANCELLED:
                return Response({'error': 'Chứng từ phải trả đã hủy, không thể thanh toán.'}, status=400)
            if amount > locked_document.remaining_amount:
                return Response({'error': f'Số tiền chi vượt số còn phải trả ({locked_document.remaining_amount:,.0f}).'}, status=400)
            tx_serializer = CashTransactionSerializer(data={
                'transaction_type': CashTransaction.TYPE_EXPENSE,
                'source_type': source_type,
                'source_cash_account': request.data.get('source_cash_account'),
                'source_bank_account': request.data.get('source_bank_account'),
                'category': category.id,
                'transaction_date': settlement_date_raw,
                'amount': str(amount),
                'object_name': (locked_document.supplier_snapshot or {}).get('company_name') or (locked_document.supplier_snapshot or {}).get('name') or '',
                'reason': f'[AP:{locked_document.id}] Chi trả công nợ {locked_document.code}',
                'note': str(request.data.get('note') or '').strip(),
            })
            tx_serializer.is_valid(raise_exception=True)
            cash_tx = tx_serializer.save(created_by=request.user)
            settlement = PayableSettlement.objects.create(
                payable_document=locked_document,
                settlement_date=settlement_date,
                amount=amount,
                source_type=source_type,
                source_cash_account_id=request.data.get('source_cash_account'),
                source_bank_account_id=request.data.get('source_bank_account'),
                cash_transaction=cash_tx,
                note=str(request.data.get('note') or '').strip(),
                created_by=request.user,
                updated_by=request.user,
            )
            refresh_payable_status(locked_document, actor=request.user)
        _log_finance_audit(
            request.user,
            action='CREATE',
            entity_type='FinancePayableSettlement',
            entity_id=int(settlement.id),
            entity_code=locked_document.code,
            old_values={},
            new_values={'payable_document_id': int(locked_document.id), 'amount': str(amount), 'settlement_date': settlement_date.isoformat()},
            changed_fields=['amount', 'settlement_date'],
        )
        return Response(PayableDocumentSerializer(locked_document).data)

    @action(detail=True, methods=['post'])
    def cancel(self, request, pk=None):
        if not _can_manage_finance(request.user):
            return Response({'error': 'Bạn không có quyền hủy chứng từ phải trả.'}, status=403)
        document = self.get_object()
        reason = str(request.data.get('reason') or request.data.get('cancel_reason') or '').strip()
        if not reason:
            return Response({'error': 'Bắt buộc nhập lý do hủy chứng từ.'}, status=400)
        _ensure_finance_month_unlocked(
            _month_from_date_obj(document.document_date),
            'Tháng tài chính đã khóa, không thể hủy chứng từ phải trả.',
        )
        if document.settled_amount > 0:
            return Response({'error': 'Chứng từ đã phát sinh thanh toán, không thể hủy.'}, status=400)
        old_status = document.status
        document.status = PayableStatus.CANCELLED
        document.note = ((document.note or '').strip() + f'\n[Huỷ] {reason}').strip()
        document.updated_by = request.user
        document.save(update_fields=['status', 'note', 'updated_by', 'updated_at'])
        _log_finance_audit(
            request.user,
            action='CANCEL',
            entity_type='FinancePayable',
            entity_id=int(document.id),
            entity_code=document.code,
            old_values={'status': old_status},
            new_values={'status': document.status, 'reason': reason},
            changed_fields=['status'],
        )
        return Response({'status': document.status})

    @action(detail=True, methods=['get'])
    def settlements(self, request, pk=None):
        document = self.get_object()
        items = document.settlements.select_related('source_cash_account', 'source_bank_account', 'cash_transaction').order_by('-settlement_date', '-id')
        return Response(PayableSettlementSerializer(items, many=True).data)


# ============== GENERAL LEDGER ==============
from decimal import Decimal as DecimalType
from django.db.models import Sum

class GeneralLedgerAccountViewSet(viewsets.ModelViewSet):
    """Chart of Accounts Management"""
    queryset = GeneralLedgerAccount.objects.all()
    serializer_class = GeneralLedgerAccountSerializer
    filter_backends = [DjangoFilterBackend, SearchFilter, OrderingFilter]
    filterset_fields = ['account_type', 'is_active']
    search_fields = ['code', 'name']
    ordering_fields = ['code', 'created_at']
    ordering = ['code']
    permission_classes = [IsAuthenticated]


class GeneralLedgerViewSet(viewsets.ReadOnlyModelViewSet):
    """General Ledger Entries - Read-only"""
    serializer_class = GeneralLedgerEntrySerializer
    filter_backends = [DjangoFilterBackend, SearchFilter, OrderingFilter]
    filterset_fields = ['account', 'posting_date', 'document_type']
    search_fields = ['document_code', 'description', 'account__code', 'account__name']
    ordering_fields = ['posting_date', 'created_at']
    ordering = ['-posting_date', '-id']
    permission_classes = [IsAuthenticated]
    
    def get_queryset(self):
        return GeneralLedgerEntry.objects.select_related("account", "created_by").all()
    
    @action(detail=False, methods=["get"])
    def trial_balance(self, request):
        """Get trial balance - sum by account"""
        date_from = request.query_params.get("date_from")
        date_to = request.query_params.get("date_to")
        
        queryset = GeneralLedgerEntry.objects.all()
        if date_from:
            queryset = queryset.filter(posting_date__gte=date_from)
        if date_to:
            queryset = queryset.filter(posting_date__lte=date_to)
        
        accounts = GeneralLedgerAccount.objects.filter(is_active=True)
        data = []
        total_debit = DecimalType(0)
        total_credit = DecimalType(0)
        
        for account in accounts:
            entries = queryset.filter(account=account)
            debit = entries.aggregate(Sum("debit_amount"))["debit_amount__sum"] or DecimalType(0)
            credit = entries.aggregate(Sum("credit_amount"))["credit_amount__sum"] or DecimalType(0)
            
            if debit > 0 or credit > 0:
                data.append({
                    "account_code": account.code,
                    "account_name": account.name,
                    "account_type": account.account_type,
                    "debit": str(debit),
                    "credit": str(credit),
                })
                total_debit += debit
                total_credit += credit
        
        data.append({
            "account_code": "TOTAL",
            "account_name": "TỔNG CỘNG",
            "account_type": "",
            "debit": str(total_debit),
            "credit": str(total_credit),
        })
        
        return Response(data)
    
    @action(detail=False, methods=["get"])
    def account_balance(self, request):
        """Get balance for specific account"""
        account_id = request.query_params.get("account_id")
        date_to = request.query_params.get("date_to")
        
        if not account_id:
            return Response({"error": "account_id required"}, status=400)
        
        queryset = GeneralLedgerEntry.objects.filter(account_id=account_id)
        if date_to:
            queryset = queryset.filter(posting_date__lte=date_to)
        
        debit_total = queryset.aggregate(Sum("debit_amount"))["debit_amount__sum"] or DecimalType(0)
        credit_total = queryset.aggregate(Sum("credit_amount"))["credit_amount__sum"] or DecimalType(0)
        balance = debit_total - credit_total
        
        return Response({
            "debit_total": str(debit_total),
            "credit_total": str(credit_total),
            "balance": str(balance),
        })
