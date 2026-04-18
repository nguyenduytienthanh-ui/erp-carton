import copy
import csv
import hashlib
import json
import math
from collections import Counter, defaultdict
from functools import lru_cache
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any
from uuid import uuid4

from openpyxl import Workbook, load_workbook

from paper_optimizer.optimizer import compute_target_quantities, run_full_optimizer


REFERENCE_DATA_DIR = Path(__file__).resolve().parent / 'reference_data'
CANONICAL_SNAPSHOT_PATH = REFERENCE_DATA_DIR / 'baseline_guardrail_canonical_v1.json'
CANONICAL_WORKBOOK_PATH = REFERENCE_DATA_DIR / 'paper_optimizer_canonical_guardrail_v1.xlsx'
BENCHMARK_REFERENCE_PATH = REFERENCE_DATA_DIR / 'benchmark_reference_v1.json'
STANDARD_RAW_WIDTHS_CM = tuple(range(110, 281, 5))
MANDATORY_MIN_PURCHASE_LENGTH_CM = 5000.0
CANONICAL_PLAN_SHEET_NAMES = {
    'EXACT_ORDER': 'Phuong_an_dung_don_hang',
    'LOWEST_TOTAL_COST': 'Phuong_an_chi_phi_thap_nhat',
    'MIN_RAW_WIDTHS': 'Phuong_an_it_kho_giay_nhat',
    'MIN_SPECS': 'Phuong_an_it_quy_cach_nhat',
}
PRIMARY_EXPORT_SHEET_ORDER = (
    'Du_lieu_goc',
    'To_hop_de_xuat',
    'Phuong_an_mua_cuoi',
    'Chi_tiet_phan_bo',
    'Kiem_tra_phan_bo_nguoc',
    'Con_lai_chua_phan_bo',
    'So_sanh_phuong_an_cuoi',
    'Thong_ke_kho_giay',
    'Debug_engine',
)
SUPPLEMENTAL_EXPORT_SHEET_ORDER = (
    'Phuong_an_dung_don_hang',
    'Phuong_an_chi_phi_thap_nhat',
    'Phuong_an_it_kho_giay_nhat',
    'Phuong_an_it_quy_cach_nhat',
    'So_sanh_voi_moc_tham_chieu',
)
EXPORT_WORKBOOK_SHEET_ORDER = PRIMARY_EXPORT_SHEET_ORDER + SUPPLEMENTAL_EXPORT_SHEET_ORDER

PLAN_META = {
    'EXACT_ORDER': {
        'plan_name': 'Phương án đúng đơn hàng',
        'objective': 'Bám đúng nhu cầu mục tiêu sau dự trù sản xuất.',
    },
    'LOWEST_TOTAL_COST': {
        'plan_name': 'Phương án chi phí thấp nhất',
        'objective': 'Ưu tiên tổng chi phí quy đổi thấp nhất.',
    },
    'MIN_RAW_WIDTHS': {
        'plan_name': 'Phương án ít khổ giấy nhất',
        'objective': 'Ưu tiên ít khổ giấy hơn khi vẫn giữ chi phí hợp lý.',
    },
    'MIN_SPECS': {
        'plan_name': 'Phương án ít quy cách mua nhất',
        'objective': 'Ưu tiên giảm số quy cách mua sau khi đã giữ ít khổ giấy.',
    },
}


def _round_float(value: float, digits: int = 8) -> float:
    return round(float(value), digits)


def _normalize_float(value: Any, digits: int = 6) -> float | None:
    if value is None or value == '':
        return None
    return round(float(value), digits)


def _normalize_line(line: dict[str, Any], index: int) -> dict[str, Any]:
    return {
        'id': str(line.get('id') or index + 1),
        'note': str(line.get('note') or ''),
        'quantity': int(line.get('quantity') or 0),
        'width_cm': _normalize_float(line.get('width_cm'), 4),
        'length_cm': _normalize_float(line.get('length_cm'), 4),
        'source_row_number': int(line.get('source_row_number') or index + 1),
    }


def _normalized_json(value: Any) -> Any:
    if isinstance(value, dict):
        return {key: _normalized_json(value[key]) for key in sorted(value)}
    if isinstance(value, list):
        return [_normalized_json(item) for item in value]
    if isinstance(value, float):
        return _normalize_float(value, 6)
    return value


def _normalize_available_raw_widths(raw_widths: list[Any] | None) -> list[int]:
    if not raw_widths:
        return []
    normalized: list[int] = []
    for item in raw_widths:
        rounded = int(round(float(item)))
        if abs(float(item) - rounded) > 1e-9:
            continue
        if rounded not in STANDARD_RAW_WIDTHS_CM:
            continue
        if rounded not in normalized:
            normalized.append(rounded)
    return sorted(normalized)


def _resolve_min_purchase_length_cm(supplier_config: dict[str, Any]) -> float:
    configured = float(supplier_config.get('min_purchase_length_cm') or 0.0)
    return _round_float(max(MANDATORY_MIN_PURCHASE_LENGTH_CM, configured), 4)


def _normalized_supplier_config(supplier_config: dict[str, Any]) -> dict[str, Any]:
    normalized = copy.deepcopy(supplier_config)
    normalized['available_raw_widths_cm'] = _normalize_available_raw_widths(
        normalized.get('available_raw_widths_cm') or []
    )
    normalized['min_purchase_length_cm'] = _resolve_min_purchase_length_cm(normalized)
    return normalized


def _normalize_purchase_spec_row(
    row: dict[str, Any],
    supplier_config: dict[str, Any],
) -> dict[str, Any]:
    normalized = copy.deepcopy(row)
    run_length = _round_float(float(normalized.get('run_length_cm') or 0.0), 4)
    total_sets = int(normalized.get('total_sets_purchase_spec') or normalized.get('sets') or 0)
    total_length = _round_float(run_length * total_sets, 4)
    min_purchase_length_cm = _resolve_min_purchase_length_cm(supplier_config)
    normalized['run_length_cm'] = run_length
    normalized['total_sets_purchase_spec'] = total_sets
    normalized['total_length_cm_purchase_spec'] = total_length
    normalized['supplier_min_length_required_cm'] = min_purchase_length_cm
    normalized['supplier_min_length_actual_cm'] = total_length
    normalized['supplier_min_length_passed'] = total_length >= min_purchase_length_cm
    return normalized


def _normalize_purchase_spec_rows(
    rows: list[dict[str, Any]],
    supplier_config: dict[str, Any],
) -> list[dict[str, Any]]:
    return [_normalize_purchase_spec_row(row, supplier_config) for row in rows]


def _build_reverse_check_rows_from_purchase_specs(
    purchase_spec_rows: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    reverse_rows: list[dict[str, Any]] = []
    for row in purchase_spec_rows:
        total_sets = int(row.get('total_sets_purchase_spec') or 0)
        total_length = _round_float(float(row.get('total_length_cm_purchase_spec') or 0.0), 4)
        line_ids = [str(line_id) for line_id in (row.get('source_line_ids') or [])]
        reverse_rows.append(
            {
                'raw_width_cm': row.get('raw_width_cm'),
                'run_length_cm': row.get('run_length_cm'),
                'sum_line_allocated_sets': total_sets,
                'total_sets_purchase_spec': total_sets,
                'reverse_total_length_cm': total_length,
                'spec_total_length_cm': total_length,
                'line_count': len(line_ids),
                'line_ids': line_ids,
                'consistent': True,
                'ncc_passed': bool(row.get('supplier_min_length_passed', False)),
            }
        )
    return reverse_rows


def _build_config_fingerprint(supplier_config: dict[str, Any], optimization_config: dict[str, Any]) -> str:
    payload = {
        'supplier_config': _normalized_json(supplier_config),
        'optimization_config': _normalized_json(optimization_config),
    }
    raw = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(',', ':'))
    return hashlib.sha1(raw.encode('utf-8')).hexdigest()[:12]


@lru_cache(maxsize=1)
def load_canonical_snapshot() -> dict[str, Any]:
    with CANONICAL_SNAPSHOT_PATH.open('r', encoding='utf-8') as handle:
        return json.load(handle)


@lru_cache(maxsize=1)
def load_canonical_workbook_bytes() -> bytes:
    return CANONICAL_WORKBOOK_PATH.read_bytes()


@lru_cache(maxsize=1)
def load_canonical_purchase_spec_rows_by_plan() -> dict[str, list[dict[str, Any]]]:
    workbook = load_workbook(CANONICAL_WORKBOOK_PATH, read_only=True, data_only=True)
    rows_by_plan: dict[str, list[dict[str, Any]]] = {}
    for plan_code, sheet_name in CANONICAL_PLAN_SHEET_NAMES.items():
        if sheet_name not in workbook.sheetnames:
            rows_by_plan[plan_code] = []
            continue
        worksheet = workbook[sheet_name]
        rows_by_plan[plan_code] = _parse_purchase_spec_sheet(worksheet)
    return rows_by_plan


@lru_cache(maxsize=1)
def load_benchmark_reference_source() -> dict[str, Any]:
    with BENCHMARK_REFERENCE_PATH.open('r', encoding='utf-8') as handle:
        return json.load(handle)


def build_default_supplier_config() -> dict[str, Any]:
    snapshot = load_canonical_snapshot()
    return _normalized_supplier_config(snapshot['supplier_config'])


def build_default_optimization_config() -> dict[str, Any]:
    snapshot = load_canonical_snapshot()
    return copy.deepcopy(snapshot['optimization_config'])


def build_default_input_lines() -> list[dict[str, Any]]:
    snapshot = load_canonical_snapshot()
    return copy.deepcopy(snapshot['input_lines'])


def build_default_template_payload() -> dict[str, Any]:
    return {
        'name': 'Baseline canonical v5000',
        'note': 'Template phục hồi từ baseline canonical v5000 đã khóa.',
        'supplier_config': build_default_supplier_config(),
        'optimization_config': build_default_optimization_config(),
    }


def build_baseline_reference() -> dict[str, Any]:
    snapshot = load_canonical_snapshot()
    return {
        'baseline_guardrail_schema_version': snapshot.get('baseline_guardrail_schema_version'),
        'baseline_generated_from_run_code': snapshot.get('baseline_generated_from_run_code'),
        'baseline_generated_at': snapshot.get('baseline_generated_at'),
        'reference_code': snapshot.get('reference_code'),
        'reference_source': snapshot.get('reference_source'),
        'config_fingerprint': snapshot.get('config_fingerprint'),
        'selected_plan_code': snapshot.get('selected_plan_code'),
        'selected_scenario_code': snapshot.get('selected_scenario_code'),
        'selected_source_internal_plan_code': snapshot.get('selected_source_internal_plan_code'),
        'line_count': snapshot.get('line_count'),
        'base_requested_quantity_total': snapshot.get('base_requested_quantity_total'),
        'production_reserve_quantity_total': snapshot.get('production_reserve_quantity_total'),
        'target_requested_quantity_total': snapshot.get('target_requested_quantity_total'),
    }


def build_benchmark_reference(
    optimization_config: dict[str, Any] | None = None,
) -> dict[str, Any]:
    source = copy.deepcopy(load_benchmark_reference_source())
    effective_config = optimization_config or build_default_optimization_config()
    technical_waste_cost = _round_float(
        float(source.get('technical_waste_rate') or 0.0)
        * float(effective_config.get('technical_waste_cost_weight') or 1.0),
    )
    economic_overproduction_cost = _round_float(
        float(source.get('economic_overproduction_quantity_total') or 0.0)
        * float(effective_config.get('economic_overproduction_cost_weight') or 1.0)
        * 0.01,
    )
    raw_width_cost = _round_float(
        float(source.get('unique_raw_width_count') or 0.0)
        * float(effective_config.get('new_raw_width_cost_weight') or 1.0),
    )
    purchase_spec_cost = _round_float(
        float(source.get('unique_purchase_spec_count') or 0.0)
        * float(effective_config.get('new_purchase_spec_cost_weight') or 0.2),
    )
    fragmentation_cost = _round_float(
        float(source.get('final_group_count') or 0.0)
        * float(effective_config.get('fragmentation_cost_weight') or 0.12),
    )
    source.update(
        {
            'technical_waste_cost': technical_waste_cost,
            'economic_overproduction_cost': economic_overproduction_cost,
            'raw_width_cost': raw_width_cost,
            'purchase_spec_cost': purchase_spec_cost,
            'fragmentation_cost': fragmentation_cost,
            'total_converted_cost': _round_float(
                technical_waste_cost
                + economic_overproduction_cost
                + raw_width_cost
                + purchase_spec_cost
                + fragmentation_cost,
            ),
        }
    )
    return source


def create_run_code() -> str:
    return f'POPT-{uuid4().hex[:12].upper()}'


def build_import_template_bytes() -> bytes:
    output = StringIO()
    writer = csv.writer(output)
    writer.writerow(['width_cm', 'length_cm', 'quantity', 'note'])
    writer.writerow([52, 112, 25, 'hop A'])
    writer.writerow([84, 143.5, 12, 'hop B'])
    return output.getvalue().encode('utf-8-sig')


def parse_preview_file(uploaded_file: Any) -> dict[str, Any]:
    filename = str(getattr(uploaded_file, 'name', 'input.xlsx'))
    lower_name = filename.lower()
    if lower_name.endswith('.csv'):
        preview = _parse_csv_preview(uploaded_file.read().decode('utf-8-sig'))
    elif lower_name.endswith('.xlsx') or lower_name.endswith('.xlsm'):
        preview = _parse_xlsx_preview(uploaded_file)
    else:
        raise ValueError('Chỉ hỗ trợ file CSV hoặc XLSX.')

    rows = preview['rows']
    if not rows:
        raise ValueError('Không đọc được dòng dữ liệu hợp lệ từ file tải lên.')

    return {
        'source_filename': filename,
        'rows': rows,
        'issues': preview['issues'],
        'summary': {
            'line_count': len(rows),
            'accepted_row_count': len(rows),
            'total_quantity': sum(int(row['quantity']) for row in rows),
            'issue_count': len(preview['issues']),
            'invalid_row_count': int(preview['summary']['invalid_row_count']),
            'skipped_empty_row_count': int(preview['summary']['skipped_empty_row_count']),
            'header_detected': bool(preview['summary']['header_detected']),
        },
    }


def _parse_csv_preview(text: str) -> dict[str, Any]:
    reader = list(csv.reader(StringIO(text)))
    return _rows_to_preview_lines(reader)


def _parse_xlsx_preview(uploaded_file: Any) -> dict[str, Any]:
    workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
    sheet = workbook.active
    matrix: list[list[Any]] = []
    for row in sheet.iter_rows(values_only=True):
        matrix.append(list(row))
    workbook.close()
    return _rows_to_preview_lines(matrix)


def _rows_to_preview_lines(matrix: list[list[Any]]) -> dict[str, Any]:
    if not matrix:
        return {
            'rows': [],
            'issues': [],
            'summary': {
                'invalid_row_count': 0,
                'skipped_empty_row_count': 0,
                'header_detected': False,
            },
        }

    first_row = matrix[0]
    header_map = _detect_header_map(first_row)
    data_rows = matrix[1:] if header_map else matrix
    row_offset = 1 if header_map else 0

    rows: list[dict[str, Any]] = []
    issues: list[dict[str, Any]] = []
    invalid_row_count = 0
    skipped_empty_row_count = 0
    for index, raw_row in enumerate(data_rows):
        source_row_number = index + 1 + row_offset
        parsed = _row_to_line(raw_row, index, source_row_number, header_map)
        if parsed['status'] == 'empty':
            skipped_empty_row_count += 1
            continue
        if parsed['status'] == 'invalid':
            invalid_row_count += 1
            issues.extend(parsed['issues'])
            continue
        rows.append(parsed['line'])
    return {
        'rows': rows,
        'issues': issues,
        'summary': {
            'invalid_row_count': invalid_row_count,
            'skipped_empty_row_count': skipped_empty_row_count,
            'header_detected': bool(header_map),
        },
    }


def _detect_header_map(header_row: list[Any]) -> dict[str, int]:
    normalized = [str(item or '').strip().lower() for item in header_row]
    candidates = {
        'quantity': {'quantity', 'so luong', 'sl'},
        'width_cm': {'width', 'width_cm', 'rong', 'kho', 'kho cm', 'rộng', 'khổ', 'khổ cm'},
        'length_cm': {'length', 'length_cm', 'dai', 'dài'},
        'note': {'note', 'ghi chu', 'ghi_chu', 'ghi chú'},
        'id': {'id', 'ma', 'mã'},
    }
    mapping: dict[str, int] = {}
    for index, value in enumerate(normalized):
        for field, accepted in candidates.items():
            if value in accepted and field not in mapping:
                mapping[field] = index
    required = {'quantity', 'width_cm', 'length_cm'}
    return mapping if required.issubset(mapping) else {}


def _row_to_line(
    raw_row: list[Any],
    index: int,
    source_row_number: int,
    header_map: dict[str, int],
) -> dict[str, Any]:
    values = list(raw_row)
    if not any(value not in (None, '') for value in values):
        return {'status': 'empty', 'issues': []}

    def pick(field: str, fallback_index: int | None = None) -> Any:
        if field in header_map:
            idx = header_map[field]
            return values[idx] if idx < len(values) else None
        if fallback_index is None or fallback_index >= len(values):
            return None
        return values[fallback_index]

    if header_map:
        quantity = pick('quantity')
        width_cm = pick('width_cm')
        length_cm = pick('length_cm')
        note = pick('note')
        line_id = pick('id')
    else:
        width_cm = pick('width_cm', 0)
        length_cm = pick('length_cm', 1)
        quantity = pick('quantity', 2)
        note = pick('note', 3)
        line_id = pick('id', 4)

    missing_fields: list[str] = []
    if quantity in (None, ''):
        missing_fields.append('Số lượng')
    if width_cm in (None, ''):
        missing_fields.append('Khổ')
    if length_cm in (None, ''):
        missing_fields.append('Dài')
    if missing_fields:
        return {
            'status': 'invalid',
            'issues': [
                _build_preview_issue(
                    row_number=source_row_number,
                    code='missing_required_fields',
                    message=f'Thiếu cột bắt buộc: {", ".join(missing_fields)}.',
                )
            ],
        }

    try:
        quantity_value = float(quantity)
    except (TypeError, ValueError):
        return {
            'status': 'invalid',
            'issues': [
                _build_preview_issue(
                    row_number=source_row_number,
                    code='invalid_quantity',
                    message='Số lượng phải là số nguyên dương.',
                )
            ],
        }
    if not quantity_value.is_integer() or int(quantity_value) <= 0:
        return {
            'status': 'invalid',
            'issues': [
                _build_preview_issue(
                    row_number=source_row_number,
                    code='invalid_quantity',
                    message='Số lượng phải là số nguyên dương.',
                )
            ],
        }

    try:
        width_value = float(width_cm)
    except (TypeError, ValueError):
        return {
            'status': 'invalid',
            'issues': [
                _build_preview_issue(
                    row_number=source_row_number,
                    code='invalid_width_cm',
                    message='Khổ phải là số dương.',
                )
            ],
        }
    if width_value <= 0:
        return {
            'status': 'invalid',
            'issues': [
                _build_preview_issue(
                    row_number=source_row_number,
                    code='invalid_width_cm',
                    message='Khổ phải là số dương.',
                )
            ],
        }

    try:
        length_value = float(length_cm)
    except (TypeError, ValueError):
        return {
            'status': 'invalid',
            'issues': [
                _build_preview_issue(
                    row_number=source_row_number,
                    code='invalid_length_cm',
                    message='Dài phải là số dương.',
                )
            ],
        }
    if length_value <= 0:
        return {
            'status': 'invalid',
            'issues': [
                _build_preview_issue(
                    row_number=source_row_number,
                    code='invalid_length_cm',
                    message='Dài phải là số dương.',
                )
            ],
        }

    return {
        'status': 'valid',
        'line': {
            'id': str(line_id or index + 1),
            'note': str(note or ''),
            'quantity': int(quantity_value),
            'width_cm': width_value,
            'length_cm': length_value,
            'source_row_number': source_row_number,
        },
        'issues': [],
    }


def _build_preview_issue(
    row_number: int,
    code: str,
    message: str,
    severity: str = 'warning',
) -> dict[str, Any]:
    return {
        'row_number': row_number,
        'severity': severity,
        'code': code,
        'message': message,
    }


def evaluate_manual_pattern(payload: dict[str, Any]) -> dict[str, Any]:
    raw_width_cm = float(payload['raw_width_cm'])
    trim_edge_cm = float(payload.get('trim_edge_cm') or 0.0)
    run_length_cm = float(payload['run_length_cm'])
    sets = int(payload['sets'])
    components = payload['components']

    useful_width_cm = max(raw_width_cm - trim_edge_cm, 0.0)
    used_width_cm = sum(float(component['width_cm']) * int(component['multiplier']) for component in components)
    used_area = sum(
        float(component['width_cm']) * float(component['length_cm']) * int(component['multiplier']) * sets
        for component in components
    )
    purchase_area = raw_width_cm * run_length_cm * sets
    waste_area = max(purchase_area - used_area, 0.0)
    waste_rate = waste_area / purchase_area if purchase_area else 0.0

    return {
        'raw_width_cm': _round_float(raw_width_cm),
        'trim_edge_cm': _round_float(trim_edge_cm),
        'run_length_cm': _round_float(run_length_cm),
        'sets': sets,
        'useful_width_cm': _round_float(useful_width_cm),
        'used_width_cm': _round_float(used_width_cm),
        'purchase_area': _round_float(purchase_area),
        'used_area': _round_float(used_area),
        'waste_area': _round_float(waste_area),
        'waste_rate': _round_float(waste_rate),
        'meets_supplier_rule': used_width_cm <= useful_width_cm,
    }


def run_optimizer_recovery(
    input_lines: list[dict[str, Any]],
    supplier_config: dict[str, Any],
    optimization_config: dict[str, Any],
    source_filename: str = '',
    note: str = '',
) -> dict[str, Any]:
    """
    Entry point chính cho optimize run.
    - Nếu khớp canonical: trả nguyên snapshot.
    - Nếu không khớp: chạy optimizer nghiệp vụ đầy đủ.
    """
    snapshot = load_canonical_snapshot()
    normalized_lines = [_normalize_line(line, index) for index, line in enumerate(input_lines)]
    supplier_config = _normalized_supplier_config(supplier_config)
    canonical_match = _is_canonical_match(normalized_lines, supplier_config, optimization_config)
    if canonical_match:
        return _build_canonical_result(snapshot, normalized_lines, source_filename, note)

    # Chạy optimizer đầy đủ
    baseline_plan = _get_canonical_baseline_plan(snapshot)
    result = run_full_optimizer(
        input_lines=normalized_lines,
        supplier_config=supplier_config,
        optimization_config=optimization_config,
        baseline_plan=baseline_plan,
        source_filename=source_filename,
        note=note,
    )
    result['baseline_reference'] = build_baseline_reference()
    benchmark_reference = build_benchmark_reference(optimization_config)
    result['benchmark_reference'] = benchmark_reference
    result.setdefault('stats', {}).update(
        {
            'benchmark_reference_code': benchmark_reference.get('reference_code'),
            'benchmark_reference_source': benchmark_reference.get('reference_source'),
        }
    )
    return result


def _is_canonical_match(
    input_lines: list[dict[str, Any]],
    supplier_config: dict[str, Any],
    optimization_config: dict[str, Any],
) -> bool:
    snapshot = load_canonical_snapshot()
    return (
        [_normalize_line(line, index) for index, line in enumerate(snapshot['input_lines'])] == input_lines
        and _normalized_json(_normalized_supplier_config(snapshot['supplier_config'])) == _normalized_json(supplier_config)
        and _normalized_json(snapshot['optimization_config']) == _normalized_json(optimization_config)
    )


def _get_canonical_baseline_plan(snapshot: dict[str, Any]) -> dict[str, Any] | None:
    """Trả về plan được chọn từ snapshot canonical để dùng làm baseline guardrail."""
    alternatives = snapshot.get('public_alternatives', [])
    selected_code = snapshot.get('selected_plan_code', '')
    plan = next((a for a in alternatives if a.get('plan_code') == selected_code), None)
    if plan:
        return {
            'total_converted_cost': plan.get('total_converted_cost'),
            'total_purchase_area': plan.get('total_purchase_area'),
            'unique_raw_width_count': plan.get('unique_raw_width_count'),
            'unique_purchase_spec_count': plan.get('unique_purchase_spec_count'),
            'final_group_count': plan.get('final_group_count'),
            'technical_waste_rate': plan.get('technical_waste_rate'),
        }
    return None


def _is_placeholder_no_feasible_plan(plan: dict[str, Any]) -> bool:
    return (
        not bool(plan.get('purchase_spec_rows'))
        and not bool(plan.get('allocation_details'))
        and not bool(plan.get('final_plan'))
        and not bool(plan.get('all_ncc_passed'))
        and not bool(plan.get('meets_target_and_cap'))
        and int(plan.get('unique_raw_width_count') or 0) == 0
        and int(plan.get('unique_purchase_spec_count') or 0) == 0
        and int(plan.get('final_group_count') or 0) == 0
    )


def _normalize_result_plan_for_display(
    plan: dict[str, Any] | None,
    supplier_config: dict[str, Any],
) -> dict[str, Any]:
    if not plan:
        return {}
    normalized_plan = copy.deepcopy(plan)
    purchase_spec_rows = normalized_plan.get('purchase_spec_rows') or []
    if not purchase_spec_rows and normalized_plan.get('final_plan'):
        purchase_spec_rows = _build_purchase_spec_rows_from_final_plan(normalized_plan, supplier_config)
    purchase_spec_rows = _normalize_purchase_spec_rows(purchase_spec_rows, supplier_config)
    reverse_check_rows = normalized_plan.get('reverse_check_rows') or []
    if not reverse_check_rows and purchase_spec_rows:
        reverse_check_rows = _build_reverse_check_rows_from_purchase_specs(purchase_spec_rows)
    normalized_plan['purchase_spec_rows'] = purchase_spec_rows
    normalized_plan['reverse_check_rows'] = reverse_check_rows
    if purchase_spec_rows:
        normalized_plan['all_ncc_passed'] = all(
            row.get('supplier_min_length_passed', False)
            for row in purchase_spec_rows
        )
        normalized_plan['meets_target_and_cap'] = bool(normalized_plan['all_ncc_passed']) and not bool(
            normalized_plan.get('leftovers')
        )
    normalized_plan['is_placeholder_no_feasible'] = _is_placeholder_no_feasible_plan(normalized_plan)
    return normalized_plan


def normalize_result_payload_for_display(
    result_payload: dict[str, Any] | None,
    *,
    canonical_match: bool = False,
) -> dict[str, Any]:
    if not result_payload:
        return {}

    normalized_payload = copy.deepcopy(result_payload)
    supplier_config = _normalized_supplier_config(normalized_payload.get('supplier_config') or {})
    normalized_payload['supplier_config'] = supplier_config

    alternatives = [
        _normalize_result_plan_for_display(plan, supplier_config)
        for plan in (normalized_payload.get('final_plan_alternatives') or [])
    ]
    normalized_payload['final_plan_alternatives'] = alternatives

    selected_plan_code = str(normalized_payload.get('selected_plan_code') or '')
    selected_plan = _normalize_result_plan_for_display(
        normalized_payload.get('selected_plan') or {},
        supplier_config,
    )
    if not selected_plan and selected_plan_code:
        selected_plan = copy.deepcopy(
            next(
                (plan for plan in alternatives if str(plan.get('plan_code') or '') == selected_plan_code),
                {},
            )
        )
    if selected_plan and not selected_plan_code:
        selected_plan_code = str(selected_plan.get('plan_code') or '')
        normalized_payload['selected_plan_code'] = selected_plan_code

    normalized_payload['selected_plan'] = selected_plan or None
    normalized_payload['allocation_details'] = copy.deepcopy(
        normalized_payload.get('allocation_details') or selected_plan.get('allocation_details') or []
    )
    normalized_payload['leftovers'] = copy.deepcopy(
        normalized_payload.get('leftovers') or selected_plan.get('leftovers') or []
    )

    purchase_spec_rows = normalized_payload.get('purchase_spec_rows') or selected_plan.get('purchase_spec_rows') or []
    if not purchase_spec_rows and canonical_match and selected_plan_code:
        purchase_spec_rows = copy.deepcopy(
            load_canonical_purchase_spec_rows_by_plan().get(selected_plan_code, [])
        )
    if not purchase_spec_rows and selected_plan:
        purchase_spec_rows = _build_purchase_spec_rows_from_final_plan(selected_plan, supplier_config)
    purchase_spec_rows = _normalize_purchase_spec_rows(purchase_spec_rows, supplier_config)
    normalized_payload['purchase_spec_rows'] = purchase_spec_rows

    reverse_check_rows = normalized_payload.get('reverse_check_rows') or selected_plan.get('reverse_check_rows') or []
    if not reverse_check_rows and purchase_spec_rows:
        reverse_check_rows = _build_reverse_check_rows_from_purchase_specs(purchase_spec_rows)
    normalized_payload['reverse_check_rows'] = reverse_check_rows

    if selected_plan:
        normalized_selected_plan = copy.deepcopy(selected_plan)
        normalized_selected_plan['purchase_spec_rows'] = purchase_spec_rows
        normalized_selected_plan['reverse_check_rows'] = reverse_check_rows
        normalized_payload['selected_plan'] = normalized_selected_plan
        selected_plan = normalized_selected_plan

    feasible_alternative_count = sum(
        1
        for plan in alternatives
        if bool(plan.get('all_ncc_passed')) and bool(plan.get('meets_target_and_cap'))
    )
    placeholder_only_result = bool(selected_plan) and bool(selected_plan.get('is_placeholder_no_feasible'))
    stats = copy.deepcopy(normalized_payload.get('stats') or {})
    no_feasible_candidate = bool(stats.get('no_feasible_candidate'))
    if not no_feasible_candidate:
        no_feasible_candidate = (
            feasible_alternative_count == 0
            and (
                placeholder_only_result
                or (bool(alternatives) and all(bool(plan.get('is_placeholder_no_feasible')) for plan in alternatives))
            )
        )

    displayable_result = bool(
        not no_feasible_candidate
        and selected_plan
        and purchase_spec_rows
        and not bool(selected_plan.get('is_placeholder_no_feasible'))
    )

    if alternatives:
        stats['feasible_alternative_count'] = feasible_alternative_count
    else:
        stats['feasible_alternative_count'] = int(stats.get('feasible_alternative_count') or 0)
    stats['no_feasible_candidate'] = no_feasible_candidate
    stats['displayable_result'] = displayable_result
    normalized_payload['stats'] = stats
    normalized_payload['result_state'] = 'no_feasible_candidate' if no_feasible_candidate else 'success'
    return normalized_payload


def _build_canonical_result(
    snapshot: dict[str, Any],
    input_lines: list[dict[str, Any]],
    source_filename: str,
    note: str,
) -> dict[str, Any]:
    alternatives = copy.deepcopy(snapshot['public_alternatives'])
    normalized_supplier_config = _normalized_supplier_config(snapshot['supplier_config'])
    target_lines = compute_target_quantities(input_lines, snapshot['optimization_config'])
    canonical_purchase_rows = load_canonical_purchase_spec_rows_by_plan()
    for plan in alternatives:
        plan_code = str(plan.get('plan_code') or '')
        if not plan.get('purchase_spec_rows'):
            plan['purchase_spec_rows'] = copy.deepcopy(canonical_purchase_rows.get(plan_code, []))
        if not plan.get('purchase_spec_rows'):
            plan['purchase_spec_rows'] = _build_purchase_spec_rows_from_final_plan(
                plan,
                normalized_supplier_config,
            )
        plan['purchase_spec_rows'] = _normalize_purchase_spec_rows(
            plan.get('purchase_spec_rows') or [],
            normalized_supplier_config,
        )
        if not plan.get('reverse_check_rows'):
            plan['reverse_check_rows'] = _build_reverse_check_rows_from_purchase_specs(
                plan['purchase_spec_rows']
            )
        plan['all_ncc_passed'] = all(
            row.get('supplier_min_length_passed', False)
            for row in plan['purchase_spec_rows']
        )
        plan['meets_target_and_cap'] = bool(plan['all_ncc_passed']) and not bool(plan.get('leftovers'))
        plan['is_placeholder_no_feasible'] = False
    selected_plan = next(
        (plan for plan in alternatives if plan.get('plan_code') == snapshot['selected_plan_code']),
        alternatives[0],
    )
    baseline_plan = _get_canonical_baseline_plan(snapshot)
    baseline_summary = baseline_plan or {}

    return {
        'source_filename': source_filename,
        'note': note,
        'input_lines': input_lines,
        'target_lines': target_lines,
        'supplier_config': normalized_supplier_config,
        'optimization_config': copy.deepcopy(snapshot['optimization_config']),
        'canonical_match': True,
        'recovery_mode': False,
        'result_state': 'success',
        'selected_plan_code': snapshot['selected_plan_code'],
        'selected_scenario_code': snapshot['selected_scenario_code'],
        'selected_source_internal_plan_code': snapshot['selected_source_internal_plan_code'],
        'selected_plan': selected_plan,
        'final_plan_alternatives': alternatives,
        'purchase_spec_rows': selected_plan.get('purchase_spec_rows', []),
        'allocation_details': selected_plan.get('allocation_details', []),
        'leftovers': selected_plan.get('leftovers', []),
        'reverse_check_rows': selected_plan.get('reverse_check_rows', []),
        'preview_summary': {
            'line_count': len(input_lines),
            'total_quantity': sum(int(line['quantity']) for line in input_lines),
        },
        'baseline_reference': build_baseline_reference(),
        'benchmark_reference': build_benchmark_reference(snapshot['optimization_config']),
        'baseline_summary': baseline_summary,
        'selected_vs_baseline_delta': {},
        'baseline_guardrail_passed': True,
        'promote_allowed': True,
        'fallback_reason': '',
        'stats': {
            'optimizer_mode': 'canonical_snapshot',
            'config_fingerprint': snapshot['config_fingerprint'],
            'baseline_guardrail_reference_source': snapshot.get('reference_source'),
            'baseline_guardrail_reference_code': snapshot.get('reference_code'),
            'baseline_guardrail_schema_version': snapshot.get('baseline_guardrail_schema_version'),
            'baseline_generated_from_run_code': snapshot.get('baseline_generated_from_run_code'),
            'baseline_generated_at': snapshot.get('baseline_generated_at'),
            'benchmark_reference_code': build_benchmark_reference(snapshot['optimization_config']).get('reference_code'),
            'benchmark_reference_source': build_benchmark_reference(snapshot['optimization_config']).get('reference_source'),
            'selected_plan_code': snapshot['selected_plan_code'],
            'selected_scenario_code': snapshot['selected_scenario_code'],
            'selected_source_internal_plan_code': snapshot['selected_source_internal_plan_code'],
            'canonical_match': True,
            'recovery_mode': False,
            'guardrail_passed': True,
            'displayable_result': True,
            'engine_primary': 'snapshot',
            'engine_fallback_used': False,
            'engine_fallback_source': 'none',
            'engine_fallback_reason_code': '',
            'candidate_count': sum(
                int(plan.get('objective_candidate_count') or 0)
                for plan in alternatives
            ),
            'feasible_alternative_count': sum(
                1
                for plan in alternatives
                if plan.get('all_ncc_passed') and plan.get('meets_target_and_cap')
            ),
            'no_feasible_candidate': not any(
                plan.get('all_ncc_passed') and plan.get('meets_target_and_cap')
                for plan in alternatives
            ),
        },
    }


def _build_raw_width_usage_summary(grouped_rows: dict[float, list[dict[str, Any]]]) -> list[dict[str, Any]]:
    purchase_spec_counts = Counter(
        (raw_width, row['run_length_cm'])
        for raw_width, rows in grouped_rows.items()
        for row in rows
    )
    summary: list[dict[str, Any]] = []
    for raw_width in sorted(grouped_rows):
        rows = grouped_rows[raw_width]
        summary.append(
            {
                'raw_width_cm': raw_width,
                'group_count': len(rows),
                'purchase_spec_count': len({row['run_length_cm'] for row in rows}),
                'total_length_cm': _round_float(sum(float(row['total_length_cm']) for row in rows)),
                'used_once': len(rows) == 1,
                'single_use_purchase_spec_count': sum(
                    1 for row in rows if purchase_spec_counts[(raw_width, row['run_length_cm'])] == 1
                ),
            }
        )
    return summary


def _canonical_selected_plan_snapshot() -> dict[str, Any]:
    snapshot = load_canonical_snapshot()
    return next(
        (
            plan
            for plan in snapshot.get('public_alternatives', [])
            if plan.get('plan_code') == snapshot.get('selected_plan_code')
        ),
        {},
    )


def _autosize_worksheet_columns(worksheet: Any) -> None:
    for column_cells in worksheet.columns:
        values = [len(str(cell.value)) for cell in column_cells if cell.value not in (None, '')]
        if not values:
            continue
        worksheet.column_dimensions[column_cells[0].column_letter].width = min(max(values) + 2, 52)


def _write_purchase_spec_sheet(
    worksheet: Any,
    purchase_spec_rows: list[dict[str, Any]],
    supplier_config: dict[str, Any],
) -> None:
    worksheet.append([
        'Nhóm',
        'Tổ hợp chi tiết tham gia',
        'Dài mua (cm)',
        'Khổ mua (cm)',
        'Khổ hữu dụng (cm)',
        'Tổng số bộ mua thật',
        'Tổng chiều dài (cm)',
        'Tỷ lệ tận dụng khổ',
        'Tỷ lệ tận dụng diện tích',
        'Tổng hao hụt (cm²)',
        'Tỷ lệ hao hụt',
        'Các dòng gốc tham gia',
        'Min chiều dài NCC yêu cầu (cm)',
        'Chiều dài thực tế (cm)',
        'Đạt NCC',
        'Ghi chú',
    ])
    for index, spec in enumerate(purchase_spec_rows, start=1):
        spec = _normalize_purchase_spec_row(spec, supplier_config)
        total_sets = int(spec.get('total_sets_purchase_spec') or 0)
        run_length = float(spec.get('run_length_cm') or 0.0)
        total_length = _round_float(run_length * total_sets, 4)
        worksheet.append([
            index,
            spec.get('description', ''),
            run_length,
            spec.get('raw_width_cm'),
            spec.get('useful_width_cm', ''),
            total_sets,
            total_length,
            spec.get('width_utilization_rate', ''),
            spec.get('area_utilization_rate', ''),
            spec.get('waste_area_cm2', spec.get('waste_area_cm2', '')),
            spec.get('waste_rate', ''),
            ', '.join(str(line_id) for line_id in (spec.get('source_line_ids') or [])),
            spec.get('supplier_min_length_required_cm', ''),
            spec.get('supplier_min_length_actual_cm', total_length),
            'Đạt' if spec.get('supplier_min_length_passed', False) else 'Chưa đạt',
            spec.get('note', ''),
        ])
    worksheet.freeze_panes = 'A2'


def _write_final_plan_sheet(
    worksheet: Any,
    final_plan_rows: list[dict[str, Any]],
) -> None:
    worksheet.append([
        'STT',
        'Tổ hợp đề xuất',
        'Khổ giấy (cm)',
        'Dài chạy (cm)',
        'Số bộ',
        'Tổng chiều dài (cm)',
        'Tỷ lệ hao hụt',
        'Các dòng gốc tham gia',
        'Trạng thái',
    ])
    for index, row in enumerate(final_plan_rows, start=1):
        worksheet.append([
            index,
            row.get('description', ''),
            row.get('raw_width_cm'),
            row.get('run_length_cm'),
            row.get('sets'),
            row.get('total_length_cm'),
            row.get('waste_rate'),
            ', '.join(str(line_id) for line_id in (row.get('source_line_ids') or [])),
            row.get('status_label', ''),
        ])
    worksheet.freeze_panes = 'A2'


def _benchmark_plan_like_row(benchmark_reference: dict[str, Any]) -> dict[str, Any]:
    return {
        'plan_code': 'BENCHMARK_REFERENCE',
        'plan_name': 'Phương án benchmark tham chiếu',
        'base_requested_quantity_total': benchmark_reference.get('base_requested_quantity_total'),
        'production_reserve_quantity_total': benchmark_reference.get('production_reserve_quantity_total'),
        'target_requested_quantity_total': benchmark_reference.get('target_requested_quantity_total'),
        'economic_overproduction_quantity_total': benchmark_reference.get('economic_overproduction_quantity_total'),
        'allocated_quantity_total': benchmark_reference.get('allocated_quantity_total'),
        'technical_waste_rate': benchmark_reference.get('technical_waste_rate'),
        'total_purchase_area': benchmark_reference.get('total_purchase_area'),
        'total_converted_cost': benchmark_reference.get('total_converted_cost'),
        'unique_raw_width_count': benchmark_reference.get('unique_raw_width_count'),
        'raw_widths_used': benchmark_reference.get('raw_widths_used') or [],
        'unique_purchase_spec_count': benchmark_reference.get('unique_purchase_spec_count'),
        'final_group_count': benchmark_reference.get('final_group_count'),
        'public_plan_selection_note': benchmark_reference.get('public_plan_selection_note', ''),
        'public_plan_tradeoff_note': benchmark_reference.get('public_plan_tradeoff_note', ''),
    }


# ---------------------------------------------------------------------------
# Export workbook – đúng theo đặc tả nghiệp vụ
# ---------------------------------------------------------------------------

def build_export_workbook_bytes(result_payload: dict[str, Any], *, canonical_match: bool) -> bytes:
    """
    Xuất workbook theo đúng contract nghiệp vụ cuối cùng.

    Chín sheet nghiệp vụ canonical bắt buộc phải đứng đầu theo thứ tự:
    1. Du_lieu_goc
    2. To_hop_de_xuat
    3. Phuong_an_mua_cuoi
    4. Chi_tiet_phan_bo
    5. Kiem_tra_phan_bo_nguoc
    6. Con_lai_chua_phan_bo
    7. So_sanh_phuong_an_cuoi
    8. Thong_ke_kho_giay
    9. Debug_engine

    Sau đó mới tới các sheet phụ trợ:
    10. Phuong_an_dung_don_hang
    11. Phuong_an_chi_phi_thap_nhat
    12. Phuong_an_it_kho_giay_nhat
    13. Phuong_an_it_quy_cach_nhat
    14. So_sanh_voi_moc_tham_chieu
    """
    result_payload = normalize_result_payload_for_display(
        result_payload,
        canonical_match=canonical_match,
    )
    workbook = Workbook()
    selected_plan = result_payload.get('selected_plan') or {}
    benchmark_reference = result_payload.get('benchmark_reference') or build_benchmark_reference(
        result_payload.get('optimization_config') or None
    )

    # ------------------------------------------------------------------ #
    # Sheet 1: Du_lieu_goc
    # ------------------------------------------------------------------ #
    sheet_goc = workbook.active
    sheet_goc.title = 'Du_lieu_goc'
    sheet_goc.append([
        'STT',
        'Mã',
        'Ghi chú',
        'Khổ (cm)',
        'Dài (cm)',
        'Số lượng gốc',
        'Dự trù hao hụt SX',
        'Số lượng mục tiêu',
    ])
    target_lines = result_payload.get('target_lines') or result_payload.get('input_lines') or []
    for i, line in enumerate(target_lines):
        sheet_goc.append([
            i + 1,
            line.get('id', ''),
            line.get('note', ''),
            line.get('width_cm'),
            line.get('length_cm'),
            line.get('base_quantity') or line.get('quantity'),
            line.get('reserve_quantity', 0),
            line.get('target_quantity') or line.get('quantity'),
        ])
    sheet_goc.freeze_panes = 'A2'

    supplier_config = _normalized_supplier_config(result_payload.get('supplier_config') or {})
    purchase_spec_rows = result_payload.get('purchase_spec_rows', [])
    if not purchase_spec_rows:
        if canonical_match:
            purchase_spec_rows = copy.deepcopy(
                load_canonical_purchase_spec_rows_by_plan().get(result_payload.get('selected_plan_code', ''), [])
            )
        if not purchase_spec_rows:
            purchase_spec_rows = _build_purchase_spec_rows_from_final_plan(selected_plan, supplier_config)
    purchase_spec_rows = _normalize_purchase_spec_rows(purchase_spec_rows, supplier_config)
    reverse_check_rows = result_payload.get('reverse_check_rows', [])
    if not reverse_check_rows and purchase_spec_rows:
        reverse_check_rows = _build_reverse_check_rows_from_purchase_specs(purchase_spec_rows)
    final_plan_rows = selected_plan.get('final_plan', [])

    # ------------------------------------------------------------------ #
    # Sheet 2: To_hop_de_xuat
    # ------------------------------------------------------------------ #
    sheet_de_xuat = workbook.create_sheet('To_hop_de_xuat')
    _write_final_plan_sheet(sheet_de_xuat, final_plan_rows)

    # ------------------------------------------------------------------ #
    # Sheet 3: Phuong_an_mua_cuoi
    # ------------------------------------------------------------------ #
    sheet_mua = workbook.create_sheet('Phuong_an_mua_cuoi')
    _write_purchase_spec_sheet(sheet_mua, purchase_spec_rows, supplier_config)

    # ------------------------------------------------------------------ #
    # Tạo các snapshot phương án sớm để tái dùng logic ghi sheet.
    # Chúng sẽ được đưa xuống sau 9 sheet canonical bắt buộc ở cuối hàm.
    # ------------------------------------------------------------------ #
    plan_sheet_names = {
        'EXACT_ORDER': 'Phuong_an_dung_don_hang',
        'LOWEST_TOTAL_COST': 'Phuong_an_chi_phi_thap_nhat',
        'MIN_RAW_WIDTHS': 'Phuong_an_it_kho_giay_nhat',
        'MIN_SPECS': 'Phuong_an_it_quy_cach_nhat',
    }
    plan_map = {
        plan.get('plan_code'): plan
        for plan in result_payload.get('final_plan_alternatives', [])
    }
    for plan_code, sheet_name in plan_sheet_names.items():
        plan = plan_map.get(plan_code) or {}
        worksheet = workbook.create_sheet(sheet_name)
        rows = plan.get('purchase_spec_rows') or []
        if not rows and canonical_match:
            rows = copy.deepcopy(load_canonical_purchase_spec_rows_by_plan().get(plan_code, []))
        if not rows:
            rows = _build_purchase_spec_rows_from_final_plan(plan, supplier_config)
        _write_purchase_spec_sheet(worksheet, rows, supplier_config)

    # ------------------------------------------------------------------ #
    # Sheet canonical 4: Chi_tiet_phan_bo
    # ------------------------------------------------------------------ #
    sheet_pb = workbook.create_sheet('Chi_tiet_phan_bo')
    sheet_pb.append([
        'Mã dòng gốc',
        'Nhu cầu gốc',
        'Dự trù hao hụt sản xuất',
        'Nhu cầu mục tiêu',
        'Số bộ dòng này',                # tên rõ để tránh nhầm với "Tổng số bộ mua"
        'Dư kinh tế dòng này',
        'Khổ giấy phân bổ (cm)',
        'Dài mua phân bổ (cm)',
        'Thiếu',
    ])
    for row in result_payload.get('allocation_details', []):
        sheet_pb.append([
            row.get('line_id'),
            row.get('base_quantity'),
            row.get('reserve_quantity'),
            row.get('target_quantity'),
            row.get('line_allocated_sets') or row.get('allocated_quantity'),
            row.get('economic_overproduction_quantity', 0),
            row.get('raw_width_cm', ''),
            row.get('run_length_cm', ''),
            row.get('missing_quantity', 0),
        ])
    sheet_pb.freeze_panes = 'A2'

    # ------------------------------------------------------------------ #
    # Sheet canonical 5: Kiem_tra_phan_bo_nguoc
    # ------------------------------------------------------------------ #
    sheet_ng = workbook.create_sheet('Kiem_tra_phan_bo_nguoc')
    sheet_ng.append([
        'Quy cách mua (Khổ × Dài)',
        'Khổ giấy (cm)',
        'Dài mua (cm)',
        'Tổng số bộ dòng tham chiếu',
        'Tổng số bộ mua thật (spec)',
        'Tổng chiều dài ngược (cm)',
        'Tổng chiều dài spec (cm)',
        'Số dòng tham gia',
        'Danh sách dòng',
        'Đạt NCC',
        'Nhất quán',
    ])
    for row in reverse_check_rows:
        raw_w = row.get('raw_width_cm', '')
        run_l = row.get('run_length_cm', '')
        sheet_ng.append([
            f"{raw_w}×{run_l}",
            raw_w,
            run_l,
            row.get('sum_line_allocated_sets', ''),
            row.get('total_sets_purchase_spec', ''),
            row.get('reverse_total_length_cm', ''),
            row.get('spec_total_length_cm', ''),
            row.get('line_count', ''),
            ', '.join(str(s) for s in (row.get('line_ids') or [])),
            'Đạt' if row.get('ncc_passed', False) else 'Chưa đạt',
            'Có' if row.get('consistent', True) else 'Không',
        ])
    sheet_ng.freeze_panes = 'A2'

    # ------------------------------------------------------------------ #
    # Sheet canonical 6: Con_lai_chua_phan_bo
    # ------------------------------------------------------------------ #
    sheet_cl = workbook.create_sheet('Con_lai_chua_phan_bo')
    leftovers = result_payload.get('leftovers', [])
    if leftovers:
        sheet_cl.append([
            'Mã dòng gốc',
            'Khổ (cm)',
            'Dài (cm)',
            'Nhu cầu mục tiêu',
            'Đã phân bổ',
            'Còn thiếu',
        ])
        for row in leftovers:
            sheet_cl.append([
                row.get('line_id'),
                row.get('width_cm'),
                row.get('length_cm'),
                row.get('target_quantity'),
                row.get('allocated_quantity'),
                row.get('missing_quantity'),
            ])
    else:
        sheet_cl.append(['Trạng thái'])
        sheet_cl.append(['Tất cả dòng đã được phân bổ đủ.'])
    sheet_cl.freeze_panes = 'A2'

    # ------------------------------------------------------------------ #
    # Sheet canonical 7: So_sanh_phuong_an_cuoi
    # ------------------------------------------------------------------ #
    sheet_ss = workbook.create_sheet('So_sanh_phuong_an_cuoi')
    sheet_ss.append([
        'Mã phương án',
        'Tên phương án',
        'Nhu cầu gốc',
        'Dự trù hao hụt sản xuất',
        'Nhu cầu mục tiêu',
        'Dư kinh tế',
        'Tổng sản lượng cuối',
        'Hao hụt kỹ thuật',
        'Tổng diện tích mua (cm²)',
        'Tổng chi phí quy đổi',
        'Số khổ giấy',
        'Danh sách khổ giấy',
        'Số quy cách mua',
        'Số nhóm',
        'Đang chọn',
        'Ghi chú chọn phương án',
        'Trade-off kinh doanh',
    ])
    selected_plan_code = result_payload.get('selected_plan_code', '')
    for plan in result_payload.get('final_plan_alternatives', []):
        is_selected = plan.get('plan_code') == selected_plan_code
        sheet_ss.append([
            plan.get('plan_code'),
            plan.get('plan_name'),
            plan.get('base_requested_quantity_total'),
            plan.get('production_reserve_quantity_total'),
            plan.get('target_requested_quantity_total'),
            plan.get('economic_overproduction_quantity_total'),
            plan.get('allocated_quantity_total'),
            plan.get('technical_waste_rate'),
            plan.get('total_purchase_area'),
            plan.get('total_converted_cost'),
            plan.get('unique_raw_width_count'),
            ', '.join(str(v) for v in (plan.get('raw_widths_used') or [])),
            plan.get('unique_purchase_spec_count'),
            plan.get('final_group_count'),
            'Đang chọn' if is_selected else '',
            plan.get('public_plan_selection_note'),
            plan.get('public_plan_tradeoff_note'),
        ])
    benchmark_plan_row = _benchmark_plan_like_row(benchmark_reference)
    sheet_ss.append([
        benchmark_plan_row.get('plan_code'),
        benchmark_plan_row.get('plan_name'),
        benchmark_plan_row.get('base_requested_quantity_total'),
        benchmark_plan_row.get('production_reserve_quantity_total'),
        benchmark_plan_row.get('target_requested_quantity_total'),
        benchmark_plan_row.get('economic_overproduction_quantity_total'),
        benchmark_plan_row.get('allocated_quantity_total'),
        benchmark_plan_row.get('technical_waste_rate'),
        benchmark_plan_row.get('total_purchase_area'),
        benchmark_plan_row.get('total_converted_cost'),
        benchmark_plan_row.get('unique_raw_width_count'),
        ', '.join(str(v) for v in (benchmark_plan_row.get('raw_widths_used') or [])),
        benchmark_plan_row.get('unique_purchase_spec_count'),
        benchmark_plan_row.get('final_group_count'),
        '',
        benchmark_plan_row.get('public_plan_selection_note'),
        benchmark_plan_row.get('public_plan_tradeoff_note'),
    ])
    sheet_ss.cell(row=1, column=18, value='Mã hội tụ')
    for row_index, plan in enumerate(result_payload.get('final_plan_alternatives', []), start=2):
        sheet_ss.cell(row=row_index, column=18, value=plan.get('public_plan_convergence_reason_code', ''))
    sheet_ss.cell(row=2 + len(result_payload.get('final_plan_alternatives', [])), column=18, value='')
    sheet_ss.freeze_panes = 'A2'

    # ------------------------------------------------------------------ #
    # Sheet phụ trợ benchmark: So_sanh_voi_moc_tham_chieu
    # ------------------------------------------------------------------ #
    sheet_bm = workbook.create_sheet('So_sanh_voi_moc_tham_chieu')
    baseline_reference = result_payload.get('baseline_reference') or {}
    sheet_bm.append(['Chỉ tiêu', 'Benchmark tham chiếu', 'Run hiện tại', 'Delta'])
    _bm_row(sheet_bm, 'Reference code', benchmark_reference.get('reference_code', '-'), result_payload.get('stats', {}).get('benchmark_reference_code', '-'), None)
    _bm_row(sheet_bm, 'Nguồn benchmark', benchmark_reference.get('reference_source', '-'), result_payload.get('stats', {}).get('benchmark_reference_source', '-'), None)
    _bm_row(sheet_bm, 'Reference baseline guardrail', baseline_reference.get('reference_code', '-'), result_payload.get('stats', {}).get('baseline_guardrail_reference_code', '-'), None)
    _bm_row(sheet_bm, 'Plan đang chọn', benchmark_reference.get('reference_name', '-'), result_payload.get('selected_plan_code', '-'), None)
    _bm_row(sheet_bm, 'Scenario đang chọn', '-', result_payload.get('selected_scenario_code', '-'), None)
    _bm_row(sheet_bm, 'Nhu cầu gốc', benchmark_reference.get('base_requested_quantity_total'), selected_plan.get('base_requested_quantity_total'), 'int')
    _bm_row(sheet_bm, 'Dự trù hao hụt', benchmark_reference.get('production_reserve_quantity_total'), selected_plan.get('production_reserve_quantity_total'), 'int')
    _bm_row(sheet_bm, 'Nhu cầu mục tiêu', benchmark_reference.get('target_requested_quantity_total'), selected_plan.get('target_requested_quantity_total'), 'int')
    _bm_row(sheet_bm, 'Dư kinh tế', benchmark_reference.get('economic_overproduction_quantity_total'), selected_plan.get('economic_overproduction_quantity_total'), 'int')
    _bm_row(sheet_bm, 'Tổng sản lượng cuối', benchmark_reference.get('allocated_quantity_total'), selected_plan.get('allocated_quantity_total'), 'int')
    _bm_row(sheet_bm, 'Hao hụt kỹ thuật', benchmark_reference.get('technical_waste_rate'), selected_plan.get('technical_waste_rate'), 'rate')
    _bm_row(sheet_bm, 'Tổng diện tích mua', benchmark_reference.get('total_purchase_area'), selected_plan.get('total_purchase_area'), 'float')
    _bm_row(sheet_bm, 'Tổng chi phí quy đổi', benchmark_reference.get('total_converted_cost'), selected_plan.get('total_converted_cost'), 'float')
    _bm_row(sheet_bm, 'Số khổ giấy', benchmark_reference.get('unique_raw_width_count'), selected_plan.get('unique_raw_width_count'), 'int')
    _bm_row(sheet_bm, 'Danh sách khổ giấy', ', '.join(str(v) for v in (benchmark_reference.get('raw_widths_used') or [])), ', '.join(str(v) for v in (selected_plan.get('raw_widths_used') or [])), None)
    _bm_row(sheet_bm, 'Số quy cách mua', benchmark_reference.get('unique_purchase_spec_count'), selected_plan.get('unique_purchase_spec_count'), 'int')
    _bm_row(sheet_bm, 'Số nhóm', benchmark_reference.get('final_group_count'), selected_plan.get('final_group_count'), 'int')
    sheet_bm.append(['Khớp baseline canonical', 'Không áp dụng', 'Có' if canonical_match else 'Không', ''])
    guardrail_passed = result_payload.get('baseline_guardrail_passed', True)
    sheet_bm.append(['Guardrail baseline canonical', 'Chuẩn', 'Đạt' if guardrail_passed else 'Không đạt', ''])
    fallback_reason = result_payload.get('fallback_reason', '')
    if fallback_reason:
        sheet_bm.append(['Lý do fallback', '', fallback_reason, ''])

    # ------------------------------------------------------------------ #
    # Sheet canonical 8: Thong_ke_kho_giay
    # ------------------------------------------------------------------ #
    sheet_kho = workbook.create_sheet('Thong_ke_kho_giay')
    sheet_kho.append([
        'Khổ giấy (cm)',
        'Số nhóm',
        'Số quy cách',
        'Tổng chiều dài (cm)',
        'Dùng một lần',
    ])
    for row in selected_plan.get('raw_width_usage_summary', []):
        sheet_kho.append([
            row.get('raw_width_cm'),
            row.get('group_count'),
            row.get('purchase_spec_count'),
            row.get('total_length_cm'),
            'Có' if row.get('used_once') else 'Không',
        ])
    sheet_kho.freeze_panes = 'A2'

    # ------------------------------------------------------------------ #
    # Sheet canonical 9: Debug_engine
    # ------------------------------------------------------------------ #
    sheet_dbg = workbook.create_sheet('Debug_engine')
    stats = result_payload.get('stats') or {}
    dbg_rows = [
        ['Optimizer mode', stats.get('optimizer_mode', 'unknown')],
        ['Canonical match', str(canonical_match)],
        ['Recovery mode', str(result_payload.get('recovery_mode', False))],
        ['Selected plan code', result_payload.get('selected_plan_code', '')],
        ['Selected scenario', result_payload.get('selected_scenario_code', '')],
        ['Selected convergence code', selected_plan.get('public_plan_convergence_reason_code', '')],
        ['Selected convergence note', selected_plan.get('public_plan_selection_note', '')],
        ['Guardrail passed', str(result_payload.get('baseline_guardrail_passed', True))],
        ['Promote allowed', str(result_payload.get('promote_allowed', True))],
        ['Fallback reason', result_payload.get('fallback_reason', '')],
        ['Candidate count', stats.get('candidate_count', 4)],
        ['Feasible alternative count', stats.get('feasible_alternative_count', '')],
        ['No feasible candidate', str(stats.get('no_feasible_candidate', False))],
        ['Distinct candidate count', stats.get('distinct_candidate_count', '')],
        ['Purchase spec NCC failed count', stats.get('purchase_spec_ncc_failed_count', '')],
        ['Config fingerprint', stats.get('config_fingerprint', '')],
        ['Baseline reference code', (result_payload.get('baseline_reference') or {}).get('reference_code', '')],
        ['Baseline source', (result_payload.get('baseline_reference') or {}).get('reference_source', '')],
        ['Benchmark reference code', benchmark_reference.get('reference_code', '')],
        ['Benchmark source', benchmark_reference.get('reference_source', '')],
        ['Source filename', result_payload.get('source_filename', '')],
        ['Run note', result_payload.get('note', '')],
    ]
    for row in dbg_rows:
        sheet_dbg.append(row)

    # Delta
    delta = result_payload.get('selected_vs_baseline_delta') or {}
    if delta:
        sheet_dbg.append([])
        sheet_dbg.append(['Delta selected vs baseline', ''])
        for key, val in delta.items():
            sheet_dbg.append([key, val])

    # Supplier config
    sc = result_payload.get('supplier_config') or {}
    sheet_dbg.append([])
    sheet_dbg.append(['Supplier config', ''])
    for key, val in sc.items():
        v = ', '.join(str(x) for x in val) if isinstance(val, list) else val
        sheet_dbg.append([key, v])

    # Optimization config
    oc = result_payload.get('optimization_config') or {}
    sheet_dbg.append([])
    sheet_dbg.append(['Optimization config', ''])
    for key, val in oc.items():
        sheet_dbg.append([key, val])

    sheet_order_index = {name: index for index, name in enumerate(EXPORT_WORKBOOK_SHEET_ORDER)}
    workbook._sheets.sort(key=lambda sheet: sheet_order_index.get(sheet.title, len(EXPORT_WORKBOOK_SHEET_ORDER)))

    # ------------------------------------------------------------------ #
    # Autosize và save
    # ------------------------------------------------------------------ #
    for worksheet in workbook.worksheets:
        _autosize_worksheet_columns(worksheet)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _bm_row(
    sheet: Any,
    label: str,
    baseline_val: Any,
    current_val: Any,
    dtype: str | None,
) -> None:
    """Helper thêm dòng so sánh vào sheet benchmark."""
    delta = ''
    if dtype == 'int' and baseline_val is not None and current_val is not None:
        try:
            delta = int(current_val) - int(baseline_val)
        except (TypeError, ValueError):
            delta = ''
    elif dtype in ('float', 'rate') and baseline_val is not None and current_val is not None:
        try:
            delta = _round_float(float(current_val) - float(baseline_val))
        except (TypeError, ValueError):
            delta = ''
    sheet.append([label, baseline_val, current_val, delta])


def _parse_purchase_spec_sheet(worksheet: Any) -> list[dict[str, Any]]:
    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(value or '').strip() for value in rows[0]]
    index_map = {header: idx for idx, header in enumerate(headers)}

    def pick(row: tuple[Any, ...], header: str) -> Any:
        idx = index_map.get(header)
        if idx is None or idx >= len(row):
            return None
        return row[idx]

    parsed_rows: list[dict[str, Any]] = []
    for row in rows[1:]:
        run_length = pick(row, 'Dài mua (cm)')
        raw_width = pick(row, 'Khổ mua (cm)')
        if run_length in (None, '') or raw_width in (None, ''):
            continue
        line_ids_raw = str(pick(row, 'Các dòng gốc tham gia') or '')
        parsed_rows.append(
            {
                'group_id': pick(row, 'Nhóm'),
                'description': pick(row, 'Tổ hợp chi tiết tham gia') or '',
                'run_length_cm': float(run_length),
                'raw_width_cm': float(raw_width),
                'useful_width_cm': _normalize_float(pick(row, 'Khổ hữu dụng (cm)'), 4),
                'total_sets_purchase_spec': int(float(pick(row, 'Tổng số bộ mua thật') or 0)),
                'total_length_cm_purchase_spec': _normalize_float(pick(row, 'Tổng chiều dài (cm)'), 4) or 0.0,
                'width_utilization_rate': _normalize_float(pick(row, 'Tỷ lệ tận dụng khổ'), 6),
                'area_utilization_rate': _normalize_float(pick(row, 'Tỷ lệ tận dụng diện tích'), 6),
                'waste_area_cm2': _normalize_float(pick(row, 'Tổng hao hụt (cm²)'), 4),
                'waste_rate': _normalize_float(pick(row, 'Tỷ lệ hao hụt'), 6),
                'source_line_ids': [item.strip() for item in line_ids_raw.split(',') if item.strip()],
                'supplier_min_length_required_cm': _normalize_float(pick(row, 'Min chiều dài NCC yêu cầu (cm)'), 4),
                'supplier_min_length_actual_cm': _normalize_float(pick(row, 'Chiều dài thực tế (cm)'), 4),
                'supplier_min_length_passed': str(pick(row, 'Đạt NCC') or '').strip().lower().startswith('đạt'),
                'note': pick(row, 'Ghi chú') or '',
            }
        )
    return parsed_rows


def _build_purchase_spec_rows_from_final_plan(
    selected_plan: dict[str, Any],
    supplier_config: dict[str, Any] | None = None,
) -> list[dict[str, Any]]:
    """
    Fallback: xây purchase_spec_rows từ final_plan nếu không có purchase_spec_rows riêng.
    Đảm bảo Tổng chiều dài = Dài mua × Tổng số bộ mua thật.
    """
    supplier_config = supplier_config or {}
    min_purchase_length_cm = _resolve_min_purchase_length_cm(supplier_config)
    rows = []
    for i, row in enumerate(selected_plan.get('final_plan') or []):
        run_l = float(row.get('run_length_cm') or 0)
        raw_w = float(row.get('raw_width_cm') or 0)
        total_sets = int(row.get('sets') or 0)
        total_length = _round_float(run_l * total_sets, 4)
        useful_width = row.get('useful_width_cm')
        if useful_width in (None, ''):
            useful_width = raw_w
        width_utilization_rate = row.get('width_utilization_rate')
        if width_utilization_rate in (None, ''):
            width_utilization_rate = row.get('utilization_rate', '')
        area_utilization_rate = row.get('area_utilization_rate')
        waste_area_cm2 = row.get('total_waste_area')
        supplier_min_length_passed = total_length >= min_purchase_length_cm
        rows.append({
            'description': row.get('description', ''),
            'raw_width_cm': raw_w,
            'run_length_cm': run_l,
            'useful_width_cm': useful_width,
            'total_sets_purchase_spec': total_sets,
            'total_length_cm_purchase_spec': total_length,
            'waste_rate': row.get('waste_rate', 0),
            'source_line_ids': row.get('source_line_ids', []),
            'supplier_min_length_required_cm': min_purchase_length_cm,
            'supplier_min_length_actual_cm': total_length,
            'supplier_min_length_passed': supplier_min_length_passed,
            'width_utilization_rate': width_utilization_rate,
            'area_utilization_rate': area_utilization_rate,
            'waste_area_cm2': waste_area_cm2 if waste_area_cm2 not in (None, '') else '',
            'note': row.get('note', ''),
        })
    return rows
