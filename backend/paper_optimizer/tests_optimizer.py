import copy
import json
from io import BytesIO, StringIO

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import TestCase
from openpyxl import load_workbook
from rest_framework.test import APIClient
from unittest.mock import patch
from core.models import Permission, Role

from paper_optimizer.models import PaperOptimizerRun, PaperOptimizerSupplierTemplate
from paper_optimizer.optimizer import build_four_public_plans
from paper_optimizer.optimizer import _select_best_plan as select_best_plan_legacy
from paper_optimizer.optimizer_v2 import _pick_plan as pick_plan_v2
from paper_optimizer.optimizer_v2 import _select_best_plan as select_best_plan_v2
from paper_optimizer.regression_corpus import (
    REGRESSION_CASE_ORDER,
    build_non_viable_public_alternatives,
    build_viable_legacy_fallback_alternatives,
    run_regression_case,
)
from paper_optimizer.services import (
    PRIMARY_EXPORT_SHEET_ORDER,
    SUPPLEMENTAL_EXPORT_SHEET_ORDER,
    load_canonical_snapshot,
)


User = get_user_model()


class PaperOptimizerRecoveryApiTest(TestCase):
    def setUp(self):
        self.client = APIClient()
        production_permission, _ = Permission.objects.get_or_create(
            resource='PRODUCTION',
            action='MANAGE',
            defaults={
                'code': 'PRODUCTION_MANAGE',
                'name': 'Manage production module',
                'description': 'Paper optimizer recovery test permission.',
            },
        )
        production_role = Role.objects.create(code='PAPER_OPT_TEST', name='Paper Optimizer Test')
        production_role.permissions.add(production_permission)
        self.user = User.objects.create_user(
            username='paper_optimizer_user',
            password='Demo123!',
        )
        self.user.roles.add(production_role)
        self.client.force_authenticate(self.user)
        self.snapshot = load_canonical_snapshot()

    def assert_export_workbook_sheet_order(self, workbook):
        self.assertEqual(
            workbook.sheetnames[:len(PRIMARY_EXPORT_SHEET_ORDER)],
            list(PRIMARY_EXPORT_SHEET_ORDER),
        )
        start = len(PRIMARY_EXPORT_SHEET_ORDER)
        end = start + len(SUPPLEMENTAL_EXPORT_SHEET_ORDER)
        self.assertEqual(
            workbook.sheetnames[start:end],
            list(SUPPLEMENTAL_EXPORT_SHEET_ORDER),
        )

    def test_defaults_and_templates_bootstrap(self):
        defaults_response = self.client.get('/api/production/paper-optimizer/runs/defaults/')
        self.assertEqual(defaults_response.status_code, 200)
        body = defaults_response.json()
        self.assertEqual(body['baseline_reference']['reference_code'], 'baseline-canonical-v5000-v1')
        self.assertEqual(body['benchmark_reference']['reference_code'], 'benchmark-manual-v1')
        self.assertAlmostEqual(body['benchmark_reference']['total_converted_cost'], 2.67223922, places=8)
        self.assertEqual(len(body['input_lines']), 7)

        templates_response = self.client.get('/api/production/paper-optimizer/templates/')
        self.assertEqual(templates_response.status_code, 200)
        templates = templates_response.json()
        self.assertGreaterEqual(len(templates), 1)
        self.assertTrue(PaperOptimizerSupplierTemplate.objects.filter(is_system_default=True).exists())

    def test_download_import_template_returns_csv_fixture(self):
        response = self.client.get('/api/production/paper-optimizer/runs/download_import_template/')
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response['Content-Type'], 'text/csv; charset=utf-8')
        self.assertIn('paper_optimizer_import_template.csv', response['Content-Disposition'])
        body = response.content.decode('utf-8-sig').splitlines()
        self.assertEqual(body[0], 'width_cm,length_cm,quantity,note')
        self.assertGreaterEqual(len(body), 3)

    def test_optimize_rejects_min_purchase_length_below_5000(self):
        supplier_config = copy.deepcopy(self.snapshot['supplier_config'])
        supplier_config['min_purchase_length_cm'] = 4500
        response = self.client.post(
            '/api/production/paper-optimizer/runs/optimize/',
            {
                'source_filename': 'below-ncc-threshold.xlsx',
                'note': 'invalid supplier threshold',
                'input_lines': self.snapshot['input_lines'],
                'supplier_config': supplier_config,
                'optimization_config': self.snapshot['optimization_config'],
            },
            format='json',
        )
        self.assertEqual(response.status_code, 400, response.content)
        self.assertIn('5.000 cm', json.dumps(response.json(), ensure_ascii=False))

    def test_templates_list_hides_other_users_items_but_keeps_system_default(self):
        other_user = User.objects.create_user(username='paper_optimizer_other', password='Demo123!')
        own_template = PaperOptimizerSupplierTemplate.objects.create(
            name='Own template',
            note='own',
            supplier_config=self.snapshot['supplier_config'],
            optimization_config=self.snapshot['optimization_config'],
            created_by=self.user,
        )
        other_template = PaperOptimizerSupplierTemplate.objects.create(
            name='Other template',
            note='other',
            supplier_config=self.snapshot['supplier_config'],
            optimization_config=self.snapshot['optimization_config'],
            created_by=other_user,
        )
        system_template = PaperOptimizerSupplierTemplate.objects.create(
            name='System template',
            note='system',
            supplier_config=self.snapshot['supplier_config'],
            optimization_config=self.snapshot['optimization_config'],
            is_system_default=True,
        )

        response = self.client.get('/api/production/paper-optimizer/templates/')
        self.assertEqual(response.status_code, 200, response.content)
        template_ids = {item['id'] for item in response.json()}
        self.assertIn(own_template.id, template_ids)
        self.assertIn(system_template.id, template_ids)
        self.assertNotIn(other_template.id, template_ids)

    def test_personal_template_can_be_updated_and_deleted(self):
        template = PaperOptimizerSupplierTemplate.objects.create(
            name='Editable template',
            note='before',
            supplier_config=self.snapshot['supplier_config'],
            optimization_config=self.snapshot['optimization_config'],
            created_by=self.user,
        )

        update_response = self.client.patch(
            f'/api/production/paper-optimizer/templates/{template.id}/',
            {
                'name': 'Updated template',
                'note': 'after',
            },
            format='json',
        )
        self.assertEqual(update_response.status_code, 200, update_response.content)
        template.refresh_from_db()
        self.assertEqual(template.name, 'Updated template')
        self.assertEqual(template.note, 'after')

        delete_response = self.client.delete(f'/api/production/paper-optimizer/templates/{template.id}/')
        self.assertEqual(delete_response.status_code, 204, delete_response.content)
        self.assertFalse(PaperOptimizerSupplierTemplate.objects.filter(id=template.id).exists())

    def test_system_template_cannot_be_updated_or_deleted(self):
        template = PaperOptimizerSupplierTemplate.objects.create(
            name='System template',
            note='system',
            supplier_config=self.snapshot['supplier_config'],
            optimization_config=self.snapshot['optimization_config'],
            is_system_default=True,
        )

        update_response = self.client.patch(
            f'/api/production/paper-optimizer/templates/{template.id}/',
            {'name': 'Blocked update'},
            format='json',
        )
        self.assertEqual(update_response.status_code, 403, update_response.content)
        self.assertIn('Template', update_response.json()['detail'])

        delete_response = self.client.delete(f'/api/production/paper-optimizer/templates/{template.id}/')
        self.assertEqual(delete_response.status_code, 403, delete_response.content)
        self.assertIn('Template', delete_response.json()['detail'])

    def test_list_runs_returns_recent_summaries(self):
        failed_run = PaperOptimizerRun.objects.create(
            code='POPT-FAILED',
            status=PaperOptimizerRun.STATUS_FAILED,
            source_filename='failed.xlsx',
            note='run loi',
            input_lines=[
                {'id': 'L1', 'quantity': 5, 'width_cm': 52, 'length_cm': 112, 'source_row_number': 1},
            ],
            supplier_config={},
            optimization_config={},
            failed_reason='Khong sinh duoc phuong an.',
            created_by=self.user,
        )
        success_run = PaperOptimizerRun.objects.create(
            code='POPT-SUCCESS',
            status=PaperOptimizerRun.STATUS_SUCCESS,
            source_filename='success.xlsx',
            note='run thanh cong',
            input_lines=[
                {'id': 'L1', 'quantity': 7, 'width_cm': 52, 'length_cm': 112, 'source_row_number': 1},
                {'id': 'L2', 'quantity': 9, 'width_cm': 84, 'length_cm': 143.5, 'source_row_number': 2},
            ],
            supplier_config={},
            optimization_config={},
            result_payload={
                'selected_plan_code': 'LOWEST_TOTAL_COST',
                'selected_scenario_code': 'ECON-02',
                'selected_plan': {'plan_name': 'Phuong an chi phi thap nhat'},
                'stats': {
                    'engine_primary': 'v2',
                    'engine_fallback_used': True,
                    'engine_fallback_source': 'legacy',
                },
            },
            created_by=self.user,
        )

        response = self.client.get('/api/production/paper-optimizer/runs/?limit=2')
        self.assertEqual(response.status_code, 200, response.content)
        body = response.json()
        self.assertEqual(body['count'], 2)
        self.assertEqual(body['limit'], 2)
        self.assertEqual(body['offset'], 0)
        self.assertEqual(len(body['results']), 2)
        self.assertEqual(body['results'][0]['id'], success_run.id)
        self.assertEqual(body['results'][0]['selected_plan_code'], 'LOWEST_TOTAL_COST')
        self.assertEqual(body['results'][0]['selected_plan_name'], 'Phuong an chi phi thap nhat')
        self.assertEqual(body['results'][0]['selected_scenario_code'], 'ECON-02')
        self.assertEqual(body['results'][0]['engine_primary'], 'v2')
        self.assertTrue(body['results'][0]['engine_fallback_used'])
        self.assertEqual(body['results'][0]['engine_fallback_source'], 'legacy')
        self.assertEqual(body['results'][0]['line_count'], 2)
        self.assertEqual(body['results'][0]['total_quantity'], 16)
        self.assertEqual(body['results'][1]['id'], failed_run.id)
        self.assertEqual(body['results'][1]['status'], PaperOptimizerRun.STATUS_FAILED)
        self.assertEqual(body['results'][1]['failed_reason'], 'Khong sinh duoc phuong an.')

    def test_list_runs_supports_filters_search_and_offset(self):
        canonical_run = PaperOptimizerRun.objects.create(
            code='POPT-CANONICAL',
            status=PaperOptimizerRun.STATUS_SUCCESS,
            source_filename='canonical.xlsx',
            note='baseline',
            input_lines=[{'id': 'L1', 'quantity': 10, 'width_cm': 52, 'length_cm': 112, 'source_row_number': 1}],
            supplier_config={},
            optimization_config={},
            result_payload={
                'selected_plan_code': 'LOWEST_TOTAL_COST',
                'selected_scenario_code': 'ECON-02',
                'selected_plan': {'plan_name': 'Baseline canonical'},
                'stats': {'engine_primary': 'snapshot', 'engine_fallback_used': False, 'engine_fallback_source': 'none'},
            },
            canonical_match=True,
            created_by=self.user,
        )
        fallback_run = PaperOptimizerRun.objects.create(
            code='POPT-FALLBACK',
            status=PaperOptimizerRun.STATUS_SUCCESS,
            source_filename='legacy.xlsx',
            note='legacy fallback',
            input_lines=[{'id': 'L2', 'quantity': 15, 'width_cm': 84, 'length_cm': 143.5, 'source_row_number': 1}],
            supplier_config={},
            optimization_config={},
            result_payload={
                'selected_plan_code': 'LOWEST_TOTAL_COST',
                'selected_scenario_code': 'ECON-02',
                'selected_plan': {'plan_name': 'Legacy fallback'},
                'stats': {'engine_primary': 'v2', 'engine_fallback_used': True, 'engine_fallback_source': 'legacy'},
            },
            canonical_match=False,
            created_by=self.user,
        )
        failed_run = PaperOptimizerRun.objects.create(
            code='POPT-FAILED-SEARCH',
            status=PaperOptimizerRun.STATUS_FAILED,
            source_filename='failed.xlsx',
            note='that bai',
            input_lines=[{'id': 'L3', 'quantity': 5, 'width_cm': 40, 'length_cm': 90, 'source_row_number': 1}],
            supplier_config={},
            optimization_config={},
            failed_reason='No feasible candidate',
            created_by=self.user,
        )

        response = self.client.get('/api/production/paper-optimizer/runs/?status=SUCCESS&engine_fallback_used=true')
        self.assertEqual(response.status_code, 200, response.content)
        body = response.json()
        self.assertEqual(body['count'], 1)
        self.assertEqual(body['results'][0]['id'], fallback_run.id)

        response = self.client.get('/api/production/paper-optimizer/runs/?canonical_match=true')
        self.assertEqual(response.status_code, 200, response.content)
        body = response.json()
        self.assertEqual(body['count'], 1)
        self.assertEqual(body['results'][0]['id'], canonical_run.id)

        response = self.client.get('/api/production/paper-optimizer/runs/?q=fallback')
        self.assertEqual(response.status_code, 200, response.content)
        body = response.json()
        self.assertEqual(body['count'], 1)
        self.assertEqual(body['results'][0]['id'], fallback_run.id)

        response = self.client.get('/api/production/paper-optimizer/runs/?limit=1&offset=1')
        self.assertEqual(response.status_code, 200, response.content)
        body = response.json()
        self.assertEqual(body['count'], 3)
        self.assertEqual(len(body['results']), 1)
        self.assertEqual(body['results'][0]['id'], fallback_run.id)
        self.assertNotEqual(body['results'][0]['id'], failed_run.id)

    def test_optimize_canonical_dataset_returns_locked_baseline(self):
        response = self.client.post(
            '/api/production/paper-optimizer/runs/optimize/',
            {
                'source_filename': 'canonical.xlsx',
                'note': 'baseline test',
                'input_lines': self.snapshot['input_lines'],
                'supplier_config': self.snapshot['supplier_config'],
                'optimization_config': self.snapshot['optimization_config'],
            },
            format='json',
        )
        self.assertEqual(response.status_code, 201, response.content)
        body = response.json()
        result = body['result_payload']
        self.assertTrue(result['canonical_match'])
        self.assertFalse(result['recovery_mode'])
        self.assertEqual(result['selected_plan_code'], 'LOWEST_TOTAL_COST')
        self.assertEqual(result['selected_scenario_code'], 'TARGET')
        self.assertEqual(result['selected_source_internal_plan_code'], 'BALANCED')
        self.assertAlmostEqual(result['selected_plan']['total_converted_cost'], 2.67223922, places=8)
        self.assertEqual(result['selected_plan']['unique_raw_width_count'], 1)
        self.assertEqual(result['selected_plan']['unique_purchase_spec_count'], 4)
        self.assertEqual(result['selected_plan']['final_group_count'], 6)
        self.assertEqual(len(result['selected_plan'].get('purchase_spec_rows') or []), 4)
        self.assertEqual(result['result_state'], 'success')
        self.assertTrue(bool(result['stats']['displayable_result']))
        self.assertFalse(bool(result['selected_plan'].get('is_placeholder_no_feasible')))
        self.assertEqual(int(result['stats']['feasible_alternative_count']), 4)
        self.assertFalse(bool(result['stats']['no_feasible_candidate']))

    @patch('paper_optimizer.views.run_optimizer_recovery', side_effect=RuntimeError('Engine no feasible plan'))
    def test_optimize_failure_creates_failed_run(self, _mock_optimizer):
        response = self.client.post(
            '/api/production/paper-optimizer/runs/optimize/',
            {
                'source_filename': 'broken.xlsx',
                'note': 'force failed run',
                'input_lines': self.snapshot['input_lines'],
                'supplier_config': self.snapshot['supplier_config'],
                'optimization_config': self.snapshot['optimization_config'],
            },
            format='json',
        )
        self.assertEqual(response.status_code, 201, response.content)
        body = response.json()
        self.assertEqual(body['status'], PaperOptimizerRun.STATUS_FAILED)
        self.assertEqual(body['source_filename'], 'broken.xlsx')
        self.assertEqual(body['note'], 'force failed run')
        self.assertIn('Engine no feasible plan', body['failed_reason'])
        self.assertEqual(body['result_payload'], {})
        self.assertEqual(len(body['preview_rows']), len(self.snapshot['input_lines']))

        run = PaperOptimizerRun.objects.get(id=body['id'])
        self.assertEqual(run.status, PaperOptimizerRun.STATUS_FAILED)
        self.assertIn('Engine no feasible plan', run.failed_reason)
        self.assertEqual(len(run.preview_rows), len(self.snapshot['input_lines']))

        export_response = self.client.get(f'/api/production/paper-optimizer/runs/{run.id}/export_excel/')
        self.assertEqual(export_response.status_code, 400)
        self.assertEqual(export_response.json()['detail'], 'Run thất bại không có workbook để tải.')

    def test_optimize_non_canonical_dataset_uses_full_optimizer(self):
        """Khi dữ liệu không khớp canonical: engine optimizer đầy đủ chạy (không phải recovery)."""
        changed_lines = [dict(item) for item in self.snapshot['input_lines']]
        changed_lines[0]['quantity'] = changed_lines[0]['quantity'] + 1
        response = self.client.post(
            '/api/production/paper-optimizer/runs/optimize/',
            {
                'source_filename': 'changed.xlsx',
                'note': 'optimizer test',
                'input_lines': changed_lines,
                'supplier_config': self.snapshot['supplier_config'],
                'optimization_config': self.snapshot['optimization_config'],
            },
            format='json',
        )
        self.assertEqual(response.status_code, 201, response.content)
        body = response.json()
        run_id = body['id']
        result = body['result_payload']
        self.assertFalse(result['canonical_match'])
        # Engine mới không phải recovery_mode thô: nó là full optimizer
        self.assertFalse(result.get('recovery_mode', False))
        self.assertIn(result['selected_plan_code'], {'EXACT_ORDER', 'LOWEST_TOTAL_COST', 'MIN_RAW_WIDTHS', 'MIN_SPECS'})
        self.assertEqual(len(result['final_plan_alternatives']), 4)
        total_cap = int(self.snapshot['optimization_config']['max_economic_overproduction_quantity_total'])
        line_cap = int(self.snapshot['optimization_config']['max_economic_overproduction_quantity_per_line'])
        feasible_alternatives = [
            alternative for alternative in result['final_plan_alternatives']
            if alternative.get('all_ncc_passed') and alternative.get('meets_target_and_cap')
        ]
        if feasible_alternatives:
            self.assertTrue(result['selected_plan']['all_ncc_passed'])
            self.assertEqual(result['result_state'], 'success')
            self.assertTrue(bool(result['stats']['displayable_result']))
            for alternative in feasible_alternatives:
                self.assertLessEqual(alternative.get('economic_overproduction_quantity_total', 0), total_cap)
                economic_by_line = {}
                for row in alternative.get('allocation_details', []):
                    economic_by_line[row['line_id']] = economic_by_line.get(row['line_id'], 0) + int(row.get('economic_overproduction_quantity') or 0)
                self.assertTrue(all(quantity <= line_cap for quantity in economic_by_line.values()))
                self.assertGreater(alternative.get('unique_purchase_spec_count', 0), 0)
                self.assertGreater(alternative.get('final_group_count', 0), 0)
        else:
            self.assertFalse(result['selected_plan']['all_ncc_passed'])
            self.assertTrue(result['stats']['no_feasible_candidate'])
            self.assertEqual(result['stats']['feasible_alternative_count'], 0)
            self.assertEqual(result['result_state'], 'no_feasible_candidate')
            self.assertFalse(bool(result['stats']['displayable_result']))
            self.assertTrue(bool(result['selected_plan'].get('is_placeholder_no_feasible')))
            for alternative in result['final_plan_alternatives']:
                self.assertEqual(alternative.get('public_plan_convergence_reason_code'), 'no_feasible_candidate')
                self.assertFalse(alternative.get('meets_target_and_cap'))

        export_response = self.client.get(f'/api/production/paper-optimizer/runs/{run_id}/export_excel/')
        if feasible_alternatives:
            self.assertEqual(export_response.status_code, 200)
            self.assertEqual(
                export_response['Content-Type'],
                'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            )
            self.assertGreater(len(export_response.content), 1024)
            workbook = load_workbook(BytesIO(export_response.content), read_only=True, data_only=True)
            self.assert_export_workbook_sheet_order(workbook)

            # Sheet names mới theo đặc tả
            self.assertIn('Du_lieu_goc', workbook.sheetnames)
            self.assertIn('To_hop_de_xuat', workbook.sheetnames)
            self.assertIn('Phuong_an_mua_cuoi', workbook.sheetnames)
            self.assertIn('Phuong_an_dung_don_hang', workbook.sheetnames)
            self.assertIn('Phuong_an_chi_phi_thap_nhat', workbook.sheetnames)
            self.assertIn('Phuong_an_it_kho_giay_nhat', workbook.sheetnames)
            self.assertIn('Phuong_an_it_quy_cach_nhat', workbook.sheetnames)
            self.assertIn('Chi_tiet_phan_bo', workbook.sheetnames)
            self.assertIn('Kiem_tra_phan_bo_nguoc', workbook.sheetnames)
            self.assertIn('Con_lai_chua_phan_bo', workbook.sheetnames)
            self.assertIn('So_sanh_phuong_an_cuoi', workbook.sheetnames)
            self.assertIn('So_sanh_voi_moc_tham_chieu', workbook.sheetnames)
            self.assertIn('Thong_ke_kho_giay', workbook.sheetnames)
            self.assertIn('Debug_engine', workbook.sheetnames)

            # Kiểm sheet Phuong_an_mua_cuoi có cột đúng
            sheet_mua = workbook['Phuong_an_mua_cuoi']
            headers = [cell.value for cell in next(sheet_mua.iter_rows(min_row=1, max_row=1))]
            self.assertIn('Tổng số bộ mua thật', headers)
            self.assertIn('Tổng chiều dài (cm)', headers)
            self.assertIn('Đạt NCC', headers)
            self.assertIn('Dài mua (cm)', headers)
            self.assertIn('Khổ mua (cm)', headers)

            compare_sheet = workbook['So_sanh_phuong_an_cuoi']
            compare_rows = list(compare_sheet.iter_rows(min_row=2, values_only=True))
            self.assertTrue(any(row[0] == 'BENCHMARK_REFERENCE' for row in compare_rows if row and row[0]))

            benchmark_sheet = workbook['So_sanh_voi_moc_tham_chieu']
            benchmark_headers = [cell.value for cell in next(benchmark_sheet.iter_rows(min_row=1, max_row=1))]
            self.assertEqual(benchmark_headers[1], 'Benchmark tham chiếu')

            # Kiểm công thức Tổng chiều dài = Dài mua × Tổng số bộ mua thật
            idx_dai = headers.index('Dài mua (cm)')
            idx_sets = headers.index('Tổng số bộ mua thật')
            idx_total = headers.index('Tổng chiều dài (cm)')
            for row in sheet_mua.iter_rows(min_row=2, values_only=True):
                if row[idx_dai] is None:
                    break
                dai = float(row[idx_dai] or 0)
                sets = int(row[idx_sets] or 0)
                total_len = float(row[idx_total] or 0)
                expected_len = round(dai * sets, 4)
                self.assertAlmostEqual(total_len, expected_len, places=2,
                    msg=f'Tổng chiều dài sai: {dai} × {sets} = {expected_len}, thực tế = {total_len}')
        else:
            self.assertEqual(export_response.status_code, 400)
            self.assertEqual(export_response.json()['detail'], 'Run chưa có phương án mua cuối hợp lệ để xuất workbook.')

    def test_retrieve_legacy_no_feasible_run_normalizes_result_payload_and_blocks_export(self):
        legacy_payload = copy.deepcopy(run_regression_case('no_feasible_candidate')['result_payload'])
        legacy_payload.pop('result_state', None)
        (legacy_payload.get('stats') or {}).pop('displayable_result', None)
        (legacy_payload.get('selected_plan') or {}).pop('is_placeholder_no_feasible', None)
        for alternative in legacy_payload.get('final_plan_alternatives') or []:
            alternative.pop('is_placeholder_no_feasible', None)

        run = PaperOptimizerRun.objects.create(
            code='POPT-LEGACY-NO-FEASIBLE',
            status=PaperOptimizerRun.STATUS_SUCCESS,
            source_filename='legacy-no-feasible.xlsx',
            note='legacy compatibility no-feasible',
            input_lines=legacy_payload.get('input_lines') or [],
            supplier_config=legacy_payload.get('supplier_config') or self.snapshot['supplier_config'],
            optimization_config=legacy_payload.get('optimization_config') or self.snapshot['optimization_config'],
            preview_rows=legacy_payload.get('input_lines') or [],
            result_payload=legacy_payload,
            recovery_mode=bool(legacy_payload.get('recovery_mode')),
            canonical_match=bool(legacy_payload.get('canonical_match')),
            created_by=self.user,
        )

        retrieve_response = self.client.get(f'/api/production/paper-optimizer/runs/{run.id}/')
        self.assertEqual(retrieve_response.status_code, 200, retrieve_response.content)
        result = retrieve_response.json()['result_payload']
        self.assertEqual(result['result_state'], 'no_feasible_candidate')
        self.assertFalse(bool(result['stats']['displayable_result']))
        self.assertTrue(bool(result['stats']['no_feasible_candidate']))
        self.assertTrue(bool(result['selected_plan']['is_placeholder_no_feasible']))
        self.assertTrue(all(bool(item['is_placeholder_no_feasible']) for item in result['final_plan_alternatives']))

        export_response = self.client.get(f'/api/production/paper-optimizer/runs/{run.id}/export_excel/')
        self.assertEqual(export_response.status_code, 400, export_response.content)
        self.assertIn('workbook', export_response.json()['detail'])

    def test_retrieve_legacy_success_run_backfills_displayable_rows(self):
        legacy_payload = copy.deepcopy(run_regression_case('non_canonical_success')['result_payload'])
        legacy_payload.pop('result_state', None)
        (legacy_payload.get('stats') or {}).pop('displayable_result', None)
        legacy_payload.pop('purchase_spec_rows', None)
        legacy_payload.pop('reverse_check_rows', None)
        (legacy_payload.get('selected_plan') or {}).pop('is_placeholder_no_feasible', None)
        for alternative in legacy_payload.get('final_plan_alternatives') or []:
            alternative.pop('is_placeholder_no_feasible', None)

        run = PaperOptimizerRun.objects.create(
            code='POPT-LEGACY-SUCCESS',
            status=PaperOptimizerRun.STATUS_SUCCESS,
            source_filename='legacy-success.xlsx',
            note='legacy compatibility success',
            input_lines=legacy_payload.get('input_lines') or [],
            supplier_config=legacy_payload.get('supplier_config') or self.snapshot['supplier_config'],
            optimization_config=legacy_payload.get('optimization_config') or self.snapshot['optimization_config'],
            preview_rows=legacy_payload.get('input_lines') or [],
            result_payload=legacy_payload,
            recovery_mode=bool(legacy_payload.get('recovery_mode')),
            canonical_match=bool(legacy_payload.get('canonical_match')),
            created_by=self.user,
        )

        retrieve_response = self.client.get(f'/api/production/paper-optimizer/runs/{run.id}/')
        self.assertEqual(retrieve_response.status_code, 200, retrieve_response.content)
        result = retrieve_response.json()['result_payload']
        self.assertEqual(result['result_state'], 'success')
        self.assertTrue(bool(result['stats']['displayable_result']))
        self.assertFalse(bool(result['selected_plan']['is_placeholder_no_feasible']))
        self.assertGreater(len(result['purchase_spec_rows']), 0)
        self.assertGreater(len(result['reverse_check_rows']), 0)

    def test_non_canonical_success_reverse_check_uses_purchase_sets(self):
        non_canonical_lines = [
            dict(self.snapshot['input_lines'][0]),
            dict(self.snapshot['input_lines'][1]),
        ]
        non_canonical_lines[0]['quantity'] = 60
        non_canonical_lines[1]['quantity'] = 40
        response = self.client.post(
            '/api/production/paper-optimizer/runs/optimize/',
            {
                'source_filename': 'non-canonical-success.xlsx',
                'note': 'reverse check purchase sets',
                'input_lines': non_canonical_lines,
                'supplier_config': self.snapshot['supplier_config'],
                'optimization_config': self.snapshot['optimization_config'],
            },
            format='json',
        )
        self.assertEqual(response.status_code, 201, response.content)
        body = response.json()
        result = body['result_payload']

        self.assertFalse(result['canonical_match'])
        self.assertTrue(result['selected_plan']['all_ncc_passed'])
        self.assertTrue(result['selected_plan']['meets_target_and_cap'])
        self.assertTrue(result['reverse_check_rows'])
        for row in result['reverse_check_rows']:
            self.assertTrue(row['consistent'])
            self.assertEqual(int(row['sum_line_allocated_sets']), int(row['total_sets_purchase_spec']))
            self.assertAlmostEqual(
                float(row['reverse_total_length_cm']),
                float(row['spec_total_length_cm']),
                places=2,
            )

        export_response = self.client.get(f"/api/production/paper-optimizer/runs/{body['id']}/export_excel/")
        self.assertEqual(export_response.status_code, 200)
        workbook = load_workbook(BytesIO(export_response.content), read_only=True, data_only=True)
        self.assert_export_workbook_sheet_order(workbook)
        reverse_sheet = workbook['Kiem_tra_phan_bo_nguoc']
        for row in reverse_sheet.iter_rows(min_row=2, values_only=True):
            if not any(value not in (None, '') for value in row):
                continue
            self.assertEqual(int(row[3] or 0), int(row[4] or 0))
            self.assertAlmostEqual(float(row[5] or 0), float(row[6] or 0), places=2)
            self.assertTrue(str(row[10] or '').strip().lower().startswith('c'))

    def test_selected_plan_prefers_cost_then_raw_widths_then_specs(self):
        alternatives = [
            {
                'plan_code': 'LOWEST_TOTAL_COST',
                'all_ncc_passed': True,
                'meets_target_and_cap': True,
                'total_converted_cost': 10.0,
                'unique_raw_width_count': 3,
                'unique_purchase_spec_count': 3,
                'final_group_count': 3,
                'technical_waste_rate': 0.10,
            },
            {
                'plan_code': 'MIN_RAW_WIDTHS',
                'all_ncc_passed': True,
                'meets_target_and_cap': True,
                'total_converted_cost': 10.0,
                'unique_raw_width_count': 2,
                'unique_purchase_spec_count': 4,
                'final_group_count': 4,
                'technical_waste_rate': 0.12,
            },
            {
                'plan_code': 'MIN_SPECS',
                'all_ncc_passed': True,
                'meets_target_and_cap': True,
                'total_converted_cost': 10.0,
                'unique_raw_width_count': 2,
                'unique_purchase_spec_count': 2,
                'final_group_count': 5,
                'technical_waste_rate': 0.14,
            },
        ]

        self.assertEqual(select_best_plan_legacy(copy.deepcopy(alternatives)), 'MIN_SPECS')
        self.assertEqual(select_best_plan_v2(copy.deepcopy(alternatives)), 'MIN_SPECS')

    def test_selected_plan_prefers_lowest_total_cost_when_metrics_converge(self):
        alternatives = [
            {
                'plan_code': 'EXACT_ORDER',
                'all_ncc_passed': True,
                'meets_target_and_cap': True,
                'total_converted_cost': 9.5,
                'unique_raw_width_count': 1,
                'unique_purchase_spec_count': 4,
                'final_group_count': 6,
                'technical_waste_rate': 0.12,
            },
            {
                'plan_code': 'LOWEST_TOTAL_COST',
                'all_ncc_passed': True,
                'meets_target_and_cap': True,
                'total_converted_cost': 9.5,
                'unique_raw_width_count': 1,
                'unique_purchase_spec_count': 4,
                'final_group_count': 6,
                'technical_waste_rate': 0.12,
            },
            {
                'plan_code': 'MIN_RAW_WIDTHS',
                'all_ncc_passed': True,
                'meets_target_and_cap': True,
                'total_converted_cost': 9.5,
                'unique_raw_width_count': 1,
                'unique_purchase_spec_count': 4,
                'final_group_count': 6,
                'technical_waste_rate': 0.12,
            },
            {
                'plan_code': 'MIN_SPECS',
                'all_ncc_passed': True,
                'meets_target_and_cap': True,
                'total_converted_cost': 9.5,
                'unique_raw_width_count': 1,
                'unique_purchase_spec_count': 4,
                'final_group_count': 6,
                'technical_waste_rate': 0.12,
            },
        ]

        self.assertEqual(select_best_plan_legacy(copy.deepcopy(alternatives)), 'LOWEST_TOTAL_COST')
        self.assertEqual(select_best_plan_v2(copy.deepcopy(alternatives)), 'LOWEST_TOTAL_COST')

    @patch('paper_optimizer.optimizer.build_purchase_plan')
    def test_legacy_min_raw_widths_keeps_cost_plan_when_candidate_does_not_reduce_widths(self, mock_build_purchase_plan):
        plan_exact = {
            'all_ncc_passed': True,
            'meets_target_and_cap': True,
            'total_converted_cost': 10.8,
            'unique_raw_width_count': 2,
            'unique_purchase_spec_count': 3,
            'final_group_count': 3,
            'technical_waste_rate': 0.18,
            'purchase_spec_rows': [{'group_id': 'exact'}],
            'allocation_details': [{'line_id': 'exact'}],
            'reverse_check_rows': [],
            'leftovers': [],
            'final_plan': [{'description': 'exact'}],
            'target_requested_quantity_total': 100,
            'production_reserve_quantity_total': 0,
        }
        plan_cost = {
            'all_ncc_passed': True,
            'meets_target_and_cap': True,
            'total_converted_cost': 10.0,
            'unique_raw_width_count': 2,
            'unique_purchase_spec_count': 3,
            'final_group_count': 3,
            'technical_waste_rate': 0.18,
            'purchase_spec_rows': [{'group_id': 'cost'}],
            'allocation_details': [{'line_id': 'cost'}],
            'reverse_check_rows': [],
            'leftovers': [],
            'final_plan': [{'description': 'cost'}],
            'target_requested_quantity_total': 100,
            'production_reserve_quantity_total': 0,
        }
        plan_widths_candidate = {
            'all_ncc_passed': True,
            'meets_target_and_cap': True,
            'total_converted_cost': 10.05,
            'unique_raw_width_count': 2,
            'unique_purchase_spec_count': 4,
            'final_group_count': 4,
            'technical_waste_rate': 0.09,
            'purchase_spec_rows': [{'group_id': 'width-candidate'}],
            'allocation_details': [{'line_id': 'width-candidate'}],
            'reverse_check_rows': [],
            'leftovers': [],
            'final_plan': [{'description': 'width-candidate'}],
            'target_requested_quantity_total': 100,
            'production_reserve_quantity_total': 0,
        }
        plan_specs = copy.deepcopy(plan_cost)
        mock_build_purchase_plan.side_effect = [plan_exact, plan_cost, plan_widths_candidate, plan_specs]

        alternatives, _ = build_four_public_plans(
            target_lines=[{'id': 'L1', 'quantity': 100, 'target_quantity': 100}],
            supplier_config={'available_raw_widths_cm': [110], 'trim_edge_cm': 0, 'min_purchase_length_cm': 5000},
            optimization_config={
                'allow_economic_overproduction': False,
                'allow_exceed_target_demand': False,
                'acceptable_cost_increase_for_fewer_raw_widths_percent': 1.0,
                'acceptable_cost_increase_for_fewer_specs_percent': 0.5,
            },
        )

        width_plan = next(item for item in alternatives if item['plan_code'] == 'MIN_RAW_WIDTHS')
        self.assertEqual(width_plan['total_converted_cost'], plan_cost['total_converted_cost'])
        self.assertEqual(width_plan['unique_raw_width_count'], plan_cost['unique_raw_width_count'])
        self.assertEqual(width_plan['unique_purchase_spec_count'], plan_cost['unique_purchase_spec_count'])

    @patch('paper_optimizer.optimizer.build_purchase_plan')
    def test_legacy_min_specs_keeps_cost_plan_when_candidate_adds_raw_widths(self, mock_build_purchase_plan):
        plan_exact = {
            'all_ncc_passed': True,
            'meets_target_and_cap': True,
            'total_converted_cost': 10.8,
            'unique_raw_width_count': 1,
            'unique_purchase_spec_count': 4,
            'final_group_count': 4,
            'technical_waste_rate': 0.18,
            'purchase_spec_rows': [{'group_id': 'exact'}],
            'allocation_details': [{'line_id': 'exact'}],
            'reverse_check_rows': [],
            'leftovers': [],
            'final_plan': [{'description': 'exact'}],
            'target_requested_quantity_total': 100,
            'production_reserve_quantity_total': 0,
        }
        plan_cost = {
            'all_ncc_passed': True,
            'meets_target_and_cap': True,
            'total_converted_cost': 10.0,
            'unique_raw_width_count': 1,
            'unique_purchase_spec_count': 4,
            'final_group_count': 4,
            'technical_waste_rate': 0.18,
            'purchase_spec_rows': [{'group_id': 'cost'}],
            'allocation_details': [{'line_id': 'cost'}],
            'reverse_check_rows': [],
            'leftovers': [],
            'final_plan': [{'description': 'cost'}],
            'target_requested_quantity_total': 100,
            'production_reserve_quantity_total': 0,
        }
        plan_widths = copy.deepcopy(plan_cost)
        plan_specs_candidate = {
            'all_ncc_passed': True,
            'meets_target_and_cap': True,
            'total_converted_cost': 10.04,
            'unique_raw_width_count': 2,
            'unique_purchase_spec_count': 3,
            'final_group_count': 3,
            'technical_waste_rate': 0.08,
            'purchase_spec_rows': [{'group_id': 'spec-candidate'}],
            'allocation_details': [{'line_id': 'spec-candidate'}],
            'reverse_check_rows': [],
            'leftovers': [],
            'final_plan': [{'description': 'spec-candidate'}],
            'target_requested_quantity_total': 100,
            'production_reserve_quantity_total': 0,
        }
        mock_build_purchase_plan.side_effect = [plan_exact, plan_cost, plan_widths, plan_specs_candidate]

        alternatives, _ = build_four_public_plans(
            target_lines=[{'id': 'L1', 'quantity': 100, 'target_quantity': 100}],
            supplier_config={'available_raw_widths_cm': [110, 115], 'trim_edge_cm': 0, 'min_purchase_length_cm': 5000},
            optimization_config={
                'allow_economic_overproduction': False,
                'allow_exceed_target_demand': False,
                'acceptable_cost_increase_for_fewer_raw_widths_percent': 1.0,
                'acceptable_cost_increase_for_fewer_specs_percent': 0.5,
            },
        )

        spec_plan = next(item for item in alternatives if item['plan_code'] == 'MIN_SPECS')
        self.assertEqual(spec_plan['total_converted_cost'], plan_cost['total_converted_cost'])
        self.assertEqual(spec_plan['unique_raw_width_count'], plan_cost['unique_raw_width_count'])
        self.assertEqual(spec_plan['unique_purchase_spec_count'], plan_cost['unique_purchase_spec_count'])

    def test_v2_objective_pick_prefers_lower_cost_before_waste_when_dimensions_match(self):
        pool = [
            {
                'plan_code': 'MIN_RAW_WIDTHS',
                'all_ncc_passed': True,
                'meets_target_and_cap': True,
                'total_converted_cost': 10.0,
                'unique_raw_width_count': 2,
                'unique_purchase_spec_count': 3,
                'final_group_count': 3,
                'technical_waste_rate': 0.2,
            },
            {
                'plan_code': 'MIN_RAW_WIDTHS',
                'all_ncc_passed': True,
                'meets_target_and_cap': True,
                'total_converted_cost': 10.1,
                'unique_raw_width_count': 2,
                'unique_purchase_spec_count': 3,
                'final_group_count': 3,
                'technical_waste_rate': 0.1,
            },
        ]

        best, stats = pick_plan_v2(pool, 'MIN_RAW_WIDTHS')
        self.assertEqual(best['total_converted_cost'], 10.0)
        self.assertEqual(stats['objective_candidate_count'], 2)

    @patch('paper_optimizer.optimizer.build_four_public_plans')
    @patch('paper_optimizer.optimizer.build_four_public_plans_v2')
    def test_non_canonical_dataset_records_legacy_fallback_metadata(self, mock_v2, mock_legacy):
        changed_lines = [dict(item) for item in self.snapshot['input_lines']]
        changed_lines[0]['quantity'] = changed_lines[0]['quantity'] + 1
        mock_v2.return_value = (build_non_viable_public_alternatives(), 'LOWEST_TOTAL_COST')
        mock_legacy.return_value = (
            build_viable_legacy_fallback_alternatives(self.snapshot),
            self.snapshot['selected_plan_code'],
        )

        response = self.client.post(
            '/api/production/paper-optimizer/runs/optimize/',
            {
                'source_filename': 'legacy-fallback.xlsx',
                'note': 'legacy fallback test',
                'input_lines': changed_lines,
                'supplier_config': self.snapshot['supplier_config'],
                'optimization_config': self.snapshot['optimization_config'],
            },
            format='json',
        )
        self.assertEqual(response.status_code, 201, response.content)
        result = response.json()['result_payload']
        stats = result['stats']

        self.assertFalse(result['canonical_match'])
        self.assertEqual(stats['engine_primary'], 'v2')
        self.assertTrue(stats['engine_fallback_used'])
        self.assertEqual(stats['engine_fallback_source'], 'legacy')
        self.assertEqual(stats['engine_fallback_reason_code'], 'no_viable_public_alternatives')
        self.assertEqual(len(result['final_plan_alternatives']), 4)
        self.assertEqual(result['selected_plan_code'], result['selected_plan']['plan_code'])
        self.assertTrue(result['selected_scenario_code'])

        export_response = self.client.get(f"/api/production/paper-optimizer/runs/{response.json()['id']}/export_excel/")
        self.assertEqual(export_response.status_code, 200)
        workbook = load_workbook(BytesIO(export_response.content), read_only=True, data_only=True)
        self.assert_export_workbook_sheet_order(workbook)
        self.assertIn('Phuong_an_mua_cuoi', workbook.sheetnames)
        self.assertIn('Debug_engine', workbook.sheetnames)

    @patch('paper_optimizer.optimizer.build_four_public_plans')
    @patch('paper_optimizer.optimizer.build_four_public_plans_v2')
    def test_non_canonical_dataset_returns_no_feasible_candidate_metadata_when_both_engines_fail(
        self,
        mock_v2,
        mock_legacy,
    ):
        changed_lines = [dict(item) for item in self.snapshot['input_lines']]
        changed_lines[0]['quantity'] = changed_lines[0]['quantity'] + 1
        non_viable = build_non_viable_public_alternatives()
        mock_v2.return_value = (non_viable, 'LOWEST_TOTAL_COST')
        mock_legacy.return_value = (non_viable, 'LOWEST_TOTAL_COST')

        response = self.client.post(
            '/api/production/paper-optimizer/runs/optimize/',
            {
                'source_filename': 'no-feasible.xlsx',
                'note': 'no feasible candidate test',
                'input_lines': changed_lines,
                'supplier_config': self.snapshot['supplier_config'],
                'optimization_config': self.snapshot['optimization_config'],
            },
            format='json',
        )
        self.assertEqual(response.status_code, 201, response.content)
        result = response.json()['result_payload']
        stats = result['stats']

        self.assertEqual(result['selected_plan_code'], 'LOWEST_TOTAL_COST')
        self.assertEqual(result['selected_scenario_code'], 'NO_FEASIBLE')
        self.assertEqual(result['result_state'], 'no_feasible_candidate')
        self.assertTrue(stats['no_feasible_candidate'])
        self.assertFalse(bool(stats['displayable_result']))
        self.assertEqual(stats['feasible_alternative_count'], 0)
        self.assertFalse(stats['engine_fallback_used'])
        self.assertEqual(stats['engine_fallback_source'], 'none')
        self.assertEqual(stats['engine_fallback_reason_code'], 'no_viable_public_alternatives')
        self.assertTrue(bool(result['selected_plan'].get('is_placeholder_no_feasible')))
        self.assertTrue(all(bool(alternative.get('is_placeholder_no_feasible')) for alternative in result['final_plan_alternatives']))

        export_response = self.client.get(f"/api/production/paper-optimizer/runs/{response.json()['id']}/export_excel/")
        self.assertEqual(export_response.status_code, 400)
        self.assertEqual(export_response.json()['detail'], 'Run chưa có phương án mua cuối hợp lệ để xuất workbook.')

    def test_optimizer_regression_corpus_benchmark_smoke(self):
        thresholds = {
            'canonical_locked': 500,
            'non_canonical_success': 10000,
            'no_feasible_candidate': 10000,
            'legacy_fallback_success': 500,
        }
        for case_name in REGRESSION_CASE_ORDER:
            payload = run_regression_case(case_name)
            expected = payload['expected']
            result_payload = payload['result_payload']
            stats = result_payload['stats']
            with self.subTest(case=case_name):
                self.assertEqual(result_payload['selected_plan_code'], expected['selected_plan_code'])
                self.assertEqual(result_payload['selected_scenario_code'], expected['selected_scenario_code'])
                self.assertEqual(bool(result_payload['canonical_match']), expected['canonical_match'])
                self.assertEqual(bool(stats['engine_fallback_used']), expected['engine_fallback_used'])
                self.assertEqual(int(stats['feasible_alternative_count']), expected['feasible_alternative_count'])
                self.assertEqual(bool(stats['no_feasible_candidate']), expected['no_feasible_candidate'])
                if expected.get('engine_fallback_source'):
                    self.assertEqual(stats['engine_fallback_source'], expected['engine_fallback_source'])
                if expected.get('engine_fallback_reason_code'):
                    self.assertEqual(stats['engine_fallback_reason_code'], expected['engine_fallback_reason_code'])
                self.assertLess(
                    payload['duration_ms'],
                    thresholds[case_name],
                    f'{case_name} benchmark too slow: {payload["duration_ms"]}ms',
                )

    def test_benchmark_paper_optimizer_command_outputs_regression_corpus_json(self):
        stdout = StringIO()
        call_command('benchmark_paper_optimizer', '--json', stdout=stdout)
        payload = json.loads(stdout.getvalue())
        self.assertEqual([item['case'] for item in payload], list(REGRESSION_CASE_ORDER))
        for item in payload:
            self.assertIn('duration_ms', item)
            self.assertIn('selected_plan_code', item)
            self.assertIn('selected_scenario_code', item)
            self.assertIn('engine_fallback_used', item)
            self.assertIn('feasible_alternative_count', item)
            self.assertIn('no_feasible_candidate', item)
            self.assertIn('guardrail_passed', item)
            self.assertIn('guardrail_decision_metric', item)
            self.assertIn('benchmark_compare', item)

    def test_benchmark_paper_optimizer_command_supports_strict_mode(self):
        stdout = StringIO()
        call_command('benchmark_paper_optimizer', '--json', '--strict', stdout=stdout)
        payload = json.loads(stdout.getvalue())
        self.assertEqual([item['case'] for item in payload], list(REGRESSION_CASE_ORDER))
        self.assertTrue(all(bool(item['guardrail_passed']) for item in payload))

    def test_optimize_recovery_mode_and_export(self):
        """Backward compat alias – gọi cùng test_optimize_non_canonical_dataset_uses_full_optimizer."""
        self.test_optimize_non_canonical_dataset_uses_full_optimizer()

    def test_canonical_dataset_workbook_purchase_spec_ncc_check(self):
        """Kiểm workbook canonical: Phuong_an_mua_cuoi phải có cột đúng và Tổng chiều dài = Dài × Bộ."""
        response = self.client.post(
            '/api/production/paper-optimizer/runs/optimize/',
            {
                'source_filename': 'canonical.xlsx',
                'note': 'ncc check test',
                'input_lines': self.snapshot['input_lines'],
                'supplier_config': self.snapshot['supplier_config'],
                'optimization_config': self.snapshot['optimization_config'],
            },
            format='json',
        )
        self.assertEqual(response.status_code, 201, response.content)
        run_id = response.json()['id']

        export_response = self.client.get(f'/api/production/paper-optimizer/runs/{run_id}/export_excel/')
        self.assertEqual(export_response.status_code, 200)
        workbook = load_workbook(BytesIO(export_response.content), read_only=True, data_only=True)
        self.assert_export_workbook_sheet_order(workbook)
        self.assertIn('Phuong_an_mua_cuoi', workbook.sheetnames)
        self.assertIn('To_hop_de_xuat', workbook.sheetnames)

        sheet = workbook['Phuong_an_mua_cuoi']
        headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        self.assertIn('Tổng số bộ mua thật', headers)
        self.assertIn('Tổng chiều dài (cm)', headers)
        self.assertIn('Đạt NCC', headers)

        benchmark_sheet = workbook['So_sanh_voi_moc_tham_chieu']
        benchmark_headers = [cell.value for cell in next(benchmark_sheet.iter_rows(min_row=1, max_row=1))]
        self.assertEqual(benchmark_headers[1], 'Benchmark tham chiếu')

        idx_dai = headers.index('Dài mua (cm)')
        idx_sets = headers.index('Tổng số bộ mua thật')
        idx_total = headers.index('Tổng chiều dài (cm)')
        idx_required = headers.index('Min chiều dài NCC yêu cầu (cm)')
        idx_actual = headers.index('Chiều dài thực tế (cm)')
        idx_ncc = headers.index('Đạt NCC')
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if row[idx_dai] is None:
                break
            dai = float(row[idx_dai] or 0)
            sets = int(row[idx_sets] or 0)
            total_len = float(row[idx_total] or 0)
            expected = round(dai * sets, 2)
            self.assertAlmostEqual(total_len, expected, places=2,
                msg=f'Công thức Tổng chiều dài sai: {dai} × {sets} ≠ {total_len}')
            required_len = float(row[idx_required] or 0)
            actual_len = float(row[idx_actual] or 0)
            self.assertGreater(required_len, 0)
            self.assertAlmostEqual(actual_len, total_len, places=2)
            expected_ncc = 'Đạt' if actual_len >= required_len else 'Chưa đạt'
            self.assertEqual(row[idx_ncc], expected_ncc)

    def test_regression_non_canonical_success_workbook_sheet_order(self):
        payload = copy.deepcopy(run_regression_case('non_canonical_success')['result_payload'])
        run = PaperOptimizerRun.objects.create(
            code='POPT-NON-CANONICAL-WORKBOOK',
            status=PaperOptimizerRun.STATUS_SUCCESS,
            source_filename='non-canonical-success.xlsx',
            note='deterministic non-canonical workbook order',
            input_lines=payload.get('input_lines') or [],
            supplier_config=payload.get('supplier_config') or self.snapshot['supplier_config'],
            optimization_config=payload.get('optimization_config') or self.snapshot['optimization_config'],
            preview_rows=payload.get('input_lines') or [],
            result_payload=payload,
            recovery_mode=bool(payload.get('recovery_mode')),
            canonical_match=bool(payload.get('canonical_match')),
            created_by=self.user,
        )

        export_response = self.client.get(f'/api/production/paper-optimizer/runs/{run.id}/export_excel/')
        self.assertEqual(export_response.status_code, 200, export_response.content)
        workbook = load_workbook(BytesIO(export_response.content), read_only=True, data_only=True)
        self.assert_export_workbook_sheet_order(workbook)

        sheet = workbook['Phuong_an_mua_cuoi']
        data_rows = [
            row for row in sheet.iter_rows(min_row=2, values_only=True)
            if any(value not in (None, '') for value in row)
        ]
        self.assertGreater(len(data_rows), 0)

    def test_upload_preview_and_manual_pattern(self):
        csv_payload = (
            "width_cm,length_cm,quantity,note\n"
            "52,112,25,hop A\n"
            "55.5,145.5,,thieu so luong\n"
            "84,143.5,1.5,so luong le\n"
            ",,,\n"
            "abc,120,8,kho loi\n"
        )
        upload_file = SimpleUploadedFile('preview.csv', csv_payload.encode('utf-8'), content_type='text/csv')
        upload_response = self.client.post(
            '/api/production/paper-optimizer/runs/upload_preview/',
            {'file': upload_file},
        )
        self.assertEqual(upload_response.status_code, 200, upload_response.content)
        upload_body = upload_response.json()
        self.assertEqual(upload_body['summary']['line_count'], 1)
        self.assertEqual(upload_body['summary']['accepted_row_count'], 1)
        self.assertEqual(upload_body['summary']['total_quantity'], 25)
        self.assertEqual(upload_body['summary']['issue_count'], 3)
        self.assertEqual(upload_body['summary']['invalid_row_count'], 3)
        self.assertEqual(upload_body['summary']['skipped_empty_row_count'], 1)
        self.assertTrue(upload_body['summary']['header_detected'])
        self.assertEqual(upload_body['rows'][0]['source_row_number'], 2)
        self.assertEqual(
            [issue['row_number'] for issue in upload_body['issues']],
            [3, 4, 6],
        )
        self.assertEqual(
            [issue['code'] for issue in upload_body['issues']],
            ['missing_required_fields', 'invalid_quantity', 'invalid_width_cm'],
        )

        eval_response = self.client.post(
            '/api/production/paper-optimizer/runs/evaluate_manual_pattern/',
            {
                'raw_width_cm': 250,
                'trim_edge_cm': 2,
                'run_length_cm': 145.5,
                'sets': 13,
                'components': [
                    {'width_cm': 52, 'length_cm': 112, 'multiplier': 2},
                    {'width_cm': 55.5, 'length_cm': 145.5, 'multiplier': 1},
                    {'width_cm': 84, 'length_cm': 143.5, 'multiplier': 1},
                ],
            },
            format='json',
        )
        self.assertEqual(eval_response.status_code, 200, eval_response.content)
        self.assertIn('waste_rate', eval_response.json())
